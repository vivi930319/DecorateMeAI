from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"C:\Users\isach\OneDrive\桌面\web_frontend\docs\project_execution_plan.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(9)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Microsoft JhengHei"
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
                    r.font.size = Pt(9)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "365F91")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(55, 55, 55)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def set_document_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Heading 1", 18, "1F4E79"),
        ("Heading 2", 14, "365F91"),
        ("Heading 3", 12, "4F81BD"),
    ]:
        st = styles[name]
        st.font.name = "Microsoft JhengHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)


def add_footer(doc):
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer.add_run("Decorate Me 專題整合文件｜BASIC 資料集、模型訓練與 API 串接計畫")
        r.font.name = "Microsoft JhengHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(120, 120, 120)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    style_table(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF2F8")
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.name = "Microsoft JhengHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string("1F4E79")
    r2 = p.add_run(body)
    r2.font.name = "Microsoft JhengHei"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    r2.font.size = Pt(9.5)
    doc.add_paragraph()


def main():
    doc = Document()
    set_document_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Decorate Me 專題整合進度與後續執行計畫")
    r.font.name = "Microsoft JhengHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    r.font.size = Pt(22)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("1F4E79")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("網頁前端、臉部分析 BASIC/PRO、BASIC 資料集、模型訓練、API 串接與七人分工")
    sr.font.name = "Microsoft JhengHei"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    sr.font.size = Pt(11)
    sr.font.color.rgb = RGBColor(90, 90, 90)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("版本日期：2026-06-13｜用途：組內進度會議、重新排程、API 串接協調")
    mr.font.name = "Microsoft JhengHei"
    mr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    mr.font.size = Pt(9.5)
    mr.font.color.rgb = RGBColor(110, 110, 110)

    add_callout(
        doc,
        "文件核心結論",
        "目前最合理的方向是先完成 BASIC 版：以正面臉圖為主，建立 raw images + labels.csv 的多欄位資料集，先穩定臉型與鼻型，再串接 Ollama 文字建議、Replicate 圖片渲染與 PostgreSQL 會員商品資料。PRO 多角度精細分類先保留入口，不要讓它拖慢 BASIC 完成度。"
    )

    doc.add_heading("1. 專題目前總架構", level=1)
    doc.add_paragraph("本專題目前分成五個主要模組，各模組由不同組員負責，最後都需要接回網頁版前端。前端不應直接連 PostgreSQL、Replicate 或 Ollama 原始服務，而是呼叫各組員包好的 API。")
    add_code_block(doc, """使用者
  ↓
網頁前端
  ↓
臉部分析端 BASIC / PRO
  ↓
analysisPackage
  ├─ Ollama 文字建議端
  ├─ Replicate 圖片渲染端
  ├─ 商品推薦端
  └─ PostgreSQL 會員 / 收藏 / 紀錄端""")

    add_table(doc, ["模組", "目前狀態", "主要負責"], [
        ["網頁前端", "已完成 BASIC/PRO UI、拍照/上傳、妝容建議、商品、收藏、分析紀錄、渲染圖預留", "你"],
        ["臉部分析端", "BASIC 可用，PRO 多角度入口已預留；鼻型已保守化", "你"],
        ["Ollama 文字建議", "由組員電腦提供 Gemma3 / LLaVA 建議 API", "Ollama 組員"],
        ["Replicate 渲染", "由組員電腦呼叫 Replicate，需回永久圖片 URL", "渲染組員"],
        ["商品端", "負責商品爬蟲、商品資料與推薦邏輯", "商品組員"],
        ["資料庫端", "負責會員、收藏、分析紀錄、生成圖片 metadata", "DB 組員"],
        ["資料/測試機動", "負責 BASIC 圖片收集、標註、API 測試與文件", "兩位機動組員"],
    ], widths=[3.2, 9.2, 3.2])

    doc.add_heading("2. 目前已完成進度", level=1)
    doc.add_heading("2.1 網頁前端", level=2)
    add_bullets(doc, [
        "前端位置：C:\\Users\\isach\\OneDrive\\桌面\\web_frontend。",
        "GitHub 分支：dev_makeup；前一版已推送，最新妝容建議渲染圖區塊仍需再推一次。",
        "已完成登入、註冊、訪客登入、會員中心、商品推薦、收藏、分析紀錄、妝容建議與妝容對比圖。",
        "臉部分析頁已拆 BASIC / PRO。BASIC 支援上傳與鏡頭拍照；PRO 預留 front、left45、right45、side 多角度上傳。",
        "分析完成後建立 analysisPackage，包含圖片資訊、臉部分析 JSON、生成式 AI 文字建議欄位、渲染結果欄位與推薦商品欄位。",
        "妝容建議頁已加入渲染後照片區；未來只要寫入 render.afterImageUrl 或 render.afterImageDataUrl 即可顯示。",
    ])

    doc.add_heading("2.2 臉部分析端", level=2)
    add_bullets(doc, [
        "後端位置：C:\\Users\\isach\\PycharmProjects\\PythonProject12。",
        "GitHub 分支：Isa；最新已推送 commit：c92361e。",
        "BASIC API：POST /analyze，預設 http://127.0.0.1:8001/analyze。",
        "PRO API：POST /analyze-pro，預設 http://127.0.0.1:8002/analyze-pro。",
        "BASIC 回傳臉型、眉型、眼型、鼻型、嘴型、膚色、嘴唇 LAB。",
        "PRO 目前先接收多角度照片，正面照沿用 BASIC 分析，側面與 45 度精細分類尚未正式演算法化。",
        "鼻型 BASIC 已改成正面可穩定判斷的標準鼻、寬鼻、窄鼻，不再以正面照判斷鷹勾鼻、塌鼻、朝天鼻。",
        "後端內部仍有 MAX_IMAGE_SIZE = 1280 的縮圖限制，後續建議改成環境變數，預設先測 2048。",
    ])

    doc.add_heading("3. BASIC 資料集策略", level=1)
    add_callout(
        doc,
        "最終資料策略",
        "不要先把圖片分到五官資料夾，也不要複製同一張圖到多個類別。BASIC 使用 raw_images + labels.csv：同一張正面照可以同時標臉型、鼻型、眼型、眉型、嘴型。"
    )
    doc.add_paragraph("資料夾建議如下：")
    add_code_block(doc, """dataset/
  basic/
    raw_images/
      raw_000001.jpg
      raw_000002.jpg
    images/
      img_000001.jpg
      img_000002.jpg
    labels.csv
    label_map.json""")
    doc.add_paragraph("raw_images 保存原始收集圖；images 保存後續統一對齊、裁切、resize 成 320x320 的正面臉圖。labels.csv 是訓練和準確度統計的核心。")

    doc.add_heading("3.1 標籤欄位", level=2)
    add_code_block(doc, "image_id,file_path,face_shape,nose_front,eye_shape,brow_shape,lip_shape,quality,note")

    add_table(doc, ["大類", "英文欄位", "類別", "中文顯示", "每類有效標籤目標"], [
        ["臉型", "face_shape", "round", "圓臉", "50"],
        ["臉型", "face_shape", "oval", "鵝蛋臉", "50"],
        ["臉型", "face_shape", "square", "方臉", "50"],
        ["臉型", "face_shape", "long", "長臉", "50"],
        ["臉型", "face_shape", "heart", "心型臉 / 瓜子臉", "50"],
        ["鼻型", "nose_front", "standard", "標準鼻", "50"],
        ["鼻型", "nose_front", "wide", "寬鼻", "50"],
        ["鼻型", "nose_front", "narrow", "窄鼻", "50"],
        ["眼型", "eye_shape", "almond", "杏仁眼", "50"],
        ["眼型", "eye_shape", "round", "圓眼", "50"],
        ["眼型", "eye_shape", "peach_blossom", "桃花眼", "50"],
        ["眼型", "eye_shape", "monolid", "單眼皮", "50"],
        ["眼型", "eye_shape", "double_eyelid", "雙眼皮", "50"],
        ["眉型", "brow_shape", "straight", "平眉", "50"],
        ["眉型", "brow_shape", "curved", "彎月眉", "50"],
        ["眉型", "brow_shape", "willow", "柳葉眉", "50"],
        ["眉型", "brow_shape", "arched", "挑眉", "50"],
        ["嘴型", "lip_shape", "standard", "標準唇", "50"],
        ["嘴型", "lip_shape", "smile", "微笑唇", "50"],
        ["嘴型", "lip_shape", "thick", "厚唇", "50"],
        ["嘴型", "lip_shape", "thin", "薄唇", "50"],
        ["嘴型", "lip_shape", "flower_petal", "花瓣唇", "50"],
    ], widths=[2.1, 3.2, 3.3, 4.1, 2.2])

    doc.add_heading("3.2 圖片數量規劃", level=2)
    add_bullets(doc, [
        "每一類目標 50 筆有效標籤，總標籤需求約 1100 筆。",
        "不代表一定要 1100 張圖片，因為同一張圖片可以同時提供 5 個欄位標籤。",
        "建議先海量收 500 張 raw images，經過篩選、unknown、遮擋剔除後，比較有機會達成各類 50 筆有效標籤。",
        "若時間不足，優先確保臉型與鼻型滿 50 筆，眼眉唇可先達到每類 30 筆，再逐步補齊。",
    ])

    add_table(doc, ["項目", "目標量", "原因"], [
        ["raw_images 原始收圖", "500 張", "會有模糊、遮擋、側臉、重複與 unknown 被剔除"],
        ["裁切後 images", "約 300-450 張", "可用於訓練或規則校正的標準臉圖"],
        ["臉型有效標籤", "250 筆", "5 類 × 50"],
        ["鼻型有效標籤", "150 筆", "3 類 × 50"],
        ["眼型有效標籤", "250 筆", "5 類 × 50"],
        ["眉型有效標籤", "200 筆", "4 類 × 50"],
        ["嘴型有效標籤", "250 筆", "5 類 × 50"],
    ], widths=[4.2, 3.2, 8.5])

    doc.add_heading("4. 標記圖片流程", level=1)
    doc.add_paragraph("標記工具建議使用 Label Studio。組員不一定需要安裝環境，可以由一台電腦啟動 Label Studio，其他組員使用瀏覽器連進來標註。")
    add_numbered(doc, [
        "收圖者先將圖片統一放入 dataset/basic/raw_images。",
        "初篩：刪除或移到 rejected 的圖片包含側臉、模糊、遮臉、多人合照、過度美顏、低解析度。",
        "將 raw_images 匯入 Label Studio。",
        "每張圖標 face_shape、nose_front、eye_shape、brow_shape、lip_shape、quality、note。",
        "不確定的欄位填 unknown，不要硬猜。",
        "匯出 Label Studio JSON。",
        "轉成 labels.csv。",
        "用 Python + MediaPipe / InsightFace 統一裁切正面臉成 320x320，輸出到 images。",
    ])
    add_callout(
        doc,
        "標註品質原則",
        "BASIC 的準確度上限很大程度取決於標籤品質。與其亂標 500 張，不如乾淨標 300 張。尤其鼻型 BASIC 只標 standard / wide / narrow，不要把塌鼻、鷹勾鼻放進 BASIC。"
    )

    doc.add_heading("5. 最終模型與訓練策略", level=1)
    doc.add_paragraph("BASIC 最終版建議使用一個 backbone + 五個分類 head 的多任務分類模型。輸入為標準化正面臉圖，不先裁眼睛、鼻子、嘴巴。")
    add_code_block(doc, """Input: 320x320 正面臉圖
Backbone: EfficientNet-B0 或 MobileNetV3
Outputs:
  face_shape head
  nose_front head
  eye_shape head
  brow_shape head
  lip_shape head""")
    add_table(doc, ["決策", "建議", "原因"], [
        ["模型工具", "PyTorch + timm", "可用預訓練模型 fine-tune，不需要從零訓練 CNN"],
        ["第一候選", "EfficientNet-B0", "準確度與大小平衡，適合第一版"],
        ["輕量候選", "MobileNetV3", "若部署速度不足，可改用較小模型"],
        ["輸入尺寸", "320x320", "五官細節比一般分類更細，224 可能偏小"],
        ["資料格式", "images + labels.csv", "支援同圖多標籤與多任務訓練"],
        ["第一階段", "先測現有規則法", "每類 50 張資料量仍偏少，先用來校正 threshold 比較務實"],
        ["第二階段", "fine-tune 臉型與鼻型", "這兩項是目前準確度痛點"],
        ["第三階段", "擴充眼眉唇", "眼眉唇受妝容、表情、遮擋影響較大"],
    ], widths=[3.0, 4.2, 8.5])

    doc.add_heading("6. 七人分工", level=1)
    add_table(doc, ["角色", "主要負責", "本週交付物", "注意事項"], [
        ["你", "網頁前端、臉部分析、整合規格", "前端 latest push、analysisPackage、BASIC/PRO API、整合流程", "你是規格與流程總控，不需要一個人做完所有功能"],
        ["Ollama 組員", "Gemma3 / LLaVA 文字建議", "GET /health、POST /suggest、request/response 範例", "不要要求前端直接打 11434，需包成自己的 API"],
        ["商品組員", "商品爬蟲與推薦演算法", "商品 JSON、GET /products、POST /recommend-products", "商品欄位需固定，包含 id/name/category/price/imageUrl/reason"],
        ["資料庫組員", "PostgreSQL、會員、收藏、紀錄", "會員 API、收藏 API、analysis_records、generated_images schema", "前端不可直接連 PostgreSQL，需提供後端 API"],
        ["Replicate 組員", "圖片渲染與永久圖片儲存", "GET /health、POST /render、afterImageUrl", "不能只回 Replicate 臨時 URL，需下載保存後回永久 URL"],
        ["機動 1", "BASIC 資料集與 Label Studio", "raw_images 500 張目標、labels.csv 第一版、每類統計", "這個角色不要叫機動，實際是資料集負責人"],
        ["機動 2", "API 測試、串接紀錄、文件支援", "api_contract.md、每日 /health 測試表、demo_flow", "不是替大家寫 API，而是確認別人真的接得起來"],
    ], widths=[2.1, 3.8, 5.5, 5.0])

    doc.add_heading("7. API 串接規格", level=1)
    doc.add_paragraph("所有提供 API 的組員都要提供 /health、正式 endpoint、request 範例、response 範例、錯誤格式與是否需要 token。")

    doc.add_heading("7.1 臉部分析 API", level=2)
    add_code_block(doc, """BASIC
POST http://你的IP或localhost:8001/analyze
Content-Type: multipart/form-data
file: image

PRO
POST http://你的IP或localhost:8002/analyze-pro
Content-Type: multipart/form-data
front: image
left45: image optional
right45: image optional
side: image optional""")

    doc.add_heading("7.2 Ollama 文字建議 API", level=2)
    add_code_block(doc, """GET /health
POST /suggest

Request:
{
  "analysisPackageId": "AN-xxx",
  "style": "Soft Baddie",
  "faceAnalysis": {},
  "compressedImage": null
}

Response:
{
  "status": "completed",
  "model": "gemma3",
  "analysisPackageId": "AN-xxx",
  "suggestion": {
    "base": "...",
    "brow": "...",
    "eye": "...",
    "blush": "...",
    "lip": "..."
  }
}""")

    doc.add_heading("7.3 Replicate 渲染 API", level=2)
    add_code_block(doc, """GET /health
POST /render

Request:
{
  "analysisPackageId": "AN-xxx",
  "style": "Soft Baddie",
  "faceAnalysis": {},
  "generativeText": {},
  "image": {
    "compressedDataUrl": "base64 or url"
  }
}

Response:
{
  "status": "completed",
  "analysisPackageId": "AN-xxx",
  "replicateTempUrl": "https://replicate.delivery/xxx.png",
  "savedImageId": 12,
  "afterImageUrl": "http://組員IP:9002/uploads/rendered/AN-xxx-after.png"
}""")

    doc.add_heading("7.4 商品與資料庫 API", level=2)
    add_table(doc, ["模組", "Endpoint", "用途"], [
        ["會員", "POST /auth/register", "註冊會員"],
        ["會員", "POST /auth/login", "登入並取得使用者狀態 / token"],
        ["會員", "GET /users/me", "取得目前會員資料"],
        ["商品", "GET /products", "取得商品列表"],
        ["商品", "GET /products/:id", "取得商品詳情"],
        ["商品", "POST /recommend-products", "根據 analysisPackage / style 推薦商品"],
        ["收藏", "GET /favorites", "取得收藏清單"],
        ["收藏", "POST /favorites", "新增收藏"],
        ["分析紀錄", "POST /analysis-records", "儲存 analysisPackage 或摘要"],
        ["分析紀錄", "GET /analysis-records", "取得歷史紀錄"],
        ["圖片紀錄", "POST /generated-images", "儲存渲染圖片 metadata"],
    ], widths=[3, 5, 8])

    doc.add_heading("8. 串接會議制度", level=1)
    doc.add_paragraph("因為所有模組最後都要接回網頁版前端，所以需要固定串接時段。這不是額外報告，而是實際測 API。")
    add_bullets(doc, [
        "每天 30 分鐘 API 小串接：確認 /health、IP、port、主要 endpoint 是否可連。",
        "每週 1 次完整流程串接：從上傳照片到分析、文字建議、渲染圖、商品推薦、收藏紀錄完整跑一次。",
        "小問題現場修：CORS、欄位名稱、response key、port、JSON 格式。",
        "大問題開任務：模型太慢、Replicate 失敗、DB schema 大改、商品推薦策略未完成。",
    ])
    add_table(doc, ["檢查項目", "成功標準", "失敗處理"], [
        ["/health", "從另一台電腦可開啟", "確認 host 是否 0.0.0.0、port、防火牆"],
        ["CORS", "前端瀏覽器可呼叫", "API 加 CORS middleware"],
        ["Request", "欄位和文件一致", "立即改 API 或前端 mapping"],
        ["Response", "固定 status / error 欄位", "統一 response schema"],
        ["圖片 URL", "前端可顯示", "Replicate 臨時 URL 需保存成本地永久 URL"],
        ["Timeout", "失敗不影響前端整體流程", "前端 fallback、API timeout、錯誤訊息"],
    ], widths=[3.5, 6, 6])

    doc.add_heading("9. 後續排程", level=1)
    add_table(doc, ["週次", "主目標", "你", "其他組員交付"], [
        ["第 1 週", "規格與 mock API 打通", "push 前端最新進度、固定 analysisPackage、開 Label Studio、定 API contract", "各組提供 /health、mock endpoint、request/response 範例；資料組開始收圖"],
        ["第 2 週", "真 API 串接", "前端接 /suggest、/render、商品、會員收藏；後端 MAX_IMAGE_SIZE 改可設定", "Ollama 真建議、Replicate 真渲染並回永久圖、DB 可存會員收藏紀錄、商品 API 可讀"],
        ["第 3 週", "BASIC 準確度修正", "跑現有規則法與 labels.csv 對照，調臉型/鼻型 threshold", "資料組補不足類別；測試組產出混淆矩陣與準確率表"],
        ["第 4 週", "展示穩定與報告", "前端畫面、fallback、demo flow 整理", "各 API 固定啟動方式、README、流程圖、測試紀錄"],
    ], widths=[2, 4, 5.2, 5.2])

    doc.add_heading("10. 風險與處理", level=1)
    add_table(doc, ["風險", "影響", "處理策略"], [
        ["BASIC 資料標籤不一致", "模型或規則校正會被錯誤資料拉歪", "使用 label_map、unknown、quality 欄位；你做最終抽查"],
        ["每類資料不足 50", "準確率難以評估", "先海量收 500 張，再用統計補不足類別"],
        ["Ollama 組員電腦沒開", "文字建議流程中斷", "前端保留 mock fallback；串接日先測 /health"],
        ["Replicate URL 過期", "妝後圖之後顯示不了", "渲染端必須下載圖片並回 afterImageUrl 永久路徑"],
        ["DB 直接存圖片 binary", "資料庫變肥、備份慢", "圖片存 uploads/storage，PostgreSQL 存 metadata"],
        ["前端直接連 PostgreSQL/Ollama", "安全與 CORS 問題，正式架構混亂", "所有外部服務都包成 API"],
        ["PRO 太早投入", "拖慢 BASIC 完成度", "PRO 只保留入口，先不做精細演算法"],
        ["API 每人格式不同", "最後整合大量重工", "API contract 先定，測試負責人每日驗證"],
    ], widths=[4.4, 5, 6.2])

    doc.add_heading("11. 立即待辦清單", level=1)
    add_numbered(doc, [
        "你：把前端最新渲染照片區 push 到 dev_makeup。",
        "你：發布 analysisPackage v1 給所有組員。",
        "機動 1：建立 raw_images 收圖規則與 Label Studio 專案。",
        "機動 1：先收 500 張 raw images，第一輪篩圖。",
        "機動 2：建立 API contract 表格並測每個組員 /health。",
        "Ollama 組：提供 /suggest mock 與真模型回傳格式。",
        "Replicate 組：提供 /render mock，確認永久圖片保存方式。",
        "商品組：提供 /products 與 /recommend-products 格式。",
        "DB 組：提供 users、favorites、analysis_records、generated_images 初版 schema。",
        "全組：固定每日 30 分鐘串接檢查，每週一次完整 demo flow。",
    ])

    doc.add_heading("12. 會議簡短講法", level=1)
    doc.add_paragraph("你可以在會議中這樣說：")
    add_callout(
        doc,
        "會議摘要稿",
        "我們目前已完成網頁前端 BASIC/PRO 流程與臉部分析 BASIC/PRO API。接下來先集中完成 BASIC，不急著做 PRO 精細分類。BASIC 資料集採用 raw_images + labels.csv，同一張正面照標臉型、鼻型、眼型、眉型、嘴型，每個形狀目標至少 50 筆有效標籤。各組 API 必須提供 /health、request、response 與錯誤格式。每天會安排短串接，確保前端能接到 Ollama、Replicate、商品與資料庫端，避免最後 demo 前才發現格式不一致。"
    )

    add_footer(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
