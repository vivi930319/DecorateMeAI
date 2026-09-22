"""Surgical recommendation supplement; retain the prior document unchanged."""
from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / '專題系統文件書_妝識你的美_初版格式_整合模型Worker前端演算法與Ollama_2026-09-06.docx'
SOURCE = Path('C:/Users/isach/Downloads/Decorate_Me_智慧彩妝推薦系統_商品推薦技術文件.docx')
OUT = ROOT / '專題系統文件書_妝識你的美_商品推薦細節補齊_待版面驗證_2026-09-08.docx'
doc, src = Document(BASE), Document(SOURCE)
original = list(doc.paragraphs)

def before(index, text):
    return original[index].insert_paragraph_before(text, style='Normal')

def insert_source(index, source_indices):
    for source_index in source_indices:
        before(index, src.paragraphs[source_index].text)

# Preserve the existing subsection and both supplied diagrams. Complete the
# missing source prose in its existing technical and validation locations.
before(251, '原推薦流程主要依 styleTags、finishTags 與固定類別權重排序。商品資料中 daily 與 natural 標籤覆蓋率過高，使不同膚色或妝容條件仍容易取得相同商品。依演算法端的修改紀錄，新版將評分訊號區分為膚色、四季型、妝容風格、Ollama 建議、臉部特徵與候選多樣性；MAC 與 NW7 不再具有固定優先權，主粉底依實際色差與個人條件選出，其他色號及品牌則列為替代選項。')
insert_source(251, [8])
before(259, '收藏清單、購物車及試妝紀錄本身屬於資料紀錄，不應直接列為推薦結果。會員行為可作為偏好推薦的輸入，但應與目前 Product Service 的臉部分析推薦分開標示。')
insert_source(260, [10])
before(261, 'HSL 妝容風格字典')
insert_source(261, [15])
before(261, '四季型與冷暖底調')
insert_source(261, [17])
before(261, 'NLP 與 Ollama 文字建議')
insert_source(261, [19])
before(261, '可重現的近分輪替')
insert_source(261, [21])
before(263, '逐商品計分與類別重排')
insert_source(263, [25])
before(263, '建議的 NLP 結構化輸出範例')
insert_source(263, [27, 28])
before(263, '上述 JSON 為建議格式，不能據此宣稱目前 Ollama 已啟用完整 JSON Schema 驗證；實際欄位與驗證狀態仍依前節 Ollama 模組的限制說明判讀。')
insert_source(267, [32])
before(271, '商品推薦的限制與後續改善')
insert_source(271, [37])
for i in range(43, 48):
    before(272, src.paragraphs[i].text)

# Source-table wording is already adapted in the target. Keep all six tables
# and their established numbering instead of appending duplicate source tables.
original[262].text = '表中依原技術文件保留各商品類別的計分項目與權重排列。此處呈現的是演算法端記載的設定；逐商品實際採用的項目與分數，仍須由 API 回傳的 scoreBreakdown 核對，不能將設定權重當作推薦准确率或驗收結果。'.replace('准确率', '準確率')

replacements = {
    158: '我們在模型開發初期先以 MediaPipe Face Mesh 與幾何比例建立規則式基線，再比較整臉 CNN 與依五官區域裁切 ROI 的 CNN。開發過程發現隨機切分可能造成身分洩漏，重複圖片亦會使驗證分數失真，因此將 identity split、五折交叉驗證、固定 holdout 與資料去重納入正式流程。',
    159: '模型更換以處理前一階段的限制為依據。規則式方法受閾值影響，可能使部分類別無法出現；整臉 CNN 中的小部位占比偏低，且容易受到背景干擾；MobileNet 在小資料情境下的部分任務仍不穩定；DINOv2 則需考量部署成本與跨部位穩定性。最終選擇 ConvNeXt-Tiny，是基於相同資料切分與訓練條件下的整體平衡。',
    208: '使用者上傳照片後，由前端將請求送至 Gateway。前端負責畫面狀態、欄位組裝、輪詢與錯誤呈現；Gateway 負責登入身分、訪客額度、路由白名單、上游服務驗證、錯誤語意轉換、私人媒體路徑及 Job 串接。通過 Gateway 後，臉部分析服務以 Face Mesh 取得關鍵點，依部位裁切 ROI，再交由對應模型產生臉型、眉型、眼型、鼻型與唇型結果。',
    218: '以下資料來自專案程式、測試 fixture、模型 manifest 與 promotion ledger。每次執行可能變動的 token、時間與雜湊識別碼以動態欄位表示，固定值則保留供來源核對。這些資料分別呈現 Gateway 請求、臉部分析工作及模型版本紀錄。',
    222: '程式契約、測試資料、模型檔案與部署 ledger 代表不同證據層級。測試 fixture 僅支持被測流程與權限情境，manifest 記錄檔案版本與完整性，promotion ledger 則記錄模型升版事件；三者不能視為同一筆線上使用者資料。',
    248: '本系統將 Ollama 定位為個人化妝容文字建議與 Prompt 組合層。其輸入為結構化五官結果、Style 與使用者需求，輸出為可讀建議；Server 端則保留輸入證據、固定 Style 規則、欄位驗證與 fallback。',
    348: '模型工作流記錄各次實驗的問題、選擇依據、執行方式、錯誤及產出檔案，以呈現資料檢查、失敗重跑、切分修正與線上驗證的演進過程。各阶段的完成狀態仍依對應證據判定。'.replace('阶段', '階段'),
    357: '本系統以使用者、管理員與維運者的操作情境規劃 UAT，檢查任務能否完成。測試紀錄依「發現問題、為什麼有問題、改善過後」整理，涵蓋畫面操作、資料、權限、錯誤訊息與後續流程。尚未進行實機驗收的項目仍保留未驗證狀態。',
    361: '每次地端訓練應保留 runId、資料版本、類別版本、訓練參數、metrics、checkpoint、ONNX、類別檔與 hash。下表列出主要成果位置；檔案存在仍不代表已部署，尚須核對 manifest、服務載入版本及對應的 Cloud Run revision。',
    364: '模型部署的完成條件包括資料與類別版本明確、訓練可重現、結果可解釋、模型與類別檔一致、服務載入正確版本、線上 smoke 通過及具備回復上一版的方式。僅完成地端訓練，不能據此宣稱線上使用者已使用新模型。',
}
for i, text in replacements.items():
    original[i].text = text

# User identified the exact sixth subsection to remove. Its complete original
# remains in BASE for slides; do not remove the following methods or Worker.
body = doc.element.body
start, end = list(body).index(original[347]._p), list(body).index(original[350]._p)
for element in list(body)[start:end]:
    body.remove(element)
for p in list(doc.paragraphs):
    if '\t' in p.text and ('六、臉部分析模型訓練工作流' in p.text or '表 5-4' in p.text and '臉部模型訓練工作流' in p.text):
        p._p.getparent().remove(p._p)
        continue
    for old, new in [('七、模型訓練方法與實驗結果','六、模型訓練方法與實驗結果'),('八、UAT 測試工作流','七、UAT 測試工作流'),('九、模型成果檔案與部署驗證','八、模型成果檔案與部署驗證'),('十、Worker 技術細節與完成界線','九、Worker 技術細節與完成界線'),('十一、出版與原始資料來源','十、出版與原始資料來源')]:
        if old in p.text:
            for t in p._p.xpath('.//w:t'):
                t.text = (t.text or '').replace(old, new)
for t in body.xpath('.//w:t'):
    if t.text:
        t.text = re.sub(r'表 5-(\d+)', lambda m: '表 5-'+str(int(m[1])-1) if int(m[1])>=5 else m[0], t.text)

# Latest instruction: table captions above, explanation below; figures below.
moved = 0
for p in list(doc.paragraphs):
    if re.match(r'^表\s*\d+-\d+', p.text) and '\t' not in p.text:
        previous = p._p.getprevious()
        if previous is not None and previous.tag == qn('w:tbl'):
            previous.addprevious(p._p)
            moved += 1
        p.paragraph_format.keep_with_next = True

def clean_text(text):
    for a, b in [('工作留', '工作流'), ('實驗數據（如果適用）', '實驗數據'),
                 ('我主要看', '我們主要評估'), ('所以我會', '因此我們會'),
                 ('我選 ConvNeXt', '我們選擇 ConvNeXt'), ('所以我在', '因此我們在')]:
        text = text.replace(a, b)
    return text

# Format all XML text including table runs and hyperlink runs, retaining URLs.
parts = [doc.part]
for rel in doc.part.rels.values():
    if not rel.is_external and any(k in rel.reltype for k in ('header', 'footer')):
        parts.append(rel.target_part)
for part in parts:
    root = part.element
    for t in root.xpath('.//w:t'):
        t.text = clean_text(t.text or '')
    for p in root.xpath('.//w:p'):
        pp = p.find(qn('w:pPr'))
        if pp is None:
            pp = OxmlElement('w:pPr'); p.insert(0, pp)
        spacing = pp.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing'); pp.append(spacing)
        spacing.set(qn('w:line'), '240'); spacing.set(qn('w:lineRule'), 'auto')
        style = pp.find(qn('w:pStyle'))
        style_id = style.get(qn('w:val'), '') if style is not None else ''
        heading = style_id.lower().startswith('heading') or style_id == 'Title'
        for r in p.xpath('.//w:r'):
            rp = r.find(qn('w:rPr'))
            if rp is None:
                rp = OxmlElement('w:rPr'); r.insert(0, rp)
            for name, attrs in [('rFonts', {'ascii':'Times New Roman','hAnsi':'Times New Roman','eastAsia':'標楷體','cs':'Times New Roman'}), ('color',{'val':'000000'}), ('sz',{'val':'28' if heading else '24'}), ('szCs',{'val':'28' if heading else '24'})]:
                el = rp.find(qn('w:'+name))
                if el is None:
                    el = OxmlElement('w:'+name); rp.append(el)
                el.attrib.clear()
                for key, value in attrs.items(): el.set(qn('w:'+key),value)
    for el in list(root.xpath('.//w:shd | .//w:highlight')):
        el.getparent().remove(el)
for style in doc.styles:
    if style.type in (1, 2):
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(14 if style.name.startswith('Heading') or style.name == 'Title' else 12)
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        rf.attrib.clear()
        for key,value in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','標楷體')]: rf.set(qn('w:'+key),value)
    for el in list(style.element.xpath('.//w:shd | .//w:highlight')):
        el.getparent().remove(el)

doc.save(OUT)
check = Document(OUT)
assert len(check.inline_shapes) == len(Document(BASE).inline_shapes)
assert len(check.tables) == len(Document(BASE).tables) - 1
all_text = '\n'.join(p.text for p in check.paragraphs)
for key in ['0.008', 'SHA-256', 'HSL', 'avoidColors', 'https://docs.python.org/3/library/hashlib.html']:
    assert key in all_text, key
print(OUT)
print(f'Tables: {len(check.tables)}; retained images: {len(check.inline_shapes)}; moved captions: {moved}')
