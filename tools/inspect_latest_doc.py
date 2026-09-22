from docx import Document
from pathlib import Path

doc = Document(next(path for path in Path.cwd().glob('*20260914.docx') if path.name.startswith('妝識你的美_最終')))
text = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'inline_images', len(doc.inline_shapes))
for token in ['錯誤! 尚未定義書籤', 'Error! Reference source not found', 'Decorate Me AI']:
    print(token, text.count(token))
print('figure_314_captions', sum('圖 3-14（' in paragraph.text for paragraph in doc.paragraphs))
