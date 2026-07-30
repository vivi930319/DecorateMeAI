from __future__ import annotations

import ast
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BACKEND = Path(r"C:\Users\user\OneDrive - 淡江大學\Desktop\PythonProject12")
FRONTEND = Path(r"C:\Users\user\OneDrive - 淡江大學\Desktop\web_frontend")
OUT = BACKEND / "docs" / "DecorateMeAI_全端系統技術文件書_2026-07-31"
TODAY = date(2026, 7, 31)

NAVY = "17365D"
BLUE = "2E74B5"
PALE = "E8EEF5"
LIGHT = "F2F4F7"
MUTED = "666666"
WHITE = "FFFFFF"
RED = "9B1C1C"
GREEN = "1F5E3B"


@dataclass
class Route:
    service: str
    method: str
    path: str
    handler: str
    line: int
    source: str


def route_inventory() -> list[Route]:
    result: list[Route] = []
    files = [
        "ai_gateway.py", "Face_analyzer_BASIC.py", "Face_analyzer_PRO.py",
        "Ollama_suggestion.py", "replicate_render_api.py", "face_feedback.py",
    ]
    for name in files:
        path = BACKEND / name
        tree = ast.parse(path.read_text(encoding="utf-8-sig"))
        for node in ast.walk(tree):
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            for dec in node.decorator_list:
                if not isinstance(dec, ast.Call) or not isinstance(dec.func, ast.Attribute):
                    continue
                method = dec.func.attr.upper()
                if method not in {"GET", "POST", "PUT", "PATCH", "DELETE"} or not dec.args:
                    continue
                try:
                    route_path = ast.literal_eval(dec.args[0])
                except Exception:
                    continue
                if isinstance(route_path, str):
                    result.append(Route(name, method, route_path, node.name, node.lineno, name))
    return sorted(result, key=lambda r: (r.service, r.path, r.method))


def js_api_inventory() -> list[tuple[str, int]]:
    text = (FRONTEND / "js" / "api.js").read_text(encoding="utf-8-sig")
    found = []
    for i, line in enumerate(text.splitlines(), 1):
        m = re.match(r"\s{4}([A-Za-z_$][\w$]*)\s*\([^)]*\)\s*\{", line)
        if m:
            found.append((m.group(1), i))
    return found


def page_inventory() -> list[tuple[str, str]]:
    labels = {
        "dashboard": "首頁與任務入口", "analysis": "BASIC／PRO 臉部分析",
        "style": "妝容風格選擇", "ai-render": "AI 妝容渲染",
        "compare": "妝前／妝後比較", "suggestion": "個人化妝容建議",
        "products": "商品推薦與商品詳情", "favorites": "收藏商品",
        "history": "分析與渲染歷史", "profile": "會員中心與點數",
        "admin": "管理中台",
    }
    return [(p.stem, labels.get(p.stem, p.stem)) for p in sorted((FRONTEND / "pages").glob("*.html"))]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[int]):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def font_run(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Microsoft JhengHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure(doc: Document, running_title: str):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, NAVY, 10, 5),
    ):
        st = styles[name]
        st.font.name = "Microsoft JhengHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font_run(header.add_run(f"DecorateMe AI｜{running_title}"), 8.5, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(footer.add_run("內部技術文件｜版本 1.0｜2026-07-31"), 8, color=MUTED)


def cover(doc: Document, volume: str, title: str, subtitle: str):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("DECORATEME AI"), 11, True, BLUE)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run(title), 25, True, NAVY)
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run(subtitle), 12, color=MUTED)
    p.paragraph_format.space_after = Pt(36)
    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    entries = [("文件編號", volume), ("文件版本", "1.0"), ("基準日期", "2026-07-31"), ("文件狀態", "正式統籌／交接基準")]
    for row, (a, b) in zip(meta.rows, entries):
        row.cells[0].text, row.cells[1].text = a, b
        set_cell_shading(row.cells[0], PALE)
        for c in row.cells:
            for r in c.paragraphs[0].runs:
                font_run(r, 10, bold=(c is row.cells[0]))
    set_table_widths(meta, [2200, 7160])
    doc.add_page_break()


def add_p(doc, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        font_run(p.add_run(bold_lead), 10.5, True)
        font_run(p.add_run(text[len(bold_lead):]), 10.5)
    else:
        font_run(p.add_run(text), 10.5)
    return p


def bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        font_run(run, 10.5)
    if not p.runs:
        font_run(p.add_run(text), 10.5)
    else:
        p.runs[0].text = text


def table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        set_cell_shading(t.rows[0].cells[i], PALE)
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    font_run(run, 9, bold=(ri == 0), color=NAVY if ri == 0 else "000000")
    set_table_widths(t, widths)
    return t


def common_front_matter(doc: Document, scope: str, audience: str):
    doc.add_heading("文件目的與適用範圍", 1)
    add_p(doc, scope)
    doc.add_heading("目標讀者", 2)
    add_p(doc, audience)
    doc.add_heading("閱讀與維護規則", 2)
    for x in (
        "本文件以原始碼與部署設定為準；若文件與執行結果不一致，先凍結上線並完成差異確認。",
        "任何端點、資料欄位、權限、環境變數或上游位址變更，都必須同步更新對應分冊、測試與版本紀錄。",
        "機密值只記錄變數名稱、保存位置與輪替程序，不得把實際金鑰、Cookie、Token 或個資寫入文件。",
        "每次正式部署後應更新基準日期、版本、已知限制、驗收結果與回復方式。",
    ):
        bullet(doc, x)
    doc.add_heading("責任分級", 2)
    table(doc, ["層級", "定義", "交付責任"], [
        ["主責", "直接設計、實作與維護", "程式、測試、部署、文件與事故處理"],
        ["整合主責", "定義契約並完成端到端驗證", "接口規格、相容性、驗收與回歸測試"],
        ["協作端", "由其他組別維護的上游或下游", "按契約提供服務並回覆異常"],
    ], [1300, 3300, 4760])


def save(doc: Document, filename: str):
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / filename
    doc.save(path)
    return path


def build_master():
    d = Document(); configure(d, "全端系統技術文件總冊")
    cover(d, "DM-AI-MASTER-001", "全端系統技術文件總冊", "前端、後端、AI 模型、資料、資安、部署與跨組接口統籌基準")
    common_front_matter(d, "本總冊定義 DecorateMe AI 的整體系統邊界、技術治理、分冊關係與端到端責任。它是專案統籌、交接、審查、Demo 與正式部署的第一入口。", "專案負責人、技術主管、前後端工程師、模型工程師、維運與協作組別。")
    d.add_heading("系統總覽", 1)
    add_p(d, "DecorateMe AI 是以瀏覽器 SPA 為使用者入口、AI Gateway 為信任邊界，串接會員／商品資料庫、BASIC／PRO 臉部分析、文字建議與妝容渲染服務的多服務系統。核心交換物件為 analysisPackage；登入身分以 Gateway 建立的 HttpOnly Session 為準。")
    table(d, ["層", "主要元件", "主要責任"], [
        ["體驗層", "web_frontend", "登入、分析旅程、建議、商品、收藏、歷史、會員與管理中台"],
        ["邊界層", "ai_gateway.py", "Session、CSRF、代理、權限、媒體授權、錯誤正規化與稽核"],
        ["AI 服務層", "BASIC、PRO、Ollama、Render", "臉部推論、多角度分析、文字生成與妝容渲染"],
        ["資料層", "Firestore、GCS、會員／商品資料庫", "工作狀態、媒體、會員與商品持久化"],
        ["工程治理", "Docker、Cloud Run、Cloud Build、CI", "建置、部署、驗證、回復與環境隔離"],
    ], [1300, 3000, 5060])
    d.add_heading("分冊地圖", 1)
    rows = [
        ["01", "系統總覽與治理", "全域架構、責任、資料流、治理與非功能需求"],
        ["02", "Web 前端", "頁面、路由、狀態、使用者旅程、API 對應與前端資安"],
        ["03", "AI Gateway", "認證、代理、權限、媒體與管理端點"],
        ["04", "臉部分析服務", "BASIC／PRO、同步／非同步、姿態與回饋"],
        ["05", "模型與資料科學", "資料集、ROI、CNN、DINOv2、ONNX、評估與發布"],
        ["06", "文字建議服務", "Ollama、提示資料、串流、逾時與降級"],
        ["07", "妝容渲染服務", "Replicate、Job、媒體生命週期與保留"],
        ["08", "資料儲存、資安與隱私", "Firestore、GCS、圖片安全、稽核與刪除"],
        ["09", "部署、測試與維運", "容器、雲端、CI、監控、事故與回復"],
        ["10", "跨組接口與驗收", "資料庫、商品、爬蟲、推薦、Ollama 的 RACI 與驗收"],
    ]
    table(d, ["冊別", "名稱", "用途"], rows, [900, 2800, 5660])
    d.add_heading("端到端主要流程", 1)
    for title, desc in (
        ("身分流程", "使用者登入 → Gateway 驗證會員上游 → 建立 HttpOnly Session → 前端以 credentials=include 呼叫受保護資源 → 登出或到期清除狀態。"),
        ("分析流程", "上傳妝前照 → 圖片安全處理 → BASIC／PRO 推論 → 建立 analysisPackage → 產生文字建議與商品推薦 → 可選擇渲染 → 保存或刪除。"),
        ("管理流程", "管理員 Session → Gateway 角色驗證 → 商品／會員／爬蟲暫存代理 → 稽核紀錄 → 成功或錯誤回前端。"),
        ("刪除流程", "使用者或管理員提出刪除 → 驗證 owner／權限 → 清除工作、GCS 媒體與資料庫紀錄 → 回報可驗證結果。"),
    ):
        d.add_heading(title, 2); add_p(d, desc)
    d.add_heading("系統級非功能要求", 1)
    for x in (
        "安全：瀏覽器不得持有上游 API 金鑰；正式環境採 Session-only、CSRF 防護、同源代理與最小權限。",
        "隱私：妝前照與渲染圖預設非公開；暫存物件必須有 TTL，保留動作須明確且可撤銷。",
        "可靠性：所有外部呼叫必須具逾時、錯誤分類與可觀測狀態；非同步工作必須可查詢且可回收。",
        "相容性：analysisPackage、分類表、端點路徑與錯誤碼採版本化契約；入口需正規化舊分類值。",
        "可交接性：每個服務必須具啟動方式、環境變數、健康檢查、測試、部署與回復程序。",
    ): bullet(d, x)
    d.add_heading("統籌者責任", 1)
    add_p(d, "本系統的統籌責任橫跨 Web 前端、Gateway、AI 服務、模型、資料／資安與部署。統籌者對端到端結果負最終責任，但協作端的內部實作仍由各組別擁有；應透過接口契約、驗收證據與 RACI 管理責任，避免將所有異常無條件歸入單一人員。")
    return save(d, "00_DecorateMeAI_全端系統技術文件總冊.docx")


def build_architecture():
    d=Document(); configure(d,"系統總覽、架構與技術治理"); cover(d,"DM-AI-ARCH-001","系統總覽、架構與技術治理規格書","系統邊界、信任區、資料流、品質屬性與變更治理")
    common_front_matter(d,"本冊說明系統組成、服務邊界、信任模型、共用契約及技術治理，是進行設計審查與重大變更評估的基準。","架構設計者、全端工程師、資安審查者、維運與專案統籌。")
    d.add_heading("架構原則",1)
    for x in ("Gateway 是瀏覽器到後端服務的唯一正式信任邊界。","前端只依賴同源公開契約，不直接保存或傳送內部服務金鑰。","分析、文字與渲染可獨立部署、健康檢查、擴縮與回復。","媒體與工作狀態分離：二進位物件進 GCS，生命週期與權限狀態進 Job Store。","跨端資料以版本化 analysisPackage 傳遞，禁止各頁面自行拼接互不相容的臨時欄位。"):
        bullet(d,x)
    d.add_heading("服務與連線關係",1)
    table(d,["來源","目的地","協定／資料","控制"],[
        ["Browser SPA","AI Gateway","HTTPS、JSON、multipart、Cookie","同源、CSRF、Session、逾時"],
        ["Gateway","Member DB","HTTPS JSON","API key、owner pinning、錯誤正規化"],
        ["Gateway","Product DB／Crawler","HTTPS JSON","白名單、管理員角色、稽核"],
        ["Gateway","Face BASIC／PRO","HTTPS multipart／JSON","上游金鑰、大小限制、job token"],
        ["Gateway","Text／Render","HTTPS JSON／stream","上游金鑰、逾時、媒體授權"],
        ["Services","Firestore／GCS","SDK／雲端 API","服務帳號、TTL、IAM"],
    ],[1800,2200,2500,2860])
    d.add_heading("信任區與威脅模型",1)
    for h,t in (("公開區","瀏覽器輸入、URL、檔案與第三方回應皆視為不可信。"),("邊界區","Gateway 驗證身分、來源、CSRF、角色、路徑與內容，再決定是否代理。"),("服務區","內部服務仍須驗證上游金鑰與輸入，不能只依賴網路位置。"),("資料區","Firestore、GCS 與 Secret Manager 採最小 IAM，媒體預設 private。")):
        d.add_heading(h,2); add_p(d,t)
    d.add_heading("共用資料契約：analysisPackage",1)
    table(d,["區段","內容","治理要求"],[
        ["identity","package id、owner、建立時間、版本","owner 不得由前端任意覆寫"],
        ["input","影像來源與處理後 metadata","不得長期保存原始 EXIF／不必要個資"],
        ["analysis","BASIC／PRO 結果、信心與品質","分類值須通過 canonical mapping"],
        ["generativeText","文字建議、模型狀態與錯誤","不得假裝 fallback 是真實模型輸出"],
        ["recommendations","商品與排序資訊","保留來源、版本與降級狀態"],
        ["render","job、妝前／妝後媒體引用、狀態","媒體只能使用受授權引用"],
        ["async","進度、錯誤、重試與時間","狀態轉移需單調且可追蹤"],
    ],[1700,4100,3560])
    d.add_heading("變更治理",1)
    for x in ("提出 RFC：說明原因、影響服務、契約差異、資料遷移、資安與回復。","建立相容層：新增欄位優先可選，移除／改名須經棄用期。","完成合約測試：前端、Gateway、上游與下游各自提供證據。","分階段部署：先非生產、再小流量，觀察錯誤率與延遲後擴大。","更新所有受影響分冊、設定範例、測試清單與故障排除。"):
        bullet(d,x)
    d.add_heading("架構驗收清單",1)
    for x in ("所有正式瀏覽器請求均經 Gateway。","每個服務有 /health，且健康不等於依賴皆健康時能清楚揭露。","工作與媒體皆具 owner、TTL、刪除與稽核路徑。","所有外部依賴具明確逾時與 fail-closed／降級策略。","環境變數和 Secret 不出現在 Git、前端 bundle 或錯誤訊息中。"):
        bullet(d,"□ "+x)
    return save(d,"01_系統總覽_架構與技術治理規格書.docx")


def build_frontend():
    d=Document(); configure(d,"Web 前端系統技術規格"); cover(d,"DM-AI-FE-001","Web 前端系統技術規格書","SPA 架構、頁面、狀態、API、使用者旅程、資安與部署")
    common_front_matter(d,"本冊完整描述 web_frontend 的瀏覽器端責任，包括頁面、路由、狀態、認證、分析旅程、管理中台、錯誤處理及 Firebase Hosting 部署。","前端與全端工程師、UI／UX、測試、資安與維運人員。")
    d.add_heading("技術基線與目錄",1)
    add_p(d,"前端採原生 HTML、CSS、ES6 JavaScript 的單頁應用架構，沒有 package.json 或框架建置鏈。index.html 提供應用殼層，pages/*.html 提供頁面片段，router.js 控制路由與畫面生命週期，api.js 封裝 Gateway／上游契約，data.js 管理靜態資料與本機資料模型。")
    table(d,["檔案／目錄","責任"],[
        ["index.html","應用殼、頂部導覽、登入／註冊容器與腳本載入"],
        ["pages/","頁面 HTML 片段，由 Router 載入並初始化"],
        ["js/router.js","hash 路由、頁面初始化、流程狀態、事件與管理中台"],
        ["js/api.js","Session、CSRF、受保護 fetch、資料正規化與各服務呼叫"],
        ["js/data.js","風格／商品 fallback、analysisPackage、本機歷史與草稿"],
        ["js/service-endpoints.js","服務路徑集中定義"],
        ["js/config.js／config.local.js","公開設定與本機覆寫；不得放秘密"],
        ["css/main.css","全站視覺、響應式、對話框與頁面元件"],
        ["firebase.json／deploy.ps1","Hosting rewrites、headers 與發布"],
        ["frontend_smoke_check.js","靜態契約、XSS、Session 與關鍵流程回歸檢查"],
    ],[3000,6360])
    d.add_heading("頁面與路由清冊",1)
    table(d,["Hash／頁面","功能","存取層級"],[[f"#{p}",label,"admin" if p=="admin" else "登入會員"] for p,label in page_inventory()],[2200,5160,2000])
    d.add_heading("核心使用者旅程",1)
    for h,steps in (
        ("註冊與 OTP",["輸入會員資料並完成前端格式檢核","向 Gateway 要求寄送 OTP","驗證 OTP；禁止任何前端 bypass","完成註冊並建立 Session","載入會員資料與導向首頁"]),
        ("臉部分析到推薦",["選擇 BASIC／PRO 並上傳合格影像","建立分析草稿與 analysisPackage","呼叫分析服務並保存結果／job 狀態","選擇妝容風格","取得文字建議與商品推薦","可進行 AI 渲染、前後比較與收藏"]),
        ("Session 生命週期",["開站先呼叫 /auth/session","credentials=include，禁止讀取 HttpOnly Cookie","遇 401 清除私有狀態並回登入","遇 owner mismatch／409 要求使用者採納新 Session 或重新登入","跨分頁以 BroadcastChannel 同步登入狀態"]),
    ):
        d.add_heading(h,2)
        for i,s in enumerate(steps,1): bullet(d,f"{i}. {s}")
    d.add_heading("前端狀態模型",1)
    table(d,["狀態","用途","持久化／清除"],[
        ["Router.currentPage","目前路由與初始化目標","記憶體；路由變更更新"],
        ["Router.analysisPackage","完整分析旅程資料","草稿保存；登出／owner 變更需清除"],
        ["analysisResult／analyzeMode","當次分析結果與 BASIC／PRO 模式","旅程內使用"],
        ["selectedStyleId","妝容風格","旅程內使用；切換需觸發離開保護"],
        ["pendingLook","待保存的妝容紀錄","保存成功或取消後清除"],
        ["generalProductCatalog","商品快取","錯誤或管理更新後失效"],
        ["Session owner","目前登入身分釘選","不得由 URL／表單任意取代"],
    ],[2500,3900,2960])
    d.add_heading("API 封裝清冊",1)
    methods=js_api_inventory()
    table(d,["方法","來源行","說明"],[[m,str(line),"api.js 公開或內部呼叫封裝；變更前需檢查呼叫者與 smoke test"] for m,line in methods[:80]],[2700,1100,5560])
    d.add_heading("前端資安控制",1)
    for x in ("所有受保護請求使用 credentials=include，瀏覽器端不保存上游 API key。","狀態改變請求附帶 Gateway 規定的 CSRF token／header。","使用者文字進入 innerHTML 前必須 escapeHtml；外部 URL 必須限制 http/https 並套 rel=noopener noreferrer。","影像 URL 使用專用安全函式驗證，不可直接插入 avatar、商品或媒體來源。","私有頁顯示前必須 validateSession；登出、到期、409 owner mismatch 時清除敏感狀態。","config.local.js 只能放本機公開位址，不得提交秘密；正式環境以同源 rewrite 為準。"):
        bullet(d,x)
    d.add_heading("錯誤、可用性與降級",1)
    table(d,["情境","使用者處理","工程處理"],[
        ["401／Session 到期","回登入並顯示可理解訊息","清除 owner-bound cache，禁止靜默重試"],
        ["403／CSRF 或角色不足","提示重新整理或權限不足","記錄路徑與 correlation id，不顯示秘密"],
        ["409／owner mismatch","要求採納新帳號或重新登入","停止使用舊會員的草稿與快取"],
        ["分析／文字／渲染逾時","保留草稿並允許重試","job 查詢需有上限與退避"],
        ["商品推薦失敗","顯示全部商品或明確維護狀態","不得把 fallback 偽裝成個人化推薦"],
    ],[2100,3500,3760])
    d.add_heading("前端測試與驗收",1)
    for x in ("11 個頁面片段可載入且 hash 路由、上一頁與重新整理正常。","登入、OTP、登出、Session 到期與跨分頁帳號切換實測通過。","BASIC／PRO 分析、風格、建議、商品、渲染、比較、收藏及歷史可完成一輪。","惡意 HTML、javascript: URL、錯誤圖片 URL 與跨 owner 快取無法造成資料外洩。","手機、平板、桌機主要斷點無水平溢出；鍵盤操作與焦點可辨識。","frontend_smoke_check.js 與瀏覽器 smoke test 通過後才可 deploy.ps1 發布。"):
        bullet(d,"□ "+x)
    return save(d,"02_Web前端系統_完整技術規格書.docx")


def route_doc(title, code, filename, services, intro, sections):
    d=Document(); configure(d,title); cover(d,code,title,"服務邊界、端點、資料流、資安、部署與驗收")
    common_front_matter(d,intro,"全端／後端工程師、服務維護者、測試、資安與維運人員。")
    for h,paras in sections:
        d.add_heading(h,1)
        for p in paras: add_p(d,p)
    relevant=[r for r in route_inventory() if r.service in services]
    d.add_heading("端點清冊",1)
    table(d,["方法","路徑","處理器／來源"],[[r.method,r.path,f"{r.handler}（{r.source}:{r.line}）"] for r in relevant],[1100,4700,3560])
    d.add_heading("每一端點的共同契約要求",1)
    for x in ("明確定義認證方式、角色、owner、CSRF、Content-Type、大小與欄位限制。","成功回應需固定狀態碼與 schema；錯誤採可機器判讀 code、可閱讀 message 與 correlation id。","所有上游呼叫具 connect／read timeout；不可無限重試或把逾時誤報為 404。","記錄方法、正規化路徑、狀態碼、延遲與 correlation id；不得記錄 Cookie、Token、API key 或完整影像。","端點變更需同步前端封裝、合約測試、OpenAPI／本冊及部署驗收。"):
        bullet(d,x)
    d.add_heading("標準驗收矩陣",1)
    table(d,["類別","必要案例","通過條件"],[
        ["正常","最小合法輸入、完整合法輸入","schema、狀態碼與副作用正確"],
        ["身分","未登入、過期、錯誤 owner、角色不足","401／403／409 分類一致且無資料洩漏"],
        ["輸入","缺欄位、錯型別、超限、惡意內容","4xx 且服務不崩潰"],
        ["依賴","上游 4xx／5xx／timeout／斷線","錯誤正規化、可觀測、可重試性正確"],
        ["併發","重複送出、輪詢、刪除競態","冪等或明確衝突，不產生孤兒資料"],
        ["隱私","跨使用者讀取、URL 猜測、刪除後讀取","拒絕存取且刪除可驗證"],
    ],[1400,4000,3960])
    return save(d,filename)


def build_gateway():
    return route_doc("AI Gateway 與認證代理技術規格書","DM-AI-GW-001","03_AI_Gateway與認證代理_完整技術規格書.docx",["ai_gateway.py"],"本冊定義 Gateway 作為瀏覽器與所有後端服務間的唯一正式入口，包括 Session、CSRF、路徑代理、管理權限、媒體授權、稽核與錯誤治理。",[
        ("角色與邊界",["Gateway 不負責重新實作會員、商品、臉部、文字或渲染業務；它負責驗證、限制、轉送、正規化與稽核。","正式環境以前端同源路徑搭配 Firebase Hosting rewrite 進入 Gateway，避免瀏覽器直連各服務。"]),
        ("認證與 Session",["登入時將會員上游結果轉為 Gateway Session；Cookie 採 HttpOnly、Secure、SameSite=Lax，TTL 由環境設定控制。","受保護請求以 Session actor 為唯一身分來源；X-Expected-Actor 用於偵測前端快取帳號與目前 Session 不一致，而不是讓前端指定 actor。"]),
        ("CSRF 與管理權限",["狀態改變請求採 double-submit CSRF。管理路徑除 Session 外必須驗證管理角色，敏感動作寫入 audit。","公開健康檢查與 public-config 必須只曝露公開資訊，不得回傳內部 URL、Secret 名稱以外的值或環境內容。"]),
        ("代理政策",["只允許明確白名單服務與路徑；重寫路徑前正規化並阻擋 path traversal。外部文字上游未允許時 fail closed。","上游錯誤轉為一致 error envelope；保留必要狀態語意，禁止把所有錯誤包成 200。"]),
        ("媒體存取",["/media/render 路徑必須確認 Session、owner 或工作 token，不能依公開 GCS URL 作授權。legacy token 只能作相容入口並具有期限。"]),
    ])


def build_face():
    return route_doc("BASIC／PRO 臉部分析服務技術規格書","DM-AI-FACE-001","04_BASIC_PRO臉部分析服務_完整技術規格書.docx",["Face_analyzer_BASIC.py","Face_analyzer_PRO.py","face_feedback.py"],"本冊說明 BASIC 正面分析、PRO 多角度分析、姿態判斷、同步／非同步工作、結果查詢與使用者回饋的完整契約。",[
        ("BASIC 服務",["BASIC 以單張正面照片為主要輸入，完成圖片安全、臉部偵測、ROI 裁切、模型推論、規則／融合校正與標準結果輸出。","同步 /analyze 與 /v1/face/analyze/basic 適合短任務；/v1/face/jobs/basic 建立可查詢工作，供 Gateway 與前端輪詢。"]),
        ("PRO 服務",["PRO 處理多角度或側臉資訊，驗證 yaw／pitch 與影像品質，再輸出側臉鼻型等進階特徵。輸入不足時應明確標示品質或缺角度，不得臆測。"]),
        ("推論管線",["接收檔案 → sanitize_upload → 解碼 → MediaPipe／InsightFace 幾何 → ROI → CNN／DINOv2／ONNX → 類別正規化 → 信心與品質 → analysisPackage。","所有模型類別順序以 *_classes.json 為準；模型二進位、類別檔與程式 mapping 必須原子發布。"]),
        ("工作狀態",["建議狀態：queued、running、succeeded、failed、expired。結果只在 succeeded 可讀；failed 提供安全錯誤碼，expired 不再提供媒體。","工作需綁 owner 或不可猜測 job token，具 TTL、併發限制與清理機制。"]),
        ("回饋端點",["回饋端點以 job_id 找到原始預測，驗證 Session／X-Job-Token 與 owner 後記錄使用者選擇。回饋資料不可直接在線修改模型，須經標註複核與版本化重訓。"]),
        ("品質拒絕條件",["無臉、多臉、解析度不足、檔案超限、姿態不符、遮擋嚴重、解碼失敗或安全檢查不通過時，回傳明確 4xx／品質狀態。"]),
    ])


def build_model():
    d=Document(); configure(d,"AI 模型、資料集與訓練治理"); cover(d,"DM-AI-ML-001","AI 模型、資料集與訓練治理規格書","標註、ROI、CNN、DINOv2、規則融合、評估、發布與監控")
    common_front_matter(d,"本冊規範臉型與五官分類模型從資料收集、標註、切分、訓練、評估、模型匯出到線上推論的完整生命週期。","模型工程師、資料標註人員、後端工程師、研究審查與品質負責人。")
    d.add_heading("任務與產物",1)
    table(d,["特徵","主要方法","正式產物"],[[x,"CNN／DINOv2／幾何或融合（依現行選模）",f"{x}.onnx、{x}_classes.json；必要時 DINOv2 head.npz"] for x in ["face_shape","eye_shape","brow_shape","nose_shape","lip_shape"]],[1900,4300,3160])
    d.add_heading("資料治理",1)
    for x in ("保留資料來源、授權／同意、身分群組、標註版本與排除原因。","train／validation／test 必須按人物身分分組，禁止同一人跨切分造成洩漏。","合併類別後重新建立 identity split、類別檔與驗證集；不得只改顯示名稱。","標註疑義進 review queue，由人工決策後以不可變紀錄套用。","原始臉照採最小保存、權限隔離與到期刪除；訓練輸出不得含可回推個資。"):
        bullet(d,x)
    d.add_heading("訓練管線",1)
    for i,x in enumerate(("檢查標註表與類別集合","建立身分映射與資料切分","產生 ROI／輪廓快取","訓練 CNN、DINOv2 head 或規則樹","交叉驗證、學習曲線與錯標診斷","依 macro F1、每類 precision／recall 與混淆矩陣選模","匯出 ONNX／NPZ 與 classes.json","shadow 驗證後切換正式答案"),1):
        bullet(d,f"{i}. {x}")
    d.add_heading("評估標準",1)
    table(d,["指標","目的","最低治理要求"],[
        ["Macro F1","避免大類掩蓋小類","必報平均與各類結果"],
        ["Precision／Recall","評估搶答與漏判","每類皆需揭露"],
        ["混淆矩陣","定位類別重疊","使用固定類別順序"],
        ["身分群組 CV","估計陌生人泛化","不得以影像隨機切分取代"],
        ["信心校準","決定低信心拒答／fallback","以獨立驗證資料調整"],
        ["延遲／大小","確保線上可部署","CPU／記憶體與冷啟動一併量測"],
    ],[1900,3300,4160])
    d.add_heading("模型發布契約",1)
    for x in ("版本包含資料版、程式 commit、超參數、隨機種子、指標與產物雜湊。","classes.json 的索引順序必須與模型輸出完全一致；發布前執行一致性檢查。","入口先正規化舊類別值，下游只接 canonical values。","新模型先 shadow，對照正式模型的差異、低信心與錯誤案例，通過後才切換。","回復需能一次恢復模型二進位、classes.json、規則與程式，不可只換單檔。"):
        bullet(d,x)
    d.add_heading("公平性與風險",1)
    add_p(d,"臉部特徵分類可能受膚色、光線、年齡、性別呈現、相機與妝容影響。評估應依可合法使用且樣本足夠的群組做切片，揭露樣本數與不確定性；結果定位為造型輔助，不作醫療、身分、人格或價值判斷。")
    d.add_heading("驗收清單",1)
    for x in ("無資料洩漏，切分可重現。","五項特徵的類別表、模型輸出與 API mapping 一致。","報告含各類指標、混淆矩陣、失敗案例與推論效能。","線上服務缺模型或載入失敗時 fail closed／明確降級。","模型卡、版本紀錄、回復包與監控門檻齊全。"):
        bullet(d,"□ "+x)
    return save(d,"05_AI模型_資料集與訓練治理_完整規格書.docx")


def build_text():
    return route_doc("Ollama 文字建議服務技術規格書","DM-AI-TEXT-001","06_Ollama文字建議服務_完整技術規格書.docx",["Ollama_suggestion.py"],"本冊定義分析資料轉為個人化妝容文字建議的服務，包括 prompt 資料最小化、同步／串流、逾時、安全輸出與降級。",[
        ("服務責任",["服務接收已結構化且最小化的臉部分析與風格資料，不接收原始臉照；將資料組成受控 prompt 並呼叫 Ollama。","服務不得代替商品推薦、會員授權或臉部分類；輸出是造型建議而非醫療／人格判定。"]),
        ("輸入治理",["僅允許 schema 內欄位、長度與枚舉值；忽略或拒絕可疑指令注入文字。Prompt 中區分 system 規則與 user data，不把使用者資料當指令。"]),
        ("同步與串流",["/suggest 適合完整 JSON 回應；/suggest/stream 應使用明確事件格式傳送增量、完成與錯誤。客戶端中斷時取消上游工作。"]),
        ("可靠性與降級",["設定連線、首 token、總時間與輸出長度上限。上游不可達時回明確服務狀態；若使用模板 fallback，回應必須標示來源，不能偽裝為模型生成。","Gateway 的外部文字上游開關預設 fail closed，正式開啟需完成資料流與供應商風險審查。"]),
        ("輸出安全",["禁止歧視、羞辱、醫療診斷、敏感屬性推斷與保證性宣稱。前端顯示時仍需 HTML escape；日誌不得保存完整個人 prompt。"]),
    ])


def build_render():
    return route_doc("Replicate 妝容渲染與媒體服務技術規格書","DM-AI-RENDER-001","07_Replicate妝容渲染與媒體服務_完整技術規格書.docx",["replicate_render_api.py"],"本冊定義妝前影像經 Replicate 產生妝後影像的同步／非同步接口，以及簽名、讀取、保留、刪除與使用者清除的媒體生命週期。",[
        ("渲染責任",["接收已授權妝前影像引用與受控風格／提示，建立第三方預測，追蹤狀態並將結果移入私有 GCS。","服務不以第三方暫時 URL 作長期保存，也不讓前端直接持有 Replicate token。"]),
        ("Job 狀態機",["queued → running → succeeded／failed；另有 cancelled／expired。狀態寫入 Job Store，重複查詢不得重複建立預測。","建立工作回 202 與 job id；查詢回進度、可重試錯誤與安全媒體引用。"]),
        ("媒體生命週期",["temporary 物件具短 TTL；使用者明確保存後移入 retained/{opaqueOwnerId}/{jobId}。妝前與妝後圖皆需 owner 授權。","signed-url 僅短效；content 代理適用不應曝露桶位址的情境。刪除工作、artifact、media 與 user 必須處理關聯物件並具冪等性。"]),
        ("第三方與隱私",["送往 Replicate 的資料應限於完成渲染所需；文件記錄供應商、區域、保存政策與刪除能力。日誌不保存圖片內容或完整 URL query token。"]),
        ("失敗與補償",["第三方失敗、下載失敗、GCS 寫入失敗或資料庫寫入失敗均需分類。若先產生媒體後寫狀態失敗，清理程序需回收孤兒物件。"]),
    ])


def build_security():
    d=Document(); configure(d,"資料儲存、資安與隱私"); cover(d,"DM-AI-SEC-001","資料儲存、資安與隱私規格書","Firestore、GCS、圖片安全、身分授權、稽核、保存與刪除")
    common_front_matter(d,"本冊統一規範系統中的資料分類、保存位置、授權、影像處理、日誌、稽核、刪除及事故應變。","全體工程師、資安、維運、資料治理與專案統籌。")
    d.add_heading("資料分類",1)
    table(d,["資料","敏感度","保存與控制"],[
        ["Session／CSRF／API key","機密","Cookie／Secret Manager；禁止日誌與前端儲存"],
        ["妝前臉照","高度敏感個資","私有 GCS、短 TTL、owner 授權、明確保存"],
        ["妝後渲染圖","敏感個資","私有 GCS、短效 URL、可刪除"],
        ["臉部分析結果","敏感衍生資料","owner 綁定、最小欄位、保存期限"],
        ["會員資料","個資","會員資料庫；Gateway 只代理必要欄位"],
        ["商品資料","內部／公開混合","商品資料庫；管理動作需角色與稽核"],
        ["操作日誌","內部敏感","遮蔽秘密與個資，限制存取與期限"],
    ],[2400,1800,5160])
    d.add_heading("圖片安全管線",1)
    for i,x in enumerate(("限制 Content-Length 與實際讀取大小","驗證 magic bytes、MIME 與允許格式","安全解碼並限制像素／尺寸，防解壓炸彈","移除 EXIF、XMP 與非必要 metadata","重新編碼為受控格式","才交給分析、渲染或儲存"),1):
        bullet(d,f"{i}. {x}")
    d.add_heading("授權模型",1)
    for x in ("Session actor 是使用者身分唯一依據；URL、body email 或 header 只能用於一致性檢查。","工作、媒體、收藏與歷史均綁 owner；管理員角色不能自動繞過所有資料最小化要求。","Job token 必須高熵、短效、用途受限，且不能出現在長期日誌。","GCS Bucket 不公開；signed URL 短效，權限撤銷與物件刪除需可驗證。"):
        bullet(d,x)
    d.add_heading("保存與刪除",1)
    table(d,["生命週期","控制","驗收"],[
        ["暫存","TTL／lifecycle、自動清理","到期後 URL、代理與物件皆不可讀"],
        ["保留","使用者明確動作、owner 路徑","會員可查看、撤銷與刪除"],
        ["工作狀態","Firestore TTL、索引與最小 payload","過期後不洩露結果"],
        ["帳號刪除","會員、歷史、收藏、工作、媒體協調刪除","回報各子系統結果並可重試"],
        ["備份／日誌","固定期限與權限","刪除政策包含延遲與例外說明"],
    ],[1700,3900,3760])
    d.add_heading("稽核與事故",1)
    for x in ("管理操作記錄 actor、動作、目標引用、結果、時間與 correlation id；不記錄秘密。","偵測大量登入失敗、跨 owner 拒絕、異常媒體讀取、管理端批次操作與錯誤率突升。","事故流程：隔離 → 保存證據 → 撤銷憑證 → 封鎖路徑 → 評估資料範圍 → 修復驗證 → 通知與事後檢討。","Secret 輪替後同步更新 Cloud Run revision，確認舊 revision 不再承接流量。"):
        bullet(d,x)
    d.add_heading("資安驗收",1)
    for x in ("前端 bundle、Git 歷史、錯誤頁與 public-config 無秘密。","跨使用者 job／媒體／歷史／收藏讀寫皆被拒絕。","CSRF、XSS、SSRF、路徑穿越、惡意影像與暴力登入測試通過。","刪除與 TTL 實測可驗證；GCS 無公開存取。","日誌遮蔽與稽核查詢經抽查。"):
        bullet(d,"□ "+x)
    return save(d,"08_資料儲存_資安與隱私_完整規格書.docx")


def build_ops():
    d=Document(); configure(d,"部署、測試、監控與維運"); cover(d,"DM-AI-OPS-001","部署、測試、監控與維運手冊","本機、容器、Cloud Run、CI、可觀測性、事故與回復")
    common_front_matter(d,"本冊提供從本機啟動、測試、容器建置、雲端部署到正式驗收、監控及事故回復的標準作業程序。","開發、測試、維運、資安與發布負責人。")
    d.add_heading("可部署單元",1)
    table(d,["單元","本機預設埠","建置／啟動"],[
        ["AI Gateway","8015","Dockerfile.gateway／python ai_gateway.py"],
        ["Face BASIC","8001","Dockerfile／python Face_analyzer_BASIC.py"],
        ["Face PRO","8002","Dockerfile／python Face_analyzer_PRO.py"],
        ["Text Suggestion","8010","python Ollama_suggestion.py"],
        ["Render","8020","Dockerfile.render／uvicorn replicate_render_api:app"],
        ["Web Frontend","Hosting","Firebase Hosting／deploy.ps1"],
    ],[2600,1800,4960])
    d.add_heading("環境設定治理",1)
    add_p(d,"設定分為公開設定、非機密運行設定與 Secret。公開設定可由 /public-config 或前端設定提供；API key、Session secret、第三方 token 與服務帳號只進 Secret Manager／受控環境。所有必要變數需在各環境有值、格式檢查與安全預設。")
    d.add_heading("測試分層",1)
    table(d,["層級","主要工具／檔案","阻擋條件"],[
        ["靜態／語法","ruff、py_compile、前端 smoke","語法、危險模式或契約檢查失敗"],
        ["單元","ai_gateway_test、image_safety_test","安全與核心邏輯失敗"],
        ["服務 smoke","backend_smoke_test、render_api_test、ollama smoke","健康、端點或錯誤契約失敗"],
        ["模型驗證","CV、混淆矩陣、class consistency","指標退化、類別不一致、產物缺失"],
        ["端到端","瀏覽器＋正式 rewrite","登入到分析／推薦／渲染流程失敗"],
        ["安全","權限、CSRF、XSS、媒體 owner","可越權、資料洩漏或 fail-open"],
    ],[1700,4000,3660])
    d.add_heading("標準發布流程",1)
    for i,x in enumerate(("凍結版本與確認 Git 工作區範圍","執行靜態、單元、服務、模型與前端測試","建立不可變 image tag／digest，記錄 commit","部署至 staging，執行健康與端到端驗收","部署新 Cloud Run revision／Firebase Hosting","小流量觀察錯誤率、p95、冷啟動與依賴狀態","擴至 100%，保存證據並更新文件","若超過門檻，將流量切回上一 revision／Hosting release"),1):
        bullet(d,f"{i}. {x}")
    d.add_heading("監控與告警",1)
    table(d,["訊號","建議觀測","告警目的"],[
        ["可用性","各服務 2xx、健康與依賴狀態","服務或上游中斷"],
        ["延遲","p50／p95／p99、冷啟動、模型推論","容量與使用體驗退化"],
        ["錯誤","按 route／code／upstream 分類","定位回歸與依賴問題"],
        ["非同步","queued/running 時間、失敗、孤兒 job","渲染／分析工作阻塞"],
        ["資安","401/403/409、rate limit、跨 owner","攻擊或帳號狀態異常"],
        ["成本","Cloud Run、GCS、Firestore、Replicate","預算與濫用"],
    ],[1600,4200,3560])
    d.add_heading("事故 Runbook",1)
    for h,t in (("上游 5xx／timeout","確認單一路徑或全面異常；啟動降級、停止重試風暴、通知協作端。"),("登入全面失敗","檢查 Gateway revision、Session secret、會員上游與 Cookie／rewrite；禁止放寬驗證作臨時修復。"),("媒體不可讀","檢查 owner、job 狀態、GCS IAM、signed URL 時間與物件是否過期。"),("模型結果異常","確認模型、classes.json、mapping 與 revision 是否同版；必要時整包回復。"),("成本突升","檢查渲染併發、輪詢、重複 job、GCS 孤兒與濫用來源，先限流再調查。")):
        d.add_heading(h,2); add_p(d,t)
    d.add_heading("正式驗收證據",1)
    for x in ("commit、image digest、Cloud Run revision、Hosting release。","所有測試輸出與端到端截圖／請求結果。","環境變數名稱與 Secret 版本（不含值）。","監控儀表板、告警與觀察窗口結果。","已知限制、回復點與負責人。"):
        bullet(d,"□ "+x)
    return save(d,"09_部署_測試_監控與維運_完整手冊.docx")


def build_integration():
    d=Document(); configure(d,"跨組接口、RACI 與驗收"); cover(d,"DM-AI-INT-001","跨組接口、RACI 與驗收規格書","前端、Gateway、資料庫、商品、爬蟲、推薦、Ollama 與渲染協作基準")
    common_front_matter(d,"本冊將各端責任、接口交付、故障歸屬、變更通知及驗收證據統一，供專案統籌與各組別對接。","專案統籌、各端負責人、測試與指導／審查人員。")
    d.add_heading("端別責任",1)
    table(d,["端別","主責","對本系統交付"],[
        ["Web 前端","頁面、流程、狀態、顯示與瀏覽器安全","同源 API 呼叫、錯誤體驗、端到端驗收"],
        ["AI Gateway","Session、CSRF、代理、權限、稽核","穩定公開契約與錯誤 envelope"],
        ["臉部／模型","BASIC、PRO、模型與推論品質","版本化結果 schema、信心與品質狀態"],
        ["會員資料庫端","會員、OTP、Session 上游、點數、收藏、歷史","欄位、狀態碼、冪等與刪除契約"],
        ["商品資料庫端","商品 CRUD、查詢與推薦資料","穩定商品 schema、圖片與分類"],
        ["爬蟲端","來源抓取、清洗與 staging","可稽核來源、安全 URL、匯入狀態"],
        ["演算法／推薦端","推薦排序與風格／色系匹配","輸入需求、排序理由與版本"],
        ["Ollama 端","模型運行與文字生成","可用性、模型資訊、串流與錯誤契約"],
        ["Replicate／渲染端","第三方預測與影像產出","job、狀態、媒體與刪除能力"],
        ["維運／雲端","IAM、Secret、Cloud Run、GCS、監控","環境、發布、告警與回復"],
    ],[1900,3400,4060])
    d.add_heading("RACI 原則",1)
    add_p(d,"R=實作者，A=最終負責，C=需諮詢，I=需知會。每項變更只設一個 A；統籌者對整體交付為 A，但協作端仍需對自身服務可用性與資料正確性負 R／A。")
    table(d,["工作","前端／Gateway","協作端","統籌"],[
        ["接口定義","R","C","A"], ["協作端內部實作","C","R/A","I"],
        ["整合測試","R","R","A"], ["正式發布","R","R/C","A"],
        ["事故定位","R","R","A"], ["文件與驗收證據","R","R","A"],
    ],[3100,2100,2100,2060])
    d.add_heading("接口變更單必填欄位",1)
    for x in ("變更編號、提出端、負責人、日期與目標版本。","現況、問題、預期行為、非目標與影響使用者。","方法、路徑、認證、request／response、錯誤碼與範例。","資料欄位型別、必填、枚舉、預設、個資與保存期限。","相容性、棄用期、資料遷移、部署順序與回復。","驗收環境、案例、證據、完成定義與簽核。"):
        bullet(d,x)
    d.add_heading("跨端驗收流程",1)
    for i,x in enumerate(("以契約範例建立 mock／fixture","協作端完成單元與合約測試","Gateway 在 staging 對真上游測試","前端完成正常、錯誤、逾時與權限情境","執行完整使用者旅程與刪除流程","保存 request id、版本、結果與簽核","變更文件狀態為已驗收並排定正式發布"),1):
        bullet(d,f"{i}. {x}")
    d.add_heading("故障歸屬判定",1)
    table(d,["現象","優先負責端","判定證據"],[
        ["瀏覽器未送正確路徑／欄位","前端","Network request 與前端版本"],
        ["Session／CSRF／rewrite 被拒絕","Gateway／部署","Gateway code、Cookie 屬性、Hosting 設定"],
        ["Gateway 正確轉送但上游錯誤","協作上游","correlation id、上游狀態與 payload 摘要"],
        ["上游成功但 Gateway 改壞 schema","Gateway","原始與正規化回應比較"],
        ["模型產物／類別不一致","模型端","產物雜湊、classes.json 與推論輸出"],
        ["媒體存在但越權／讀不到","Render／Gateway／GCS","owner、job、IAM 與物件狀態"],
    ],[2800,2200,4360])
    d.add_heading("統籌儀表板建議",1)
    for x in ("各端版本、環境、健康、最後驗收日期與 owner。","P0／P1 問題、阻擋端、下一動作與期限。","契約變更、棄用項目與部署順序。","端到端成功率、錯誤率、延遲、模型品質與成本。","文件版本、測試證據、事故與回復演練狀態。"):
        bullet(d,x)
    return save(d,"10_跨組接口_RACI與驗收_完整規格書.docx")


def main():
    outputs = [build_master(), build_architecture(), build_frontend(), build_gateway(), build_face(), build_model(), build_text(), build_render(), build_security(), build_ops(), build_integration()]
    manifest = OUT / "文件清單.txt"
    manifest.write_text("DecorateMe AI 全端系統技術文件書\n版本：1.0\n基準日：2026-07-31\n\n" + "\n".join(p.name for p in outputs), encoding="utf-8")
    print("\n".join(str(p) for p in outputs))


if __name__ == "__main__":
    main()
