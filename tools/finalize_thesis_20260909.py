from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / '專題系統文件書_妝識你的美_商品推薦細節補齊_待版面驗證_2026-09-08.docx'
OUTPUT = ROOT / '專題系統文件書_妝識你的美_正文引用與格式修訂版_2026-09-09.docx'
doc = Document(SOURCE)
body = doc.element.body

def drop(e):
    e.getparent().remove(e)

def find(text):
    return next(p for p in doc.paragraphs if p.text.startswith(text) and not p.text.endswith('\t—'))

def add_before(anchor, text, style='Normal'):
    return anchor.insert_paragraph_before(text, style)

# Remove the old static TOCs; rebuilt below from actual headings and captions.
first = find('目錄')._p
last = find('摘要')._p
elements = list(body)
for e in elements[elements.index(first):elements.index(last)]: drop(e)

# Publication sources belong after chapter six, not in a source audit table.
first = find('十、出版與原始資料來源')._p
last = find('陸、')._p
elements = list(body)
for e in elements[elements.index(first):elements.index(last)]: drop(e)

# The final five reference-page screenshots are archived in the original file.
# Replace this non-editable bibliography, not the system diagrams/screenshots.
last = find('系統穩定性部分，未來可持續')._p
elements = list(body)
for e in elements[elements.index(last)+1:]:
    if e.tag != qn('w:sectPr'): drop(e)
for p in list(doc.paragraphs):
    if p.text.strip() in ['- 簡短描述專題的主要內容及結果', '（組別、老師與成員請依實際資料填寫）']:
        drop(p._p)

# Restore precise academic voice while keeping uncertainty and measured values.
rewrites = {
 '本系統最大的特色在於': '本系統結合五官分析、個人化妝容建議、虛擬試妝及商品推薦，並整合一般電商平台的商品瀏覽與收藏功能。使用者可從了解五官特徵、選擇妝容、預覽試妝結果，接續查看對應商品。',
 '本研究主要探討如何運用': '本研究探討人工智慧於線上美妝試妝及導購流程的應用。首先，透過臉部關鍵點與影像分類分析五官及膚色特徵；其次，經由 Replicate API 串接圖片生成模型，依原始照片與妝容提示產生試妝影像；最後，以 CIEDE2000 色差及商品條件比對推薦候選。本研究以建立完整流程為目標；圖片身分保留及推薦準確度仍須依實際評估結果判定。',
 '同時透過 AI 技術協助使用者': '本系統將臉部分析結果與使用者選擇的妝容條件用於商品比對，串接分析、試妝與推薦流程，以協助使用者判斷商品。使用者留存率及購物體驗改善幅度尚未量測，不作為已驗證成果。',
 '這張表的重點是把資料來源': '各推薦入口依資料來源與個人化程度區分。商品清單、商品詳情與臉部分析結果頁採用的資料與排序依據不同，介面應保留來源標示，避免將一般商品排序誤認為個人化推薦。',
 '目前要把兩個 Port 分開講清楚': 'Ollama 原生推論服務使用 Port 11434，Suggestion Service 使用 Port 8010。開發測試期間可透過 Cloudflare Quick Tunnel 將本機 127.0.0.1:8010 暫時提供給雲端 Gateway。Quick Tunnel 僅作連線通道，並不代表 Ollama 部署於 Cloudflare，也不等同正式 production 架構。',
}
for p in doc.paragraphs:
    for prefix,text in rewrites.items():
        if p.text.startswith(prefix): p.text = text

# Add author-date citations to research claims, not as proof of our deployment.
citations = {
 'MediaPipe FaceMesh 主要負責': 'MediaPipe 提供組合式感知流程框架（Lugaresi et al., 2019）；臉部關鍵點功能可參考 Google（n.d.）的官方說明。',
 '臉部分析模型的開發過程曾比較': 'ConvNeXt 的架構依據為 Liu et al.（2022）；ONNX 模型交換格式參考 ONNX（n.d.）。',
 '本系統於個人化妝容建議模組中採用': '結構化輸出與推論 API 參考 Ollama（n.d.-a, n.d.-b），LLaVA 的方法背景參考 Liu et al.（2023）。',
 '文字妝容建議與圖片生成在本系統中': '服務串接方式參考 Replicate（n.d.）；該文件不作為本系統妝容影像品質的驗證證據。',
 '本系統的商品推薦主要採用': '內容式推薦的理論背景參考 Lops et al.（2011）；CIEDE2000 的實作與測試依據參考 Sharma et al.（2005）及 International Commission on Illumination（2022）。',
 '商品資料擷取使用 Requests': 'HTTP 存取與 HTML 解析分別參考 Python Software Foundation（n.d.-c）與 Richardson（n.d.）。',
 '由於本系統同時包含臉部分析': '物件存取及生命週期的公開技術背景參考 Google Cloud（n.d.-a, n.d.-b），前端託管背景參考 Firebase（n.d.）。',
 '商品十六進位色碼先轉換為 HSL': '色彩空間轉換的 API 說明參考 Python Software Foundation（n.d.-a）。',
 '為避免不同使用者長期取得完全相同商品': 'SHA-256 的函式介面參考 Python Software Foundation（n.d.-b）；0.008 為本專案演算法端記載的設定，不是函式庫預設值。',
 '建議 Ollama 輸出下列 JSON': 'JSON 型別約束參考 JSON Schema（n.d.），結構化生成介面參考 Ollama（n.d.-b）。',
 '四季型屬於本系統的影像分類': '影像色差對量測及觀看條件的要求可參考 International Commission on Illumination（2011）。',
}
for p in doc.paragraphs:
    for prefix,text in citations.items():
        if p.text.startswith(prefix): p.add_run(' '+text)

# Explicitly scope supplementary implementation documents; no invented tests.
anchor = find('（六）商品推薦演算法與排序流程')
anchor = doc.paragraphs[doc.paragraphs.index(anchor)+1] if False else anchor
add_before(find('商品推薦端新增的重點'), '本節整合演算法端的商品推薦技術紀錄（妝識你的美專題團隊，2026a），涵蓋推薦入口、資料欄位、計分規則與測試範圍。紀錄中的測試通過僅支持所列情境，不代表完整 E2E 已驗收。')
add_before(find('Ollama 在本系統中不是'), '本節依 Ollama 模組技術文件整理服務分工及輸入輸出（妝識你的美專題團隊，2026b）。')
add_before(find('模型訓練和模型換版不是'), 'Worker 流程與限制依專案技術紀錄整理（妝識你的美專題團隊，2026c），未完成及未驗證狀態均保留於本節。')

# Fix captions interrupted by previously inserted prose.
for prefix,header in [('表 3-15','資料區段'),('表 3-16','商品類別')]:
    caption=find(prefix)
    table=next(t for t in doc.tables if t.cell(0,0).text.strip()==header)
    table._tbl.addprevious(caption._p)
for p in list(doc.paragraphs):
    if re.match(r'^表\s*\d+-\d+',p.text):
        node = p._p.getnext()
        while node is not None and node.tag != qn('w:tbl'):
            if node.tag == qn('w:p') and re.match(r'^[圖表]\s*\d+-\d+', ''.join(node.xpath('.//w:t/text()'))): break
            node = node.getnext()
        if node is not None and node.tag == qn('w:tbl'):
            node.addprevious(p._p)
        p.paragraph_format.keep_with_next = True

# Crop only the embedded top title area using editable Word crop metadata.
for prefix,crop in [('圖 3-16',6500),('圖 3-17',6000)]:
    p = find(prefix)
    previous=p._p.getprevious()
    if previous is not None:
        # Older builder painted over part of the start node. Re-embed the
        # intact user-supplied image and crop only its original title band.
        asset=ROOT/'docs/figures'/('algorithm_supplement_1.png' if prefix=='圖 3-16' else 'algorithm_supplement_2.png')
        rid,_=doc.part.get_or_add_image(str(asset))
        for blip in previous.xpath('.//a:blip'): blip.set(qn('r:embed'),rid)
        for fill in previous.xpath('.//pic:blipFill'):
            rect=fill.find(qn('a:srcRect'))
            if rect is None:
                rect=OxmlElement('a:srcRect'); fill.insert(1,rect)
            rect.set('t',str(crop))
        for ext in previous.xpath('.//wp:extent | .//a:xfrm/a:ext'):
            ext.set('cy',str(round(int(ext.get('cy'))*(1-crop/100000))))

# Keep the activity figure and caption together within the existing page.
activity=find('圖 3-9')._p.getprevious()
for ext in activity.xpath('.//wp:extent | .//a:xfrm/a:ext'):
    for dim in ['cx','cy']:
        ext.set(dim,str(round(int(ext.get(dim))*0.92)))
from docx.text.paragraph import Paragraph
Paragraph(activity,doc).paragraph_format.keep_with_next=True

# Long operational descriptions should be body text, not 14-point headings.
for p in list(doc.paragraphs):
    if p.style.name.startswith('Heading') and len(p.text)>80:
        short=re.split(r'[:：]',p.text,1)[0]
        if len(short)<42: add_before(p,short,'Heading 3')
        p.style='Normal'
        for r in p.runs: r.bold=False
        if ':' in p.text or '：' in p.text:
            p.text=re.split(r'[:：]',p.text,1)[1]

for p in list(doc.paragraphs):
    if p.text.strip()=='。': drop(p._p)
    if p.text.startswith('六、模型分析與使用者回饋資料流程'):
        p.text=p.text.replace('六、','十一、',1)
    if p.text in ['HSL 妝容風格字典','四季型與冷暖底調','NLP 與 Ollama 文字建議','可重現的近分輪替','逐商品計分與類別重排','建議的 NLP 結構化輸出範例','商品推薦的限制與後續改善']:
        p.style='Heading 3'
        p.paragraph_format.keep_with_next=True
    if p.text.startswith('使用者如果未註冊帳號'):
        p.text='使用者可由登入頁的「立即註冊」進入註冊流程，填寫資料並完成電子郵件驗證。忘記密碼時，使用者可申請驗證碼，驗證通過後設定新密碼。會員個資與密碼的儲存方式應依各欄位分別說明，不將全部會員資料概括為雜湊資料。管理員可透過後台查看授權範圍內的會員權限與妝容紀錄。'
    if p.text.startswith('最後一個頁面為模型修正複核介面'):
        p.text='模型修正複核頁呈現使用者回饋及已同意提供的影像，由管理員進行人工複核並決定是否採用。採用後建立送訓工作，由地端 Worker 執行；電腦休眠、重新啟動或工作中斷後的行為，仍須依 Worker 狀態與測試紀錄判斷，不將介面操作視為已完成训练。'.replace('训练','訓練')
    if p.text.startswith('本系統目前已完成從前端操作'):
        p.text=p.text.replace('本系統目前已完成從前端操作、臉部分析、妝容文字建議、圖片渲染、商品推薦、會員收藏到後台複核的完整流程。','本系統已建立前端操作、臉部分析、妝容文字建議、圖片渲染、商品推薦、會員收藏與後台複核的主要模組；各項完成狀態仍依對應測試與驗收紀錄判定。')

# References are unnumbered, editable and alphabetized. Italic spans use *...*.
refs = [
 ('Firebase. (n.d.). *Firebase Hosting*. https://firebase.google.com/docs/hosting'),
 ('Google. (n.d.). *Face landmark detection guide*. Google AI Edge. https://ai.google.dev/edge/mediapipe/solutions/vision/face_landmarker'),
 ('Google Cloud. (n.d.-a). *Object Lifecycle Management*. https://cloud.google.com/storage/docs/lifecycle'),
 ('Google Cloud. (n.d.-b). *Signed URLs*. https://cloud.google.com/storage/docs/access-control/signed-urls'),
 ('International Commission on Illumination. (2011). *Methods for evaluating colour differences in images* (CIE 199:2011). https://www.cie.co.at/publications/methods-evaluating-colour-differences-images'),
 ('International Commission on Illumination. (2022). *Colorimetry—Part 6: CIEDE2000 colour-difference formula* (ISO/CIE 11664-6:2022). https://www.cie.co.at/publications/colorimetry-part-6-ciede2000-colour-difference-formula-1'),
 ('JSON Schema. (n.d.). *Type-specific keywords*. https://json-schema.org/understanding-json-schema/reference/type'),
 ('Liu, H., Li, C., Wu, Q., & Lee, Y. J. (2023). *Visual instruction tuning*. arXiv. https://doi.org/10.48550/arXiv.2304.08485'),
 ('Liu, Z., Mao, H., Wu, C.-Y., Feichtenhofer, C., Darrell, T., & Xie, S. (2022). *A ConvNet for the 2020s*. arXiv. https://doi.org/10.48550/arXiv.2201.03545'),
 ('Lops, P., de Gemmis, M., & Semeraro, G. (2011). Content-based recommender systems: State of the art and trends. In F. Ricci, L. Rokach, B. Shapira, & P. B. Kantor (Eds.), *Recommender systems handbook* (pp. 73–105). Springer. https://doi.org/10.1007/978-0-387-85820-3_3'),
 ('Lugaresi, C., Tang, J., Nash, H., McClanahan, C., Uboweja, E., Hays, M., Zhang, F., Chang, C.-L., Yong, M. G., Lee, J., Chang, W.-T., Hua, W., Georg, M., & Grundmann, M. (2019). *MediaPipe: A framework for building perception pipelines*. arXiv. https://doi.org/10.48550/arXiv.1906.08172'),
 ('Ollama. (n.d.-a). *Generate a response*. https://docs.ollama.com/api/generate'),
 ('Ollama. (n.d.-b). *Structured outputs*. https://docs.ollama.com/capabilities/structured-outputs'),
 ('ONNX. (n.d.). *Introduction to ONNX*. https://onnx.ai/onnx/intro/'),
 ('Python Software Foundation. (n.d.-a). *colorsys—Conversions between color systems*. Python documentation. https://docs.python.org/3/library/colorsys.html'),
 ('Python Software Foundation. (n.d.-b). *hashlib—Secure hashes and message digests*. Python documentation. https://docs.python.org/3/library/hashlib.html'),
 ('Python Software Foundation. (n.d.-c). *Requests: HTTP for humans*. https://requests.readthedocs.io/en/latest/'),
 ('Ramírez, S. (n.d.). *Request body*. FastAPI. https://fastapi.tiangolo.com/tutorial/body/'),
 ('Replicate. (n.d.). *Documentation*. https://replicate.com/docs'),
 ('Richardson, L. (n.d.). *Beautiful Soup documentation*. https://www.crummy.com/software/BeautifulSoup/bs4/doc/'),
 ('Sharma, G., Wu, W., & Dalal, E. N. (2005). The CIEDE2000 color-difference formula: Implementation notes, supplementary test data, and mathematical observations. *Color Research & Application, 30*(1), 21–30. https://doi.org/10.1002/col.20070'),
 ('The PostgreSQL Global Development Group. (n.d.). *Array functions and operators*. PostgreSQL documentation. https://www.postgresql.org/docs/current/functions-array.html'),
 ('妝識你的美專題團隊（2026a）。*妝識你的美商品推薦系統與演算法技術文件*［未出版技術文件］。'),
 ('妝識你的美專題團隊（2026b）。*妝識你的美個人化妝容建議模組詳細技術說明*［未出版技術文件］。'),
 ('妝識你的美專題團隊（2026c）。*Worker 技術細節*［未出版技術文件］。'),
]
heading=doc.add_paragraph('參考文獻','Heading 1')
heading.paragraph_format.page_break_before=True
for reference in refs:
    p=doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent=Inches(.5)
    p.paragraph_format.first_line_indent=Inches(-.5)
    p.paragraph_format.space_after=Pt(6)
    for i,chunk in enumerate(reference.split('*')):
        r=p.add_run(chunk); r.italic=bool(i%2)

# Cite supplemental data/API sources at their actual technical discussion.
add_before(find('商品推薦 API、前端呈現與結果證據'), '推薦端的請求資料建模及資料庫陣列操作，其公開介面說明分別參考 Ramírez（n.d.）與 The PostgreSQL Global Development Group（n.d.）；本系統實際契約仍以表列端點及回傳欄位為準。')

# Preserve original local locations and deployed URLs, separated from bibliography.
p=doc.add_paragraph('原始資料連結','Heading 2')
locations=[
 ('商品推薦技術文件', 'C:/Users/isach/Downloads/Decorate_Me_智慧彩妝推薦系統_商品推薦技術文件.docx'),
 ('Ollama 技術文件', 'C:/Users/isach/Downloads/DECORATE_ME_Ollama_個人化妝容建議模組_詳細技術說明.md'),
 ('Worker 技術細節', 'C:/Users/isach/PycharmProjects/PythonProject12/docs/專案管理與交接/Worker技術細節.md'),
 ('歷史流程更改追蹤', 'C:/Users/isach/PycharmProjects/PythonProject12/docs/專案管理與交接/歷史流程更改追蹤.md'),
 ('串接卡點歷史', 'C:/Users/isach/PycharmProjects/PythonProject12/docs/專案管理與交接/串接卡點歷史.md'),
 ('前端原始封存', 'C:/Users/isach/Downloads/專題前端 (Remix) (Remix) (1).zip'),
 ('前端部署原始入口', 'https://decorate-me.web.app'),
 ('Face Basic 原始健康檢查入口', 'https://face-basic-258021445391.asia-east1.run.app/health'),
 ('Face Pro 原始健康檢查入口', 'https://face-pro-258021445391.asia-east1.run.app/health'),
 ('Render 原始健康檢查入口', 'https://replicate-render-258021445391.asia-east1.run.app/health'),
]
doc.add_paragraph('以下保留技術文件、歷史紀錄及服務的原始查核位置。歷史網址或本機路徑不代表目前服務可用、模型已部署或完整 UAT 已通過。')
for label,url in locations: doc.add_paragraph(label+'：'+url)

# Normalize typo fragments and style properties without touching evidence states.
fixes={'工作留':'工作流','覆写':'覆寫','截近':'捷徑','一般店商':'一般電商','特色道選擇':'特色到選擇','建議即個人化':'建議及個人化','我覺得':'團隊評估','我認為':'團隊認為','我做了什麼':'執行工作','我發現的問題':'發現的問題','我這邊實際處理的內容':'實際處理內容','我把 jobId':'本系統將 jobId'}
for t in body.xpath('.//w:t'):
    for old,new in fixes.items(): t.text=(t.text or '').replace(old,new)

# Build tables of contents with stable bookmarks and genuine PAGEREF fields.
anchor=find('摘要')
headings=[p for p in doc.paragraphs if p.style.name.startswith('Heading')]
captions=[p for p in doc.paragraphs if re.match(r'^[圖表]\s*\d+-\d+',p.text)]
bookmark_id=10000
def bookmark(p):
    global bookmark_id
    bookmark_id+=1
    name=f'thesis_{bookmark_id}'
    s=OxmlElement('w:bookmarkStart'); s.set(qn('w:id'),str(bookmark_id)); s.set(qn('w:name'),name)
    e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),str(bookmark_id))
    p._p.insert(1 if p._p.pPr is not None else 0,s); p._p.append(e)
    return name
for title,items in [('目錄',headings),('圖目錄',[p for p in captions if p.text.startswith('圖')]),('表目錄',[p for p in captions if p.text.startswith('表')])]:
    h=add_before(anchor,title)
    h.paragraph_format.page_break_before=True
    h.paragraph_format.keep_with_next=True
    h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for target in items:
        p=add_before(anchor,re.sub(r'\s+',' ',target.text).strip()+'\t')
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.left_indent=Pt(12 if target.style.name=='Heading 2' else 24 if target.style.name=='Heading 3' else 0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0),2,1)
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGEREF '+bookmark(target)+' \\h')
        r=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text='0'; r.append(text); field.append(r); p._p.append(field)
anchor.paragraph_format.page_break_before=True

# Full font enforcement, including styles, fields, tables, headers and footers.
roots=[doc.element, doc.styles.element]
for rel in doc.part.rels.values():
    if not rel.is_external and any(x in rel.reltype for x in ('header','footer')): roots.append(rel.target_part.element)
for root in roots:
    for blip in root.xpath('.//a:blip'):
        if blip.find(qn('a:grayscl')) is None: blip.append(OxmlElement('a:grayscl'))
    for el in list(root.xpath('.//w:shd | .//w:highlight')): drop(el)
    for color in root.xpath('.//w:color'):
        color.attrib.clear(); color.set(qn('w:val'),'000000')
    for border in root.xpath('.//w:tblBorders/* | .//w:tcBorders/*'):
        border.set(qn('w:color'),'000000')
        for attr in ['themeColor','themeTint','themeShade']:
            border.attrib.pop(qn('w:'+attr),None)
    for p in root.xpath('.//w:p'):
        pp=p.find(qn('w:pPr'))
        if pp is None: pp=OxmlElement('w:pPr'); p.insert(0,pp)
        ps=pp.find(qn('w:pStyle'))
        sid=ps.get(qn('w:val'),'') if ps is not None else ''
        txt=''.join(p.xpath('.//w:t/text()'))
        is_heading=sid.lower().startswith('heading') or txt in ['專題系統文件書','妝識你的美','目錄','圖目錄','表目錄']
        spacing=pp.find(qn('w:spacing'))
        if spacing is None: spacing=OxmlElement('w:spacing'); pp.append(spacing)
        spacing.set(qn('w:line'),'240'); spacing.set(qn('w:lineRule'),'auto')
        for r in p.xpath('.//w:r'):
            rp=r.find(qn('w:rPr'))
            if rp is None: rp=OxmlElement('w:rPr'); r.insert(0,rp)
            specs=[('rFonts',{'ascii':'Times New Roman','hAnsi':'Times New Roman','cs':'Times New Roman','eastAsia':'標楷體'}),('color',{'val':'000000'}),('sz',{'val':'28' if is_heading else '24'}),('szCs',{'val':'28' if is_heading else '24'})]
            for tag,attributes in specs:
                el=rp.find(qn('w:'+tag))
                if el is None: el=OxmlElement('w:'+tag); rp.append(el)
                el.attrib.clear()
                for key,value in attributes.items(): el.set(qn('w:'+key),value)
for table in doc.tables:
    for row in table.rows:
        for h in list(row._tr.xpath('./w:trPr/w:trHeight')): drop(h)
        pr=row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None: pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next=False
                p.paragraph_format.widow_control=True
    if table.rows:
        pr=table.rows[0]._tr.get_or_add_trPr()
        if pr.find(qn('w:tblHeader')) is None: pr.append(OxmlElement('w:tblHeader'))
doc.settings.element.append(OxmlElement('w:updateFields'))
doc.settings.element[-1].set(qn('w:val'),'true')
doc.save(OUTPUT)
print(OUTPUT)
print('References',len(refs),'Tables',len(doc.tables),'Images',len(doc.inline_shapes))
