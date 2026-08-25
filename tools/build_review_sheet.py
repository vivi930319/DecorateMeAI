"""用**現在**這一版模型跑一批照片，產生人工校對表。

為什麼要重做一份
----------------
校對表裡的「模型_XX」欄是**產生表格當下**的預測，模型換過之後它就過期了。
2026-08-26 查到既有的 `亞洲人臉10000筆_ConvNeXt_待人工校對.xlsx` 就是這樣：
它的眉型只有三類（沒有「挑眉」），而線上早就換成四類的 20260824_brow4。
拿那份表去談準確率，談的是一個已經不存在的模型。

所以這支腳本做兩件事：
1. 讀 `models/basic_features_roi/` 底下**現在**的 ONNX 重跑一次推論；
2. **把模型身分寫進表格裡**——每個部位的類別清單、ONNX 的 sha256、產生時間，
   都存進第二個工作表。下次有人看到這份表，可以直接判斷它有沒有過期，
   而不必像這次一樣靠「眉型少了一類」才發現。

輸出的欄位跟既有校對表一致（檔名／看圖／模型_／信心_／次選_／人工_／狀態_），
所以 `tools/import_manual_review.py` 與 `tools/score_against_manual.py` 都吃得下。

分布摘要要看
------------
跑完會印每個部位的預測分布。這個數字比準確率更早發現問題：如果某一類佔了九成，
那通常不是資料不平衡，是模型在這批照片上塌掉了——它把所有東西推進同一類。

用法
----
    python tools/build_review_sheet.py --images data/kaggle_asian_faces/generated_yellow-stylegan2 \
        --out 亞洲人臉一萬筆人工校對/亞洲人臉10000筆_現行模型_待人工校對.xlsx
    python tools/build_review_sheet.py --images <dir> --out <xlsx> --limit 200   # 先試小批
"""
from __future__ import annotations

import _bootstrap  # noqa: F401
import argparse
import hashlib
import json
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import cv2
import mediapipe_ascii  # noqa: F401
import mediapipe as mp
import numpy as np

import basic_roi_shadow
from basic_roi_shadow import PART_TO_FIELD

MAX_IMAGE_SIZE = 1024
PART_ORDER = ["face_shape", "brow_shape", "eye_shape", "nose_shape", "lip_shape"]


def model_fingerprint() -> dict:
    """這份表是哪一版模型跑的。少了它，表格過期就只能靠人眼發現。"""
    model_dir = Path(basic_roi_shadow.MODEL_DIR)
    info = {"modelDir": str(model_dir), "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "parts": {}}
    for part in PART_ORDER:
        onnx = model_dir / f"{part}.onnx"
        classes = model_dir / f"{part}_classes.json"
        entry: dict = {}
        if classes.is_file():
            entry["classes"] = json.loads(classes.read_text(encoding="utf-8")).get("classes", [])
        if onnx.is_file():
            digest = hashlib.sha256()
            with onnx.open("rb") as handle:
                for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                    digest.update(chunk)
            entry["sha256"] = digest.hexdigest()
            entry["bytes"] = onnx.stat().st_size
        if entry:
            info["parts"][part] = entry
    return info


def predict_top2(frame, points, sessions) -> dict:
    """每個部位回 (最可能, 信心, 次可能)。

    次選對校對的人很有用：模型不確定的時候，正確答案通常就是次選那一個，
    校對的人可以直接挑，不必自己想。
    """
    out = {}
    for part, (session, classes) in sessions.items():
        try:
            crop = basic_roi_shadow.crop_roi(frame, points, part)
            logits = session.run(None, {"input": basic_roi_shadow.roi_to_tensor(crop)})[0][0]
            probs = basic_roi_shadow._softmax(logits)
            order = np.argsort(probs)[::-1]
            out[part] = (classes[int(order[0])], round(float(probs[order[0]]), 3),
                         classes[int(order[1])] if len(order) > 1 else "")
        except Exception:
            # 單一部位失敗不影響其他部位——一張照片有四個部位判得出來仍然值得校對。
            continue
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--images", required=True, help="照片目錄")
    ap.add_argument("--out", required=True, help="輸出的 xlsx")
    ap.add_argument("--limit", type=int, default=0, help="只跑前 N 張（0 = 全部）")
    ap.add_argument("--ext", default=".png,.jpg,.jpeg", help="要收的副檔名")
    args = ap.parse_args()

    import openpyxl

    img_dir = Path(args.images)
    if not img_dir.is_dir():
        raise SystemExit(f"找不到照片目錄：{img_dir}")
    exts = {e.strip().lower() for e in args.ext.split(",") if e.strip()}
    # 按檔名的數字排序，讓 0.png 排在 10.png 前面——校對的人是照順序看的。
    files = sorted((p for p in img_dir.iterdir() if p.suffix.lower() in exts),
                   key=lambda p: (len(p.stem), p.stem))
    if args.limit:
        files = files[:args.limit]
    if not files:
        raise SystemExit(f"{img_dir} 裡沒有符合 {sorted(exts)} 的檔案")

    fingerprint = model_fingerprint()
    print(f"照片 {len(files)} 張｜模型目錄 {fingerprint['modelDir']}")
    for part, entry in fingerprint["parts"].items():
        print(f"  {part:12} {entry.get('classes')}  sha256={entry.get('sha256', '')[:12]}")

    sessions = basic_roi_shadow._load()
    if not sessions:
        raise SystemExit("載不到 ONNX 模型，確認 models/basic_features_roi/ 內容")

    fields = [PART_TO_FIELD[p] for p in PART_ORDER if p in sessions]
    header = ["檔名", "看圖"]
    for field in fields:
        header += [f"模型_{field}", f"信心_{field}", f"次選_{field}", f"人工_{field}"]
    header += [f"狀態_{field}" for field in fields] + ["資料來源"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "校對"
    ws.append(header)
    ws.freeze_panes = "C2"

    mesh = mp.solutions.face_mesh.FaceMesh(static_image_mode=True, max_num_faces=1,
                                           refine_landmarks=False, min_detection_confidence=0.5)
    dist = {field: Counter() for field in fields}
    no_face = unreadable = 0
    started = time.time()

    for i, path in enumerate(files, 1):
        # 用 imdecode 而不是 imread：路徑含中文時 imread 在 Windows 會回 None，
        # 而那個 None 跟「檔案損毀」長得一模一樣，很難查。
        frame = cv2.imdecode(np.fromfile(str(path), dtype=np.uint8), cv2.IMREAD_COLOR)
        if frame is None:
            unreadable += 1
            continue
        h0, w0 = frame.shape[:2]
        scale = MAX_IMAGE_SIZE / max(h0, w0)
        if scale < 1:
            frame = cv2.resize(frame, (int(w0 * scale), int(h0 * scale)), interpolation=cv2.INTER_AREA)
        h, w = frame.shape[:2]
        res = mesh.process(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
        if not res.multi_face_landmarks:
            no_face += 1
            continue
        pts = np.array([[int(lm.x * w), int(lm.y * h)]
                        for lm in res.multi_face_landmarks[0].landmark], dtype=np.int32)
        pred = predict_top2(frame, pts, sessions)

        row = [path.name, f'=HYPERLINK("{path.resolve().as_posix()}","看圖")']
        for part in PART_ORDER:
            if PART_TO_FIELD.get(part) not in dist:
                continue
            field = PART_TO_FIELD[part]
            label, conf, second = pred.get(part, ("", "", ""))
            row += [label, conf, second, ""]
            if label:
                dist[field][label] += 1
        row += [""] * len(fields) + ["ConvNeXt-Tiny 預測／待人工校對"]
        ws.append(row)

        if i % 250 == 0:
            rate = i / max(time.time() - started, 1e-9)
            left = (len(files) - i) / max(rate, 1e-9)
            print(f"  {i}/{len(files)}　{rate:.1f} 張/秒　剩約 {left / 60:.0f} 分", flush=True)

    mesh.close()

    # 第二個工作表：這份表是誰跑的。這正是既有那份表缺的東西。
    meta = wb.create_sheet("模型版本")
    meta.append(["產生時間", fingerprint["generatedAt"]])
    meta.append(["模型目錄", fingerprint["modelDir"]])
    meta.append(["照片目錄", str(img_dir.resolve())])
    meta.append(["張數", len(files)])
    meta.append([])
    meta.append(["部位", "類別", "ONNX sha256", "位元組"])
    for part, entry in fingerprint["parts"].items():
        meta.append([part, "、".join(entry.get("classes", [])),
                     entry.get("sha256", ""), entry.get("bytes", "")])

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    written = ws.max_row - 1
    print(f"\n完成：{written} 列寫入 {out_path}")
    if no_face or unreadable:
        print(f"  略過 {no_face} 張偵測不到人臉、{unreadable} 張讀不到")

    print("\n各部位的預測分布（某一類佔九成通常代表模型在這批照片上塌掉了，不是資料不平衡）：")
    for field in fields:
        total = sum(dist[field].values()) or 1
        top = dist[field].most_common()
        line = "　".join(f"{label} {n}（{n / total:.0%}）" for label, n in top)
        flag = "  ← 偏斜" if top and top[0][1] / total > 0.7 else ""
        print(f"  {field:6} {line}{flag}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
