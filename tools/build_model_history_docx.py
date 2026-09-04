from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\isach\Downloads\專題系統文件書 (1).docx")
OUTPUT = ROOT / "專題系統文件書_模型歷程詳版.docx"


def set_cell_shading(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=False, color=None, name="Microsoft JhengHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def write_cell(cell, text, bold=False, size=9.2):
    p = clear_cell(cell)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def make_table(doc, created, headers, rows, widths=None, size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        write_cell(table.rows[0].cells[idx], header, bold=True, size=size)
        set_cell_shading(table.rows[0].cells[idx], "C6E0B4")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            write_cell(cells[idx], value, size=size)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    created.append(table._tbl)
    return table


def add_paragraph(doc, created, text="", style=None, bold=False, size=10.5,
                  align=None, keep_with_next=False, code=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, name="Consolas" if code else "Microsoft JhengHei")
    created.append(p._p)
    return p


def add_heading(doc, created, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 11.5,
                 bold=True)
    created.append(p._p)
    return p


def add_bullet(doc, created, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=10.2)
    created.append(p._p)
    return p


def add_code(doc, created, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.7)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=9, name="Consolas")
    created.append(p._p)
    return p


def add_page_break(doc, created):
    p = doc.add_paragraph()
    p.add_run().add_break()
    created.append(p._p)
    return p


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(f"找不到來源文件：{SOURCE}")

    # 保留使用者原始檔，所有新增內容寫入另一份成果檔。
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    created = []

    add_page_break(doc, created)
    add_heading(doc, created, "附錄一、臉部分析模型完整開發歷程與訓練工作留", 1)
    add_paragraph(
        doc, created,
        "這一章是我把臉部分析模型從一開始怎麼想、怎麼試、怎麼失敗，到最後為什麼選現在這個版本，全部整理成一條可以被追溯的歷程。這裡不只放最後的分數，也把我中間遇到的資料問題、切分問題、模型選擇問題、地端訓練流程，以及使用者回饋如何回到訓練資料裡面寫清楚。這樣在展示或口試時，我可以說明的不只是「我們用了什麼模型」，而是「我為什麼這樣做、我怎麼知道它比較好、哪些結果我沒有直接相信」。"
    )
    add_paragraph(
        doc, created,
        "本章的歷史數字來自不同日期、不同資料版本與不同評估協定，所以我不把所有數字硬湊成一條看起來很漂亮的上升曲線。每個表格都會標示它的資料切分、類別版本和用途；如果是探索性實驗，我會直接寫明它不能當成正式準確率。現行可執行模型則以工作區裡的模型檔、類別檔與 manifest 為準。"
    )

    add_heading(doc, created, "一、我先把模型到底要解決什麼定義清楚", 2)
    add_paragraph(
        doc, created,
        "我的臉部分析不是一個模型一次把整張臉回答完，而是拆成五個互相獨立的多分類任務。這樣做是因為臉型、眉型、眼型、鼻型、唇型的可見區域不同，資料量和難度也不同。如果把五個答案塞進同一個分類器，模型很容易依賴背景、妝容或整張照片的風格；拆開後，我可以針對每個部位做 ROI、標籤清理、類別合併和錯誤分析。"
    )
    make_table(
        doc, created,
        ["任務", "現行主要類別", "輸入區域", "我在訓練時關注的問題"],
        [
            ["face_shape", "圓形臉、心形臉、方形臉、長形臉、鵝蛋臉", "正面臉部 ROI", "臉型邊界主觀、髮型與背景容易干擾"],
            ["brow_shape", "一字眉、彎月眉、挑眉、落尾眉", "雙眉 ROI", "早期版本為三類，後來補上挑眉，必須區分版本"],
            ["eye_shape", "下垂眼、圓眼、桃杏眼、鳳眼", "雙眼 ROI", "眼型名稱容易重疊，合併前後不能直接比較"],
            ["nose_shape", "寬鼻、標準鼻", "鼻部 ROI", "窄鼻樣本定義不穩定，後來併入標準鼻"],
            ["lip_shape", "厚唇、微笑唇、花瓣唇、薄唇", "嘴唇 ROI", "M 型唇與花瓣唇的界線不一致"],
        ],
        widths=[2.5, 5.1, 3.2, 7.0],
    )
    add_paragraph(
        doc, created,
        "資料整理時，我從原始照片去除檔名重複後得到 2,311 張照片，對應約 1,015 個身份。五個任務實際可用的標註數量不同，早期統計約為臉型 530、眉型 441、眼型 676、鼻型 269 或 542、唇型 269 或 542；鼻型與唇型的數字會因為資料版本、合併方式與是否包含後續補充資料而不同。因此後面看到鼻型或唇型的結果時，一定要跟著版本與 protocol 看，不能只看小數點後面的數字。"
    )

    add_heading(doc, created, "二、整條模型管線實際怎麼走", 2)
    add_paragraph(
        doc, created,
        "使用者上傳照片後，系統先判斷工作類型，再由 MediaPipe Face Mesh 找出臉部關鍵點。關鍵點不是直接拿來當最後答案，而是用來定位臉、眉、眼、鼻、唇的 ROI，接著把各個 ROI 送進對應的分類器。分類結果會經過格式化與信心資訊整理，再交給建議服務，最後才串到產品推薦或虛擬試妝。也就是說，模型的錯誤會一路影響後面的建議，因此我不能只把模型當成獨立的 notebook 實驗。"
    )
    add_code(
        doc, created,
        "照片 → Face Mesh landmarks → ROI 擷取與正規化 → 五個 ConvNeXt-Tiny 分類器 → analysisPackage → 建議／商品推薦／回饋"
    )
    add_paragraph(
        doc, created,
        "這也是我後來把模型輸出固定成結構化 analysisPackage 的原因。前端不應該自己猜模型回傳的文字，Gateway 也不應該把不同服務的欄位任意拼接；每個部位要有 label、confidence、modelVersion、是否使用 fallback 等資訊，後面才能追蹤某一次結果到底是由哪個模型、哪一版類別定義產生。"
    )

    add_heading(doc, created, "三、模型世代與我為什麼一直換模型", 2)
    make_table(
        doc, created,
        ["世代", "當時的做法", "我遇到的問題", "最後處理"],
        [
            ["第 0 代", "MediaPipe 比例、幾何規則、if-else", "規則可解釋但對亞洲臉、妝容與角度不穩，閾值常常根本碰不到", "保留作 fallback／診斷，不當主要答案"],
            ["第 1 代", "整張臉的 CNN，160×160", "鼻子、眉毛等小部位比例太小；隨機切分還會產生身份洩漏", "淘汰整臉輸入，重新做 ROI 與 identity split"],
            ["第 2 代", "ROI CNN、MobileNetV3-small、ImageNet 預訓練", "模型變輕且方向正確，但部位標籤、類別不平衡和資料切分仍影響很大", "作為基準線，繼續比較更強架構"],
            ["第 3 代", "資料清理、身份分組、五折交叉驗證，加入 DINOv2 比較", "重複圖片、舊 identity map、窄鼻定義造成結果失真", "修資料與 protocol，不能只追模型分數"],
            ["第 4 代", "調 epoch、learning rate、class merge、augmentation", "小資料集的結果有 noise，部分類別本身就難以一致標註", "40 epochs、重新建立身份 map，合併不穩定類別"],
            ["第 5 代", "比較 MobileNet、EfficientNet、ResNet、ConvNeXt、AlexNet", "大模型較準但檔案和推論成本變高；單一部位最高不等於整體最好", "選 ConvNeXt-Tiny 作為 BASIC／PRO 主要架構"],
            ["第 6 代", "ONNX 化、manifest、版本化與可回滾 promotion", "訓練完成不代表線上已換模，部署與模型檔可能不同步", "把訓練、上傳、部署、驗證、回滾接成可追蹤流程"],
        ],
        widths=[2.0, 5.0, 7.2, 5.0],
        size=8.7,
    )

    add_heading(doc, created, "3.1 第 0 代：先用 MediaPipe 與幾何規則建立可解釋基線", 3)
    add_paragraph(
        doc, created,
        "我一開始其實沒有直接訓練深度模型，而是先用 MediaPipe Face Mesh 取得臉部關鍵點，再計算臉寬、臉高、眼睛比例、鼻翼寬度等幾何特徵，最後用 if-else 判斷。這個版本的優點是我可以直接解釋：例如某個比例超過多少就判成某一類，而且不需要 GPU，遇到模型載入失敗時也可以當作備援。"
    )
    add_paragraph(
        doc, created,
        "但實際測試後，我發現規則不是只有閾值要調，而是特徵和標籤本身不一定對得上。例如眉型測試約只有 33% 的使用者同意，而且規則幾乎每次都預測成彎月眉；眼睛的「瞇縫眼」規則使用 ratio_to_face < 0.12，但實際資料的最小值約 0.194，等於這條規則幾乎不可能被觸發。鼻子的窄鼻條件也有相同問題。這讓我知道，把網路上或其他資料集的閾值直接搬過來，不能代表適合我們的資料。"
    )
    add_paragraph(
        doc, created,
        "所以幾何規則沒有完全消失，但角色改成 fallback、資料診斷與特徵解釋。它可以幫我檢查 landmark 是否合理，也可以在模型真的無法使用時維持基本功能；正式的主要分類結果則不能再依賴一組寫死的比例。"
    )

    add_heading(doc, created, "3.2 第 1 代：整張臉 CNN，先做出能訓練的版本", 3)
    add_paragraph(
        doc, created,
        "接下來我用整張臉圖片訓練 CNN，統一縮放到 160×160。這一版的目的不是追求最後準確率，而是確認資料能不能被模型學到，以及整個 train、validation、export 流程是否能跑通。"
    )
    add_paragraph(
        doc, created,
        "跑完後我看到一個很明顯的問題：鼻子和眉毛在整張臉裡面只佔很小的區域，模型可能靠髮型、膚色、背景或拍攝風格猜答案，而不是看真正的部位。更嚴重的是，早期使用隨機切分時，同一個人的照片可能同時出現在 train 和 validation，這會讓分數看起來比真正遇到新使用者時好很多。這一代的數字不能當正式成果，我把它當成「發現問題的基線」。"
    )

    add_heading(doc, created, "3.3 第 2 代：改成 ROI CNN，並開始處理身份洩漏", 3)
    add_paragraph(
        doc, created,
        "我後來把流程改成先切 ROI 再分類：臉型使用比較大的臉部區域，眉、眼、鼻、唇使用各自的區域。當時的基準模型是 ImageNet 預訓練的 MobileNetV3-small，使用 AdamW、cosine scheduler、label smoothing、WeightedRandomSampler 與基本的亮度／對比度增強。這一代比整臉 CNN 更符合任務本身，因為模型看到的畫面裡，真正有用的部位比例變大了。"
    )
    add_paragraph(
        doc, created,
        "我也把切分方式從單純 random split 改成 identity split。同一個身份的照片只能出現在同一個 fold，這樣測試的是模型能不能泛化到沒有看過的人，而不是記住某個人的臉。這個改動一開始讓分數下降，但這個下降反而是重要的，因為它把原本被資料洩漏掩蓋的問題暴露出來。"
    )
    make_table(
        doc, created,
        ["任務", "Random split", "Identity split", "表面差距", "我對結果的解讀"],
        [
            ["臉型", "0.588", "0.475", "+0.113", "原本的高分有身份／風格洩漏疑慮"],
            ["眉型", "0.590", "0.589", "+0.000", "切分影響小，但標籤一致性仍是問題"],
            ["眼型", "0.481", "0.378", "+0.103", "眼型很容易學到個人特徵或拍攝風格"],
            ["鼻型", "0.771", "0.666", "+0.105", "鼻子有可學訊號，但樣本與類別定義要再檢查"],
            ["唇型", "0.479", "0.452", "+0.027", "差距較小，但類別重疊仍然明顯"],
        ],
        widths=[2.2, 2.7, 2.8, 2.4, 8.0],
    )
    add_paragraph(
        doc, created,
        "這張表對我來說比單看最高準確率更有用。臉型與眼型的差距超過 0.10，表示如果我繼續用 random split，很可能會高估模型。從這裡之後，我的正式比較都盡量使用身份分組、固定 fold 和相同訓練設定。"
    )

    add_heading(doc, created, "3.4 第 3 代：先修資料，才有資格比較模型", 3)
    add_paragraph(
        doc, created,
        "在比較更大的模型之前，我先做資料稽核。檔案去重後，我找到 314 個重複檔案，包含 254 組 byte-identical 群組，也發現有 195 組相同圖片被標到不同類別。這不只是資料量問題，而是會直接破壞 validation 的可信度；如果同一張圖或幾乎一樣的圖跑到不同切分，模型會被錯誤地獎勵。"
    )
    add_paragraph(
        doc, created,
        "我也發現舊的 identity map 是用絕對路徑當 key。7 月 29 日移動資料、做類別合併後，很多路徑改變，沒有對上的樣本被默默標成 identity -1。結果 315 張照片幾乎全部留在 train，眼型的桃杏眼總共有 221 張，但 validation 只剩 1 張，某次結果因此出現 recall 1.000。這個數字看起來很好，但其實是驗證集太少，不是模型真的學得很好。"
    )
    add_paragraph(
        doc, created,
        "我的修正方式是重新建立身份群組，確認 identity -1 從 315 張降到只剩 2 張，並確認桃杏眼的 221 張都有合理的 validation 分布。這件事讓我學到，資料切分錯誤不一定會讓程式報錯，反而最危險的是程式順利跑完、只留下看起來漂亮但不能解釋的結果。"
    )
    add_paragraph(
        doc, created,
        "另外，窄鼻資料的幾何分析也和原本的命名不一致：鼻翼寬度中位數是寬鼻約 0.310、窄鼻約 0.291、標準鼻約 0.283。窄鼻不但沒有穩定地落在一個獨立區間，還和標準鼻重疊，因此我把窄鼻併入標準鼻，避免模型被迫學一個資料本身分不清楚的類別。"
    )

    add_heading(doc, created, "3.5 第 4 代：調整訓練設定與類別定義", 3)
    add_paragraph(
        doc, created,
        "資料清理完成後，我才開始比較 epoch、learning rate、augmentation 與類別合併。這一階段我沒有把每次變動混在一起，而是盡量固定資料切分，只改一個主要變因。對小型資料集來說，一次訓練的結果會有 noise，所以我也做了同設定重跑，觀察到大約 ±0.016 的波動；差距小於這個範圍時，我不會直接說某個設定一定比較好。"
    )
    make_table(
        doc, created,
        ["設定", "臉型", "眉型", "眼型", "鼻型", "唇型", "平均"],
        [
            ["20 epochs / 3e-4", "0.470", "0.490", "0.388", "0.827", "0.485", "0.532"],
            ["40 epochs / 6e-4", "0.520", "0.514", "0.434", "0.867", "0.497", "0.566"],
            ["40 epochs / 3e-4", "0.516", "0.521", "0.402", "0.881", "0.492", "0.562"],
            ["40 epochs / 1e-4", "0.487", "0.502", "0.387", "0.833", "0.473", "0.537"],
            ["20 epochs / 1e-4", "0.443", "0.482", "0.378", "0.809", "0.458", "0.514"],
        ],
        widths=[4.0, 2.1, 2.1, 2.1, 2.1, 2.1, 2.1],
    )
    add_paragraph(
        doc, created,
        "從這一輪來看，40 epochs 通常比 20 epochs 穩定；learning rate 3e-4 和 6e-4 在早期結果差距不一定大，但在修正 identity map 後，6e-4 在 18／25 個 fold 的五個任務中都比較常勝出，所以後續以 6e-4 作為主要設定。這不是說 6e-4 對任何新資料都保證最好，而是目前固定 protocol 下比較有一致性。"
    )
    add_paragraph(
        doc, created,
        "類別合併也是模型工作的一部分。眼型曾經做過「杏仁眼與桃花眼合併到圓眼」的探索性實驗，CNN 從 0.388 提升到約 0.523；後來我根據標註語意改成「杏仁眼與桃花眼合併為桃杏眼，圓眼保留」，重新建立 identity map 後，DINOv2 約 0.542、CNN 約 0.463。這兩個數字不能混成同一個結果，因為合併目標和資料版本不同。唇型把 M 型唇併入花瓣唇後，CNN 約從 0.492 提升到 0.560 ± 0.070，提升幅度大於一般重跑 noise，因此保留這個方向。"
    )

    add_heading(doc, created, "四、我實際採用的訓練方式", 2)
    add_heading(doc, created, "4.1 ROI、影像前處理與增強", 3)
    add_paragraph(
        doc, created,
        "目前的分類器不是直接讀原始大圖，而是由 Face Mesh 關鍵點定位後，裁出不同部位的 ROI。臉型需要保留較多外輪廓資訊，其他部位則盡量集中在眉、眼、鼻、唇。這樣可以減少背景與其他五官干擾，也讓同一個模型輸入的語意比較一致。"
    )
    add_paragraph(
        doc, created,
        "早期 ROI CNN 的尺寸是臉型 128×128、其他部位 96×96；後續 ConvNeXt-Tiny 的訓練與匯出使用 224×224 的標準輸入。影像會轉成 RGB、依模型需求 resize，再使用與 ImageNet 預訓練一致的 normalization。augmentation 主要是亮度／對比度 0.85～1.15、約 ±0.06 的小幅色彩變化，以及機率式的小位移；我刻意不使用水平翻轉，因為眉毛高低、唇形不對稱與局部方向可能是標籤的一部分。"
    )

    add_heading(doc, created, "4.2 Identity split、五折交叉驗證與 holdout", 3)
    add_paragraph(
        doc, created,
        "正式比較時，我使用 identity-aware 的五折交叉驗證。先用照片和身份群組建立分組，再讓同一個人只出現在同一個 fold，最後輪流把每一折當 validation。這樣可以減少某一張照片或某一個人的特徵被模型記住。除了五折 CV，我也保留 holdout 概念，用完全不參與調參的資料確認模型是否真的能泛化。"
    )
    add_paragraph(
        doc, created,
        "我看的主要指標不是只有 accuracy，而是 Macro Accuracy、Balanced Accuracy、Macro-F1、每一類的 precision／recall，以及 confusion matrix。類別不平衡時，單看 accuracy 很容易被樣本最多的類別拉高；Macro-F1 可以讓每一類都被算進來，confusion matrix 則可以看出到底是哪兩個標籤互相混淆。"
    )
    add_code(
        doc, created,
        "每個 fold：identity split → train → validation → 每類 precision/recall → Macro Accuracy / Balanced Accuracy / Macro-F1 → 平均 ± 標準差"
    )

    add_heading(doc, created, "4.3 訓練超參數與類別不平衡處理", 3)
    make_table(
        doc, created,
        ["項目", "目前工作紀錄的設定", "我這樣設定的原因"],
        [
            ["主要 backbone", "ConvNeXt-Tiny；同時保留其他架構比較結果", "在五個部位的整體平衡較好，且能輸出 ONNX"],
            ["輸入尺寸", "224×224（ConvNeXt／DINO比較）；早期 ROI CNN 有 128／96", "兼顧局部細節與預訓練模型輸入規格"],
            ["optimizer", "AdamW", "比單純 SGD 更容易在小資料與預訓練 backbone 上微調"],
            ["learning rate", "主要比較 1e-4、3e-4、6e-4，後續以 6e-4 為主", "透過固定 fold 比較收斂速度與跨 fold 一致性"],
            ["epoch", "20 與 40；目前比較偏向 40", "20 epochs 容易還沒有收斂，40 epochs 的平均表現較好"],
            ["class imbalance", "WeightedRandomSampler、label smoothing 0.05", "避免模型只預測樣本數最多的類別"],
            ["augmentation", "亮度／對比度、小位移；不做水平翻轉", "保留五官方向與非對稱特徵"],
            ["輸出", "PyTorch checkpoint、metrics、ONNX、classes JSON、run summary", "讓訓練成果能被回溯、部署與回滾"],
        ],
        widths=[4.0, 6.5, 8.3],
        size=9.0,
    )

    add_heading(doc, created, "五、不同模型的公平比較結果", 2)
    add_paragraph(
        doc, created,
        "我最後比較的重點不是哪一個模型在某一個部位拿到最高，而是五個任務是否使用同一套資料切分、同一組 5-fold、同樣的 40 epochs／learning rate 6e-4，然後看整體是否穩定。以下是同一 protocol 下的 Macro Accuracy 比較，數字是歷史模型選型報告中的結果。"
    )
    make_table(
        doc, created,
        ["架構", "參數量約", "臉型", "眉型", "眼型", "鼻型", "唇型"],
        [
            ["MobileNetV3-small", "1.5M", "0.521", "0.511", "0.562", "0.849", "0.553"],
            ["EfficientNet-B0", "4.0M", "0.564", "0.612", "0.607", "0.855", "0.561"],
            ["ResNet-50", "23.5M", "0.604", "0.587", "0.660", "0.877", "0.578"],
            ["ConvNeXt-Tiny", "27.8M", "0.561", "0.603", "0.693", "0.881", "0.640"],
            ["AlexNet", "57M", "0.494", "0.422", "0.628", "0.867", "0.396"],
        ],
        widths=[4.0, 2.4, 2.1, 2.1, 2.1, 2.1, 2.1],
    )
    add_paragraph(
        doc, created,
        "比較結果是：大型架構相對 MobileNet 在 15 個部位／任務格子中大約有 15 次勝出，但 ConvNeXt-Tiny 並不是每個單項都最高，臉型其實是 ResNet-50 的 0.604 比 ConvNeXt-Tiny 的 0.561 高。我的選擇不是因為 ConvNeXt 每一格都第一，而是它在眼型、鼻型、唇型和整體平衡上比較適合作為統一架構；同時保留單部位結果，避免把「整體選擇」說成「每個部位都最優」。"
    )
    add_paragraph(
        doc, created,
        "這個選擇也有代價。ConvNeXt-Tiny 的 ONNX 檔案約 110 MB，MobileNet 約 6 MB，推論和部署成本都比較高，所以 Gateway、Cloud Run 記憶體、啟動時間和模型快取都要一起考慮。這也是我把模型版本、檔案 hash、健康檢查和 rollback 寫進 deployment 流程的原因。"
    )

    add_heading(doc, created, "六、DINOv2、混合模型與為什麼沒有直接全部上線", 2)
    add_heading(doc, created, "6.1 DINOv2 實驗", 3)
    add_paragraph(
        doc, created,
        "我也測過 DINOv2。做法是使用 frozen ViT-S/14 把 224×224 的 ROI 轉成 384 維 embedding，再訓練 Logistic Regression 或 SVM 的線性分類器。眼型最初 SVM 約 0.625，高於 Logistic Regression 約 0.592；在後續修正資料與類別後，DINOv2 的眼型約 0.542，在 4／5 個 fold 比 CNN 好約 0.047。"
    )
    add_paragraph(
        doc, created,
        "但 DINOv2 對其他部位沒有一致地贏過 CNN，而且原本如果直接保存 joblib，會遇到 sklearn 版本相容性問題。因此我把線性 head 的 coef 和 intercept 匯出成 npz，用 X @ coef.T + intercept 重新計分，降低部署依賴。最後 DINOv2 沒有當成現行主要線上模型，而是保留作為實驗、shadow model 或後續研究比較。這點我要說清楚：有做過實驗，不等於它就是現在服務回傳的模型。"
    )

    add_heading(doc, created, "6.2 Hybrid、majority vote 與 feature fusion", 3)
    add_paragraph(
        doc, created,
        "我也試過把 CNN、DINOv2 和幾何規則混合。結果顯示，三個模型不是每次都在同一張圖答錯，因此多數決和 feature fusion 有機會改善；但把規則硬塞進去的 hybrid 並沒有穩定贏過純 CNN。最明顯的是鼻型，feature fusion 約 0.8920，比單一模型更好；臉型只約 0.5449 的小幅提升；眉、眼、唇則沒有穩定改善。"
    )
    make_table(
        doc, created,
        ["任務", "純模型基準", "Hybrid", "多數決", "Feature fusion", "我最後的判斷"],
        [
            ["臉型", "約 0.5431", "較低", "0.5017", "0.5449", "提升太小，不能宣稱大幅改善"],
            ["眉型", "約 0.5296", "較低", "0.5427", "0.5398", "差異接近 noise"],
            ["眼型", "約 0.6244", "較低", "0.6382", "0.6367", "有提升但要看版本與 fold"],
            ["鼻型", "約 0.8482", "較低", "0.8558", "0.8920", "最有機會保留，但仍需持續驗證"],
            ["唇型", "約 0.5790", "較低", "0.5573", "0.5135", "不採用融合結果"],
        ],
        widths=[2.5, 3.0, 2.5, 2.6, 3.0, 6.0],
        size=8.8,
    )
    add_paragraph(
        doc, created,
        "鼻型 fusion 的分數是把幾何特徵先標準化，再和模型分數加權：score = we × e + wg × (g − μ)／σ + b。匯出時會把平均值、標準差、權重折進 coef 和 intercept，重新計算的最大 decision error 約 5.74e-7，正負方向一致。雖然鼻型在 5 個 fold 贏了 4 次、p 約 0.07，我還是用比較保守的說法，把它當作有重複正向訊號但尚未完全定案的融合實驗。"
    )

    add_heading(doc, created, "七、類別與資料標註的完整變化", 2)
    add_paragraph(
        doc, created,
        "模型換架構只是其中一部分，真正花時間的是把類別定義弄到可以訓練。我的資料裡有些名稱在一般語意上很接近，但標註者未必用同樣標準；如果名稱看起來不同，實際照片卻沒有穩定的視覺差異，模型會學得很痛苦，分數也會受到標註者偏好影響。"
    )
    make_table(
        doc, created,
        ["部位", "早期／探索類別", "後續調整", "我為什麼調整"],
        [
            ["眼型", "六類，另做過杏仁＋桃花→圓眼", "目前四類：下垂、圓、桃杏、鳳", "減少高度重疊類別，保留能說清楚的桃杏眼"],
            ["鼻型", "寬、標準、窄等多種命名", "目前 BASIC 兩類：寬鼻、標準鼻；PRO 側面另五類", "窄鼻分布與標準鼻重疊，不能假裝是清楚獨立類別"],
            ["唇型", "五類，含 M 型唇", "M 型唇併入花瓣唇，現為四類", "M 型唇 recall 約 0.296，錯誤分散，標註邊界不穩"],
            ["眉型", "歷史報告常見三類", "現行類別檔為四類，增加挑眉", "正式文件必須註明版本，不把三類與四類分數混用"],
        ],
        widths=[2.5, 6.0, 5.2, 5.8],
        size=8.9,
    )
    add_paragraph(
        doc, created,
        "PRO 側面鼻型是另一個任務，不應和 BASIC 正面鼻型混在一起。目前類別檔是塌鼻、直挺鼻、翹鼻、蒜頭鼻、駝峰鼻，輸入 224×224，ConvNeXt-Tiny，訓練紀錄顯示 40 epochs、636 筆 train count。因為這一組在目前工作區沒有完整一致的 holdout 指標，我不在文件裡填一個猜測的準確率；它目前可以說有模型、類別檔與部署路徑，但還需要補齊正式評估表。"
    )

    add_heading(doc, created, "八、現行模型到底是哪一版，模型成果放在哪裡", 2)
    add_paragraph(
        doc, created,
        "截至 2026 年 9 月 3 日，工作區的正式 runtime manifest 是 tools/face_models_manifest.json，版本標記為 20260903_brow_nose，模型來源前綴是 gs://decorate-me-models/20260903_brow_nose。BASIC 的臉、眉、眼、鼻、唇與 PRO 側面鼻型目前都以 ConvNeXt-Tiny 為主要 runtime 架構；DINOv2 檔案仍在資料夾裡作為實驗／shadow 資產，但不應直接說成線上主要答案。"
    )
    make_table(
        doc, created,
        ["成果類型", "實際位置", "用途"],
        [
            ["BASIC ONNX", "models/basic_features_roi/face_shape.onnx、brow_shape.onnx、eye_shape.onnx、nose_shape.onnx、lip_shape.onnx", "線上或服務端推論使用的五個主要模型"],
            ["BASIC 類別定義", "models/basic_features_roi/*_classes.json", "記錄 label index 與類別名稱，部署時必須和模型一致"],
            ["PRO 側面鼻型", "models/pro_nose_side/nose_shape_side.onnx 與 nose_shape_side_classes.json", "側面鼻型獨立任務，不能跟 BASIC 正面鼻型混用"],
            ["訓練批次成果", "models/training_runs/<runId>/", "每次 run 的 summary、metrics、checkpoint／ONNX 等產物"],
            ["目前模型清單", "tools/face_models_manifest.json", "記錄版本、GCS prefix、檔案與部署要載入的模型"],
            ["訓練程式", "training/train_basic_cnn_roi.py", "ROI、架構、訓練、CV、export 等主要流程"],
            ["推廣與回滾", "tools/promote_model.py、tools/promotion_worker.py", "檢查模型、上傳、部署、寫 promotion ledger、失敗時回滾"],
        ],
        widths=[4.0, 9.0, 6.0],
        size=8.8,
    )
    add_paragraph(
        doc, created,
        "目前地端訓練批次資料夾裡可以看到多次實際 run，例如 TR-0f4d71078deb5828、TR-ef5d050c5ecd690c、TR-0b26dbf8226765c8、TR-3d51867215768c25、TR-2dd3595461f3c82b 等。不同 run 的部位數不一定相同，有的只有 2 個或 4 個部位，不能看到日期最新就直接當成完整 final model；要以 run summary、metrics、類別檔、manifest 和 promotion ledger 一起核對。"
    )
    add_paragraph(
        doc, created,
        "目前工作區還有一份 cv_report_summary.json，裡面的數字是另一個現行評估產物：臉型 accuracy 約 0.492、Macro-F1 約 0.483；眉型 accuracy 約 0.537、Macro-F1 約 0.507；眼型 accuracy 約 0.580、Macro-F1 約 0.569；鼻型 accuracy 約 0.840、Macro-F1 約 0.839；唇型 accuracy 約 0.573、Macro-F1 約 0.556。這些不是前面 ConvNeXt 選型表的同一種指標與同一版 protocol，所以我把它們標成「現行工作區評估產物」，不和歷史 Macro Accuracy 直接比較。"
    )

    add_heading(doc, created, "九、我在地端怎麼訓練，從按下開始到成果回來", 2)
    add_paragraph(
        doc, created,
        "目前訓練沒有直接放在 Cloud Run 裡面跑，主要原因是 ConvNeXt 單一部位一次可能需要 1～2 小時，曾經量到 TR-f5ae340ef767e1b8 約 112 分 31 秒，而且雲端服務不能直接穩定連到我自己的地端 GPU／資料環境。因此我把「資料和任務在雲端管理」與「重訓練在地端執行」拆開。"
    )
    add_code(
        doc, created,
        "後台建立 training job → Firestore 排隊 → 地端 worker polling → 下載已審核資料 → 合併到 *_plus_feedback → identity split／holdout → 訓練與評估 → 寫入 models/training_runs/<runId>/ → 回傳狀態"
    )
    add_paragraph(
        doc, created,
        "地端 worker 會持續輪詢後台任務，抓取已同意、已修正且通過審核的資料，依五官部位合併進訓練集。訓練時會記錄 runId、資料版本、類別版本、seed、模型架構、epoch、learning rate、fold、每類指標與輸出 hash。訓練完成後並不會自動把線上模型換掉，必須再經過 promotion 檢查，這是我刻意保留的安全閘門。"
    )
    add_paragraph(
        doc, created,
        "地端可靠性也是模型工作留的一部分。訓練期間我會防止電腦睡眠，並用 Task Scheduler 每 5 分鐘檢查 worker；worker 如果因為 DNS、網路或程序中斷而沒有 heartbeat，要能被重新啟動或告警。曾經有 DNS 問題讓 worker 靜默中斷 11 小時，這件事讓我把 watchdog、heartbeat、failed run 與重啟流程列成正式設計，而不是只靠我記得去看終端機。"
    )

    add_heading(doc, created, "十、使用者回饋為什麼重要，以及它怎麼回到模型", 2)
    add_paragraph(
        doc, created,
        "使用者回饋不是單純的滿意度按鈕。對模型來說，它有三個用途：第一是估計真實使用情境下的接受度；第二是找出模型最容易錯的部位和類別；第三是在使用者明確同意的前提下，產生下一輪訓練可以使用的資料。"
    )
    add_code(
        doc, created,
        "模型分析 → 使用者確認／修正 → 明確同意 → Firestore face_feedback + GCS 圖片 → 後台人工審核 → training job → 地端訓練 → 評估 → promotion"
    )
    add_paragraph(
        doc, created,
        "在資料規則上，我沒有把所有回饋都直接拿去訓練。只有使用者明確勾選同意、真的有修正標籤、而且圖片存在時，才會形成可送訓的 contribution。確認但沒有修改的資料可以放進 acceptance evaluation，用來看使用者是否同意系統答案，但不能冒充新的 ground truth。沒有同意時可以保存必要的標籤紀錄，但不保存臉部影像作為訓練資料。"
    )
    add_paragraph(
        doc, created,
        "回饋資料的 owner 由原本的 job 查出來，不信任前端自己送進來的 userId；GCS 路徑會包含 specversion、part、label 與 jobId，例如 user_contributed/<specversion>/<part>/<label>/<jobId>.png。管理者可以採用、修改或拒絕回饋，只有採用後才排進 training job。整條關聯用 jobId → feedbackId → runId 串起來，之後才能回答「這個模型改動到底用了哪些使用者資料」。"
    )
    make_table(
        doc, created,
        ["資料類型", "是否進評估", "是否進訓練", "條件"],
        [
            ["使用者只確認原答案", "可以", "不可以", "可用於 acceptance rate，不代表真實標註"],
            ["使用者修正但未同意影像使用", "依規則保存標籤紀錄", "不可以", "隱私同意優先"],
            ["使用者修正且明確同意、圖片存在", "可以", "可以排隊", "還要經過後台人工審核"],
            ["管理者拒絕的回饋", "不作正式 ground truth", "不可以", "保留決策紀錄以便追蹤"],
            ["PRO 側面鼻型回饋", "獨立計算", "不自動混入 BASIC", "使用自己的類別與模型任務"],
        ],
        widths=[6.0, 3.0, 3.0, 7.0],
        size=8.9,
    )
    add_paragraph(
        doc, created,
        "截至既有紀錄，系統曾累積 126 筆回饋；其中的 acceptance 是「使用者是否接受系統答案」，不是嚴格意義的模型 accuracy。既有 Wilson 95% 區間顯示，眉型接受率約 78.6%、鼻型約 87.3%，眼型約 59.5%、臉型約 61.9%、唇型約 65.9%。這些數字很適合拿來找優先改善順序，但不能直接和五折分類準確率畫等號。"
    )

    add_heading(doc, created, "十一、Gateway、前端、後台和模型的銜接", 2)
    add_paragraph(
        doc, created,
        "我一開始讓前端直接呼叫不同服務，結果 API URL、金鑰、身份驗證和錯誤處理散落在各個地方。後來改成由 AI Gateway 當主要入口，前端只需要對 Gateway 說明要做哪一種工作，Gateway 再依 allowlist 將請求送到 face analysis、text suggestion、render 或 product recommendation。這次換 Gateway 不是為了多一層而多一層，而是要把模型版本、session、CSRF、IAM、timeout、錯誤回應和媒體存取集中管理。"
    )
    add_paragraph(
        doc, created,
        "Gateway 和模型工作留的關係是：模型訓練成果先落在 training_runs，經過 promotion 產生版本化 manifest，再由服務端載入。Gateway 回傳的 analysisPackage 應該帶有 modelVersion，這樣前端看到使用者回饋時，後台才知道那筆回饋是針對哪一版模型。若只記錄最後的 label，沒有版本、runId 和原始 predicted label，之後就無法判斷是模型退化、類別變更，還是使用者只是不同意主觀分類。"
    )
    add_paragraph(
        doc, created,
        "目前 Gateway 還處理幾個模型以外但很重要的問題：Firebase CDN 只穩定轉送 __session，所以系統在 session cookie 中保存 JWT 與上游加密 cookie；多使用者情境用 session slot 與 actor 隔離；寫入請求使用 double-submit CSRF；路由使用 allowlist 與 fullmatch；文字建議失敗時 fail closed；私有 GCS 媒體不直接公開，而由後端 proxy。這些設計會影響模型結果能不能安全地到達正確使用者。"
    )
    make_table(
        doc, created,
        ["層", "主要技術／責任", "和模型工作留的關係"],
        [
            ["前端", "Next.js／React、上傳、結果確認、修正與 consent", "呈現 label、confidence、版本與回饋入口"],
            ["Gateway", "FastAPI、session、CSRF、IAM、route allowlist、timeout", "統一模型 API 入口，保存可追蹤的 job 與版本資訊"],
            ["模型服務", "Cloud Run／ONNX Runtime／Face Mesh", "載入 manifest 指定的 ONNX 與類別檔，產生 analysisPackage"],
            ["後台", "任務佇列、回饋審核、training run、promotion", "決定哪些回饋可以進訓練，以及哪個模型可以上線"],
            ["資料層", "Firestore、GCS", "保存 job、feedback、圖片、run 關聯與模型檔"],
        ],
        widths=[3.0, 8.0, 8.0],
        size=8.9,
    )

    add_heading(doc, created, "十二、我在模型過程中遇到的問題與實際修正", 2)
    make_table(
        doc, created,
        ["問題", "為什麼會發生", "我怎麼發現", "修正與留下的防線"],
        [
            ["隨機切分分數過高", "同一身份照片跨 train／validation", "random 與 identity split 差距大", "identity-aware 五折、holdout、禁止身份跨 fold"],
            ["重複圖片跨類別", "檔案複製或標註重複", "byte hash 與 image hash 稽核", "去重、列出衝突群組，不能只刪一個檔名"],
            ["identity map 過期", "移動檔案後絕對路徑 key 失效", "identity -1 暴增、某類 validation 只剩 1 張", "重建 map、檢查 -1、檢查每類每 fold 樣本數"],
            ["類別名稱重疊", "不同標註者用不同語意邊界", "confusion matrix 與低 recall", "合併類別、補標註規則、在文件標示版本"],
            ["cache 維度不一致", "只檢查 cache 長度沒有檢查 embedding 維度", "DINO／模型切換後輸出異常", "cache 同時驗證長度、維度和模型版本"],
            ["訓練跑完但沒有上線", "訓練產物和線上 manifest 是兩件事", "本地 metrics 好但服務版本未變", "promotion ledger、manifest hash、health／smoke／rollback"],
            ["worker 靜默停止", "DNS、睡眠、鎖定或網路中斷", "heartbeat 中斷 11 小時", "Task Scheduler、restart loop、cloud alert、training keep-awake"],
            ["LLM 回傳格式不固定", "文字模型有時漏欄位或混入額外文字", "parser 讀不到必要 section", "schema validation、必要欄位檢查、失敗時不送錯誤建議"],
        ],
        widths=[4.0, 5.0, 5.0, 6.0],
        size=8.5,
    )
    add_paragraph(
        doc, created,
        "我後來把這些問題都當成模型系統的一部分，而不是訓練腳本以外的雜事。因為只要資料洩漏、版本對不上或 worker 靜默停止，最後展示給使用者的結果就算形式上有回傳，也不能說是可信的模型成果。"
    )

    add_heading(doc, created, "十三、我把介面修改整理成 UAT 測試工作留", 2)
    add_paragraph(
        doc, created,
        "其實我前面反覆要求調整的很多介面，不是單純改顏色或換按鈕位置，而是使用者驗收測試（UAT）的概念：我用實際使用者的角度操作流程，先找出畫面上看起來可以用、但實際上會誤導人、資料會錯、或下一步做不下去的問題，再把問題原因和改善結果留下來。下面這張表把這些修改用同一種格式整理，讓人可以看出每一次介面調整背後對應的是哪一個使用情境。"
    )
    add_paragraph(
        doc, created,
        "我把 UAT 的判斷標準定成三件事：第一，使用者看得懂現在系統在哪個狀態；第二，使用者的操作不會讓資料重複、遺失或送錯；第三，遇到服務失敗時，畫面要說真正的原因，不能把系統錯誤包裝成登入失效、沒有資料或模型已經成功。"
    )
    make_table(
        doc, created,
        ["發現問題", "為什麼有問題", "改善過後"],
        [
            ["臉部分析結果頁：分析尚未完成時，前端容易把 409 或暫時狀態當成失敗。", "背景分析本來就需要時間，使用者如果看到錯誤就會重送照片，可能產生多個 job，也不知道原本的分析其實還在跑。", "我把 jobId、狀態查詢與結果查詢分開，尚未完成就維持等待／輪詢；只有真正失敗才顯示失敗，並讓前端保留目前工作狀態。"],
            ["妝容渲染：第一次請求逾時後，前端自動重試，畫面卻顯示渲染失敗。", "伺服器其實已建立第一個 job，第二次請求只是撞到同一個進行中的工作；如果 409 沒有回傳原始 jobId，前端就無法接續輪詢。", "我讓重複請求回傳原始 jobId 與 resultToken，前端可以接續原本的 job，不把正在處理的渲染誤判成失敗。"],
            ["妝前／妝後圖片：管理員要檢查妝容是否真的保留本人，但妝前圖被權限擋住。", "只看妝後圖分不出模型畫得好，還是把使用者的臉換掉；但如果直接放寬所有人，又會讓陌生人讀到生物特徵資料。", "我只放寬管理員且保留 audit 記錄；本人仍可看自己的妝前圖，陌生人沒有 admin 旗標仍然被擋，並要求 Gateway 和圖片服務的權限判斷一致。"],
            ["臉部結果回饋：使用者按一次送出後又修改，系統可能出現兩筆回饋。", "前端提示「已送出，可再修改」，所以修改是正常流程；如果每次都新增文件，評估分母會被灌大，訓練集也會留下已經改回去的舊標籤。", "我用 jobId 作為回饋文件識別，重送同一個 job 會覆寫最後狀態；改回原答案時會清除 corrected 與 training record，不重複計算。"],
            ["使用者同意訓練：沒有勾選同意、沒有圖片，或 GCS 實際一張都沒存，後台仍可能顯示可以送訓。", "標籤回饋存在不代表影像貢獻成功；如果 contributed 說謊，管理員會看到空白樣本，還可能把零樣本批次送進訓練。", "我把 contributed 定義成實際成功保存至少一張影像；無同意、無圖片、GCS 失敗或 ROI 無法擷取時，回饋仍可保存，但不當成可訓練影像。"],
            ["後台回饋列表：舊資料沒有 reviewStatus 時，畫面可能把它當成已採用。", "舊資料沒有狀態不代表管理員審核過，若預設成 approved，未審核的使用者標籤會直接進入送訓候選。", "我把缺少 reviewStatus 的資料統一視為 pending，最新回饋排在前面，並只把真的有 predicted／correction 差異的欄位列在 changes。"],
            ["後台看樣本影像：回饋列表有資料，但點開是空白圖。", "GCS 可能讀取失敗、檔案超過單次下載限制，或 ROI 根本沒有成功保存；如果畫面只顯示一個空白框，管理員會誤以為是前端壞掉，也可能誤送訓。", "我把影像保存結果和 contributed 分開記錄，沒有實際樣本就不能正常送訓；後台要顯示可辨識的缺圖／讀取失敗狀態，而不是假裝有圖。"],
            ["後台採用／退回：所有上游錯誤都被顯示成同一種 503。", "「這筆資料不存在」、「權限不足」和「臉部服務掛掉」的處理方式不同，全部顯示服務錯誤會讓管理員對不存在的資料一直重試。", "我保留可處理的 4xx 原因，只有真正的上游不可用才回 503，並在 Gateway 測試中確認管理員身分、CSRF 和 actor header 都要通過。"],
            ["後台訓練批次：訓練機失敗時，heartbeat 還在跳，畫面卻沒有失敗原因。", "原本監控只知道機器還活著，不知道訓練程序已經炸掉；管理員會以為批次仍在跑，無法判斷要不要重試。", "我先把 status=failed 和 error 寫回 Firestore，再盡力推送失敗指標；錯誤會截斷到可保存的長度，空錯誤也補上可讀訊息。"],
            ["後台訓練批次：管理員找不到某次訓練的完整前後模型紀錄，甚至以為批次被刪了。", "訓練 run 是模型工作留的唯一證據，若可以刪除或只留下畫面顯示，就無法證明當時用了哪些資料、模型從哪版換到哪版。", "我把 training runs 做成 append-only，保留 startedAt、finishedAt、modelBefore、modelAfter、metrics 與 error；後台只提供建立／查詢／失敗重試，不提供刪除路徑。"],
            ["會員儲存稽核：會員名冊讀取失敗時，畫面把全部渲染資料列成孤兒。", "空名冊和「真的沒有會員」不是同一件事；管理員照著錯誤清單刪資料，可能刪掉仍屬於會員的圖片。訪客也不應被算成遺失會員。", "我在名冊讀不到時直接停止稽核並回報 MEMBER_DIRECTORY_UNREADABLE；結果分開顯示會員、沒有儲存資料的會員、訪客與真正 orphan，也不把 email 回傳給瀏覽器。"],
            ["商品後台：商品上游回 401／403 時，前端顯示「請重新登入」。", "管理員 session 其實已經驗證成功，錯的是商品服務金鑰或 URL；一直重新登入不會修好上游拒絕，還會讓人誤判會員系統壞掉。", "我把上游授權拒絕轉成 PRODUCT_UPSTREAM_REJECTED 的 502，訊息明確說明登入有效、問題在商品服務連線或金鑰；404、409、422 等資料層錯誤則保留原意。"],
            ["商品公開 API：前端呼叫新跨品牌色號路徑時，Gateway 自己回 404。", "Gateway 路由白名單沒有放行時，錯誤長得像商品上游根本沒有端點，前端和後台會把兩個問題混在一起。", "我補上明確的 route allowlist 與方法檢查，並用 Gateway 測試確認路徑與 HTTP method 都正確，讓「Gateway 沒開」和「上游沒做」能分開判斷。"],
            ["前端建議頁：畫面可能直接顯示固定 style.advice，或繼續切割長篇英文渲染指令。", "固定文字會冒充 AI 已完成；前端自己切割 LLM 長文也會受格式變更影響，且可能把英文 prompt 露給使用者或錯誤保存。", "我把正式方向改成由 Ollama 回傳結構化欄位，前端只呈現中文建議與必要結果；但目前正式 200 response 契約仍未完全驗證，所以這一列標記為待完成，暫不把現況誤報成 PASS。"],
            ["文字建議服務：Ollama 不可達時，系統仍可能顯示上一段或固定的建議。", "使用者會以為這次分析成功，實際上建議可能不是這張照片的結果，後續渲染 prompt 也會跟錯。", "我讓 Gateway／文字服務在沒有可用上游時 fail closed，回傳明確的 EXTERNAL_TEXT_UPSTREAM_DISABLED 或 OLLAMA_UNAVAILABLE，不使用假 fallback；正式端到端仍需補正確 API key 才能驗收成功回應。"],
            ["後台送訓：畫面上的 sampleCount 容易被理解成照片數或使用者數。", "實際 selections 可能是一個回饋裡的多個部位，sampleCount 計的是部位數；如果管理員誤讀，會錯估訓練資料量。", "我在批次紀錄保留 selections 快照，並把 sampleCount 的定義寫清楚，另外核對實際 GCS 物件數、部位數、feedbackIds 與 runId，不用一個數字假裝代表所有數量。"],
            ["回饋統計：使用者確認系統答案和使用者修正被混成同一種 accuracy。", "確認代表接受目前答案，修正才代表使用者不同意；兩者混在一起會讓模型看起來比實際好，也無法知道哪個部位最常被改。", "我把 acceptance evaluation 和 training contribution 分開，保留 predicted、agreed、corrected、contributed，並提醒 acceptance rate 不能直接當成分類模型 accuracy。"],
            ["資料匯總：基本五官五筆和額外的側臉鼻型資料加在一起，總數對不起來。", "如果只看總數，會出現例如 649 對不上 645 的問題，但不知道多出的資料屬於哪個部位，演算法端無法追查。", "我在彙總結果增加 extraParts 與 consistencyChecks，把 BASIC 五官固定五部位，PRO 側臉鼻型獨立列出，並在缺少 feedback 時留 null，不自己猜標籤。"],
            ["模型部署：本地訓練完成後，線上服務仍載入舊模型。", "訓練輸出、GCS 上傳、manifest 更新和 Cloud Run revision 是不同步驟；只看到本地 metrics 好，不能證明使用者已經用到新模型。", "我增加 modelVersion、SHA-256、manifest、promotion ledger、health check、smoke test 與 rollback；只有完成 promotion 並核對 revision，才把它稱為已上線。"],
        ],
        widths=[6.1, 7.1, 7.1],
        size=8.0,
    )
    add_paragraph(
        doc, created,
        "這張表也說明為什麼我把很多看起來像前端問題的修改，最後都放進後台、Gateway、資料庫或測試裡一起處理。UAT 不是只看畫面有沒有跑出來，而是從使用者按下按鈕開始，一路確認狀態、權限、資料保存、錯誤訊息和後續流程都沒有被誤導。對我來說，改善過後不是「畫面看起來比較漂亮」，而是使用者真的能完成任務，管理員能知道下一步，系統也留下可以追蹤的證據。"
    )
    add_paragraph(
        doc, created,
        "對應的自動化驗收證據主要放在 tests/ai_gateway_test.py、tests/render_api_test.py、tests/face_feedback_test.py、tests/face_contributions_test.py、tests/training_run_store_test.py、tests/export_feedback_aggregate_test.py，以及前端的 node --check 與 smoke test。這些測試不是取代人工 UAT，而是把我曾經實際發現過的問題固定下來，避免下次重構又回到原本的錯誤。"
    )

    add_heading(doc, created, "十四、現階段結果要怎麼解讀，哪些話不能亂講", 2)
    add_paragraph(
        doc, created,
        "歷史選型報告中，ConvNeXt-Tiny 五折 Macro Accuracy 是臉型 0.561 ± 0.038、眉型 0.603 ± 0.022、眼型 0.693 ± 0.041、鼻型 0.881 ± 0.056、唇型 0.640 ± 0.032。這可以支持「在當時固定的資料版本與 protocol 下，ConvNeXt-Tiny 是整體較平衡的選擇」，但不能延伸成「每一張臉都能正確判斷」，也不能忽略主觀標籤與類別不平衡。"
    )
    make_table(
        doc, created,
        ["部位", "五折 Macro Accuracy", "Wilson 使用者接受率", "我會怎麼說"],
        [
            ["臉型", "0.561 ± 0.038", "61.9%（53.2–69.9%）", "仍要改善，臉型受主觀與外輪廓影響大"],
            ["眉型", "0.603 ± 0.022", "78.6%（70.6–84.8%）", "使用者接受度較高，但類別版本要講清楚"],
            ["眼型", "0.693 ± 0.041", "59.5%（50.8–67.7%）", "分類分數不低，真實使用接受度仍需改善"],
            ["鼻型", "0.881 ± 0.056", "87.3%（80.4–92.0%）", "目前最穩定，但 BASIC／PRO 要分開"],
            ["唇型", "0.640 ± 0.032", "65.9%（57.2–73.6%）", "仍受標籤重疊與妝容影響"],
        ],
        widths=[2.5, 4.0, 6.3, 7.2],
        size=8.7,
    )
    add_paragraph(
        doc, created,
        "我會把「模型準確率」和「使用者接受率」分開報告。前者是在固定標註與驗證切分下，模型和資料標籤的一致程度；後者是使用者是否覺得答案符合自己，會受到審美、命名理解和展示方式影響。兩者不一樣，但一起看才知道系統在實際使用時哪裡需要改善。"
    )
    add_paragraph(
        doc, created,
        "目前還有一個文件工作要持續補：現行四類眉型與歷史三類眉型、現行模型 manifest 與舊報告、BASIC 正面鼻型與 PRO 側面鼻型，都要在每份 metrics 旁邊綁定 modelVersion、classVersion、datasetVersion、protocol。沒有這些欄位時，我不應該把不同檔案的結果合成一個「最新準確率」。"
    )

    add_heading(doc, created, "十五、如何從地端成果重現一次模型工作", 2)
    add_paragraph(
        doc, created,
        "如果我要讓別人重現我的訓練，不是只把一個 .onnx 檔交出去，而是要按照以下順序把條件固定。這也是我認為模型工作留最重要的部分：別人能知道我輸入了什麼、怎麼切分、跑了什麼程式、最後哪個檔案真的被部署。"
    )
    for item in [
        "先確認資料版本、五個任務的類別檔和身份 map，並檢查重複圖片、跨類別衝突與 identity -1。",
        "固定 seed、fold、輸入尺寸、normalization、augmentation、epoch、batch size、learning rate 和 backbone。",
        "執行 identity-aware 五折與 holdout，不把 random split 的結果拿來和正式結果混在一起。",
        "輸出每個 fold 的 confusion matrix、每類 precision／recall、Macro Accuracy、Balanced Accuracy、Macro-F1 與平均／標準差。",
        "把 checkpoint、ONNX、classes JSON、metrics、training summary、環境版本和 git commit 放在同一個 runId 目錄。",
        "訓練完成後先做 class consistency、檔案 hash、ONNX load、health check 和 smoke test，再交給 promotion。",
        "部署後確認線上 manifest、Cloud Run revision、modelVersion 與實際回傳欄位一致；失敗時用 ledger 指定的上一版回滾。",
    ]:
        add_bullet(doc, created, item)
    add_code(
        doc, created,
        "python training/train_basic_cnn_roi.py --part <part> --arch convnext_tiny --epochs 40 --lr 6e-4 --identity-split --run-id <runId>"
    )
    add_paragraph(
        doc, created,
        "實際執行參數要以當次 run summary 為準，上面的指令是工作流程示意。因為現在 repository 裡同時存在歷史模型、實驗模型、DINO head、不同類別版本和多個 training_runs，所以我不把一條命令寫成永遠不變的真相，而是要求每次 run 自己保存完整參數。"
    )

    add_heading(doc, created, "十六、我最後的總結：模型不是換到最大，而是換到比較能負責任地解釋", 2)
    add_paragraph(
        doc, created,
        "回頭看，我的模型歷程不是單純從小模型一路換到大模型。真正的順序是：先用規則建立可解釋基線，再發現整臉 CNN 會受小部位與背景影響，改成 ROI；接著發現 random split 和重複資料會讓分數失真，所以改做 identity split、資料去重和五折；再來處理類別定義、訓練超參數和不同 backbone；最後才把模型產物、回饋送訓、Gateway、部署與 rollback 接起來。"
    )
    add_paragraph(
        doc, created,
        "我最後選 ConvNeXt-Tiny，不是因為它在每一個部位都第一，而是因為在同一套公平 protocol 下，它對眼型、鼻型、唇型和整體表現較平衡，並且能和目前的 ONNX／manifest／Cloud Run 流程接起來。DINOv2、幾何規則、feature fusion 也沒有被我說成完全沒用，而是根據結果放到適合的位置：規則做 fallback，DINO 做實驗或 shadow，鼻型 fusion 保留研究方向，純 ConvNeXt 做現行主要 runtime。"
    )
    add_paragraph(
        doc, created,
        "目前我最清楚的限制是資料量仍小、標籤有主觀性、不同版本的類別不完全一致，且「模型分數高」不代表使用者一定接受。因此下一步不是盲目再換一個更大的模型，而是補足身份分組資料、固定 holdout、補齊 PRO 側面鼻型評估、把 dataset／class／protocol 版本綁進所有 metrics，再用經過審核的使用者回饋做可追蹤的增量訓練。"
    )
    add_paragraph(
        doc, created,
        "這樣整理後，我在口頭報告時可以完整回答：我換過哪些模型、為什麼換、哪次結果不能相信、我如何修正資料和切分、我在哪台地端訓練、成果實際存在哪裡、使用者回饋怎麼影響下一輪訓練，以及為什麼最後要透過 Gateway 和 promotion 才能讓模型上線。"
    )

    # 將新增內容整段移到原文件的「結論及未來發展」之前，保留原始章節順序。
    target = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("陸、結論及未來發展"):
            target = p
            break
    if target is None:
        # 若來源文件文字略有差異，退而放在原內容最後。
        target = None
    else:
        for element in created:
            target._p.addprevious(element)

    doc.save(OUTPUT)

    # 重新讀取做結構性基本檢查，避免輸出檔未寫完整。
    check = Document(OUTPUT)
    appendix_found = any("臉部分析模型完整開發歷程與訓練工作留" in p.text for p in check.paragraphs)
    if not appendix_found:
        raise RuntimeError("輸出文件中找不到新增模型歷程章節")
    print(f"OUTPUT={OUTPUT}")
    print(f"PARAGRAPHS={len(check.paragraphs)}")
    print(f"TABLES={len(check.tables)}")
    print("APPENDIX_FOUND=True")


if __name__ == "__main__":
    main()
