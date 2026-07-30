from __future__ import annotations

import ast
import re
import textwrap
from pathlib import Path

try:
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, Ellipse, Circle
except ImportError:
    plt = FancyBboxPatch = Ellipse = Circle = None
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    Document = None

ROOT = Path.cwd()
FRONT = Path(r"C:\Users\user\OneDrive - 淡江大學\Desktop\web_frontend")
OUT = next((ROOT / "docs").glob("*2026-07-31"))
FIG = OUT / "系統圖集_PNG"
FIG.mkdir(exist_ok=True)

SERVICES = ["ai_gateway.py", "Face_analyzer_BASIC.py", "Face_analyzer_PRO.py", "Ollama_suggestion.py", "replicate_render_api.py", "face_feedback.py"]
COLORS = {"front":"#E8F1FB","gateway":"#DCE6F1","service":"#E9F4EA","data":"#FFF2CC","external":"#FCE4D6","line":"#44546A"}

def setup_plot(title, w=16, h=9):
    plt.rcParams["font.family"] = ["Microsoft JhengHei", "DejaVu Sans"]
    fig, ax = plt.subplots(figsize=(w,h), dpi=180)
    ax.set_xlim(0,100); ax.set_ylim(0,100); ax.axis("off")
    ax.set_title(title, fontsize=20, fontweight="bold", pad=16, color="#17365D")
    return fig, ax

def box(ax,x,y,w,h,text,color,fs=10):
    p=FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.5,rounding_size=1.6",facecolor=color,edgecolor="#5B6573",linewidth=1.2)
    ax.add_patch(p); ax.text(x+w/2,y+h/2,text,ha="center",va="center",fontsize=fs,wrap=True)
    return (x,y,w,h)

def arrow(ax,a,b,label="",style="-"):
    ax.annotate("",xy=b,xytext=a,arrowprops=dict(arrowstyle="->",color=COLORS["line"],lw=1.4,linestyle=style))
    if label: ax.text((a[0]+b[0])/2,(a[1]+b[1])/2+1.5,label,ha="center",fontsize=8,color="#404040",bbox=dict(facecolor="white",edgecolor="none",pad=1))

def savefig(fig,name):
    p=FIG/name; fig.tight_layout(); fig.savefig(p,bbox_inches="tight",facecolor="white"); plt.close(fig); return p

def architecture():
    fig,ax=setup_plot("DecorateMe AI 系統架構圖")
    box(ax,3,42,15,14,"使用者瀏覽器\nWeb SPA",COLORS["front"],12)
    box(ax,23,42,18,14,"Firebase Hosting\n同源 Rewrite",COLORS["front"],11)
    box(ax,46,40,18,18,"AI Gateway\nSession／CSRF／Proxy\n權限／稽核",COLORS["gateway"],11)
    services=[("Face BASIC",72,76),("Face PRO",72,59),("Ollama 建議",72,42),("Replicate Render",72,25),("會員／商品 API",72,8)]
    for t,x,y in services: box(ax,x,y,20,11,t,COLORS["service"] if "API" not in t else COLORS["external"],10)
    box(ax,45,10,19,12,"Firestore Job Store",COLORS["data"],10); box(ax,45,70,19,12,"GCS 私有媒體",COLORS["data"],10)
    arrow(ax,(18,49),(23,49),"HTTPS"); arrow(ax,(41,49),(46,49),"同源 HTTPS")
    for _,x,y in services: arrow(ax,(64,49),(x,y+5.5),"API key／JSON")
    arrow(ax,(55,40),(55,22),"Job 狀態"); arrow(ax,(55,58),(55,70),"媒體")
    ax.text(4,90,"信任邊界：瀏覽器不持有上游金鑰；所有正式請求由 Gateway 驗證與代理",fontsize=11,color="#9B1C1C")
    return savefig(fig,"01_系統架構圖.png")

def user_flow():
    fig,ax=setup_plot("使用者端到端流程圖",18,8)
    steps=["進入網站","登入／註冊\nOTP","上傳妝前照","BASIC／PRO\n臉部分析","選擇妝容風格","文字建議／\n商品推薦","AI 妝容渲染","前後比較","收藏／歷史／\n會員中心"]
    xs=[2,13,24,35,47,59,71,82,91]
    for i,(x,t) in enumerate(zip(xs,steps)):
        box(ax,x,45,8.5 if i<8 else 8,16,t,COLORS["front"] if i<3 else COLORS["service"],9)
        if i: arrow(ax,(xs[i-1]+(8.5 if i-1<8 else 8),53),(x,53))
    box(ax,35,17,12,10,"品質不合格\n重新上傳",COLORS["external"],9); arrow(ax,(39,45),(41,27),"無臉／姿態／畫質"); arrow(ax,(35,22),(28,45),"重試")
    box(ax,71,17,12,10,"渲染失敗\n保留分析草稿",COLORS["external"],9); arrow(ax,(75,45),(77,27),"timeout／上游錯誤"); arrow(ax,(71,22),(64,45),"仍可看建議")
    return savefig(fig,"02_使用者流程圖.png")

def api_flow():
    fig,ax=setup_plot("API 串接與認證資料流圖",16,10)
    lanes=[("Browser",88,COLORS["front"]),("Hosting",72,COLORS["front"]),("Gateway",56,COLORS["gateway"]),("AI／資料服務",38,COLORS["service"]),("Firestore／GCS",20,COLORS["data"])]
    for name,y,c in lanes:
        box(ax,3,y-4,16,8,name,c,10); ax.plot([21,97],[y,y],color="#D0D5DC",lw=0.8)
    events=[(25,"1. GET /auth/session",88,56),(37,"2. Set-Cookie __session\n+ CSRF",56,88),(49,"3. POST /face-basic/...\nCookie + X-CSRF-Token",88,56),(61,"4. X-API-Key + sanitized input",56,38),(73,"5. job/result",38,56),(83,"6. owner-authorized media",56,20),(93,"7. normalized response",56,88)]
    for x,label,ya,yb in events:
        arrow(ax,(x,ya),(x,yb),label)
    ax.text(3,7,"重要：Session actor 是唯一身分來源；前端 email／owner 只能做一致性檢查。錯誤以 code、message、correlationId 正規化。",fontsize=10,color="#17365D")
    return savefig(fig,"03_API串接圖.png")

def use_case():
    fig,ax=setup_plot("DecorateMe AI 使用案例圖",16,10)
    # actors
    for x,label in [(8,"訪客"),(8,"會員"),(92,"管理員"),(92,"外部服務")]:
        ax.add_patch(Circle((x,78 if label in ["訪客","管理員"] else 30),2,fill=False,color="#44546A")); y=78 if label in ["訪客","管理員"] else 30
        ax.plot([x,x],[y-2,y-10],color="#44546A"); ax.plot([x-4,x+4],[y-5,y-5],color="#44546A"); ax.plot([x,x-4],[y-10,y-15],color="#44546A"); ax.plot([x,x+4],[y-10,y-15],color="#44546A"); ax.text(x,y-20,label,ha="center",fontsize=10)
    cases=[("註冊／登入／OTP",30,80),("執行 BASIC／PRO 分析",52,80),("取得妝容建議",32,60),("取得商品推薦",52,60),("建立 AI 渲染",72,60),("收藏與查看歷史",32,40),("管理會員／商品",70,80),("審核爬蟲暫存商品",70,40),("刪除帳號與媒體",52,25)]
    for t,x,y in cases:
        e=Ellipse((x,y),22,9,facecolor="#F4F7FB",edgecolor="#5B6573"); ax.add_patch(e); ax.text(x,y,t,ha="center",va="center",fontsize=8.5)
    links=[((12,75),(19,80)),((12,30),(21,40)),((12,30),(21,60)),((12,30),(41,80)),((12,30),(41,25)),((88,75),(81,80)),((88,75),(81,40)),((88,30),(83,60)),((88,30),(63,25))]
    for a,b in links: ax.plot([a[0],b[0]],[a[1],b[1]],color="#7A8491",lw=1)
    return savefig(fig,"04_使用案例圖.png")

def gantt():
    fig,ax=setup_plot("專題開發與整合甘特圖（文件化重建）",16,9)
    phases=[("需求／接口盤點",0,2),("Web 前端與會員流程",1,5),("BASIC／PRO 分析",2,6),("模型資料與訓練",3,8),("Ollama 建議",5,3),("Replicate 渲染",5,4),("Gateway／資安整合",7,4),("端到端測試與修復",9,3),("雲端部署／Demo",11,2),("文件與面試準備",10,4)]
    ax.set_xlim(-3.8,14); ax.set_ylim(-1,len(phases)); ax.set_yticks([]); ax.invert_yaxis(); ax.set_xticks(range(14)); ax.set_xticklabels([f"W{i+1}" for i in range(14)]); ax.grid(axis="x",color="#D9DEE5",lw=.8)
    for i,(task,start,dur) in enumerate(phases):
        ax.text(-3.65,i,task,ha="left",va="center",fontsize=9,color="#24364B")
        ax.barh(i,dur,left=start,height=.58,color="#5B9BD5" if i<7 else "#70AD47",edgecolor="white")
        ax.text(start+dur/2,i,f"{dur} 週",ha="center",va="center",fontsize=8,color="white",fontweight="bold")
    ax.set_xlabel("專題週次（依實際校程可調整）")
    return savefig(fig,"05_甘特圖.png")

def activity():
    fig,ax=setup_plot("臉部分析、建議與渲染活動圖",12,13)
    nodes=[("開始",50,93,10,5,"start"),("驗證 Session／CSRF",50,84,28,7,"box"),("驗證圖片與移除 metadata",50,73,32,7,"box"),("品質是否合格？",50,62,25,8,"decision"),("執行 BASIC／PRO 推論",50,50,30,7,"box"),("建立 analysisPackage",50,40,28,7,"box"),("並行：文字建議／商品推薦",50,30,36,7,"box"),("使用者是否要求渲染？",50,20,30,8,"decision"),("建立 Render Job／輪詢",27,9,28,7,"box"),("保存／比較／結束",72,9,25,7,"box")]
    for t,x,y,w,h,kind in nodes:
        if kind=="decision":
            ax.scatter([x],[y],s=2800,marker="D",facecolor="#FFF2CC",edgecolor="#5B6573"); ax.text(x,y,t,ha="center",va="center",fontsize=8)
        elif kind=="start": ax.add_patch(Ellipse((x,y),w,h,facecolor="#D9EAD3",edgecolor="#5B6573")); ax.text(x,y,t,ha="center",va="center",fontsize=9)
        else: box(ax,x-w/2,y-h/2,w,h,t,COLORS["service"],9)
    for a,b,l in [((50,90.5),(50,87.5),""),((50,80.5),(50,76.5),""),((50,69.5),(50,66),""),((50,58),(50,53.5),"是"),((50,46.5),(50,43.5),""),((50,36.5),(50,33.5),""),((50,26.5),(50,24),""),((45,17),(32,12.5),"是"),((55,17),(67,12.5),"否"),((41,62),(15,62),"否：回傳品質錯誤")]: arrow(ax,a,b,l)
    box(ax,3,58,22,8,"提示重拍／重新上傳",COLORS["external"],8); arrow(ax,(14,58),(14,73),"重試")
    return savefig(fig,"06_活動圖.png")

def all_figures(): return [architecture(),user_flow(),api_flow(),use_case(),gantt(),activity()]

def set_font(run,size=10.5,bold=False,color="000000"):
    run.font.name="Microsoft JhengHei"; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)

def config(doc,title):
    s=doc.sections[0]; s.top_margin=s.bottom_margin=Inches(.75); s.left_margin=s.right_margin=Inches(.8)
    st=doc.styles["Normal"]; st.font.name="Microsoft JhengHei"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); st.font.size=Pt(10); st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.18
    for n,z in [("Heading 1",17),("Heading 2",14),("Heading 3",11.5)]:
        x=doc.styles[n]; x.font.name="Microsoft JhengHei"; x._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei"); x.font.size=Pt(z); x.font.bold=True; x.font.color.rgb=RGBColor.from_string("2E74B5"); x.paragraph_format.keep_with_next=True
    p=s.header.paragraphs[0]; set_font(p.add_run("DecorateMe AI｜"+title),8,color="666666")
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("專題技術答辯基準｜2026-07-31"),8,color="666666")

def title(doc,text,sub):
    for _ in range(3): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(text),26,True,"17365D")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run(sub),12,color="666666")
    doc.add_page_break()

def add(doc,text,style=None):
    p=doc.add_paragraph(style=style); set_font(p.add_run(text),10); return p

def endpoint_records():
    out=[]
    for fn in SERVICES:
        src=(ROOT/fn).read_text(encoding="utf-8-sig"); lines=src.splitlines(); tree=ast.parse(src)
        for node in ast.walk(tree):
            if not isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef)): continue
            routes=[]
            for dec in node.decorator_list:
                if isinstance(dec,ast.Call) and isinstance(dec.func,ast.Attribute) and dec.func.attr in {"get","post","put","patch","delete"} and dec.args:
                    try: path=ast.literal_eval(dec.args[0])
                    except: continue
                    routes.append((dec.func.attr.upper(),path))
            if routes:
                segment="\n".join(lines[node.lineno-1:min(getattr(node,"end_lineno",node.lineno+80),node.lineno+140)])
                sig=ast.unparse(node.args)
                calls=sorted(set(re.findall(r"\b(?:await\s+)?([A-Za-z_]\w*)\(",segment)) - {node.name})
                for method,path in routes: out.append(dict(service=fn,method=method,path=path,handler=node.name,line=node.lineno,signature=sig,calls=calls[:18],source=segment))
    return sorted(out,key=lambda x:(x["service"],x["path"],x["method"]))

def infer(rec):
    s=rec["source"]
    auth=[]
    for key,label in [("session","Session"),("csrf","CSRF"),("api_key","API key"),("job_token","Job token"),("admin","Admin role"),("owner","Owner binding")]:
        if key in s.lower(): auth.append(label)
    errors=sorted(set(re.findall(r"status_code\s*=\s*(\d{3})",s)))
    env=sorted(set(re.findall(r"os\.getenv\([\"']([A-Z0-9_]+)",s)))
    return ", ".join(auth) or "依部署邊界／上游驗證", ", ".join(errors) or "由共用錯誤處理或上游決定", ", ".join(env) or "無端點內直接讀取"

def api_book():
    d=Document(); config(d,"API 端點逐項技術字典"); title(d,"API 端點逐項技術字典","面試用：逐端點來源、函式簽章、驗證、依賴、狀態碼、測試與答辯重點")
    add(d,"本冊由目前程式碼自動抽取 FastAPI route 與函式簽章，再補上安全及測試解讀。它比摘要型規格更接近實際實作；若程式變更，應重新產生並人工審查。")
    records=endpoint_records(); d.add_heading("端點統計",1)
    for svc in SERVICES: add(d,f"{svc}：{sum(1 for r in records if r['service']==svc)} 個 route", "List Bullet")
    for svc in SERVICES:
        rs=[r for r in records if r["service"]==svc]
        d.add_page_break(); d.add_heading(svc,1)
        for idx,r in enumerate(rs,1):
            d.add_heading(f"{idx}. {r['method']} {r['path']}",2)
            auth,errors,env=infer(r)
            fields=[("處理器",f"{r['handler']}（{r['service']}:{r['line']}）"),("函式簽章",r['signature']),("驗證／授權線索",auth),("明確狀態碼",errors),("端點內環境變數",env),("主要呼叫",", ".join(r['calls']) or "無可抽取呼叫")]
            for k,v in fields:
                p=d.add_paragraph(); set_font(p.add_run(k+"："),10,True,"17365D"); set_font(p.add_run(v),9.5)
            d.add_heading("請求與處理流程",3)
            for x in ("解析 path／query／header／cookie／body 或上傳檔案。","執行身分、owner、角色、CSRF、API key、大小與內容檢查。","呼叫本地模型、Job Store、GCS 或外部上游；套用逾時與錯誤正規化。","輸出固定 schema／串流／媒體內容，或建立非同步工作供後續查詢。"):
                add(d,x,"List Bullet")
            d.add_heading("面試可能追問",3)
            for x in (f"為什麼此端點使用 {r['method']}，是否具冪等性？",f"此端點的信任邊界在哪裡，如何避免越權與跨 owner？","上游 timeout、重複送出或部分成功時怎麼補償？","哪些資訊可以進日誌，哪些必須遮蔽？","如何以正常、未授權、邊界輸入、依賴失敗與併發案例驗收？"):
                add(d,x,"List Bullet")
    path=OUT/"11_API端點逐項技術字典_面試版.docx"; d.save(path); return path

def interview_book(figs):
    d=Document(); config(d,"專題面試技術答辯大全"); title(d,"專題面試技術答辯大全","從架構、前端、後端、AI、資料、資安到部署的深度問答與系統圖集")
    chapters=[
    ("系統定位與架構",[("一句話怎麼介紹？","這是一套以 Web SPA 為體驗入口、AI Gateway 為安全邊界，整合會員／商品資料、BASIC／PRO 臉部推論、Ollama 文字建議與 Replicate 妝容渲染的端到端 AI 美妝系統。"),("為什麼要 Gateway？","把 Session、CSRF、上游金鑰、路徑白名單、錯誤格式、媒體授權與稽核集中，避免瀏覽器直連多個服務並暴露秘密；也讓前端只依賴同源穩定契約。"),("單體還是微服務？","前端與 Gateway 是整合核心，臉部、文字、渲染可獨立部署，屬服務化架構。優點是資源與故障隔離；代價是契約、觀測、部署順序與分散式錯誤更複雜。")]),
    ("前端",[("為什麼用 Vanilla JS？","專題規模下可降低建置鏈與框架學習成本，Firebase Hosting 可直接部署；代價是 router.js 容易膨脹，因此需以頁面片段、API 封裝、資料模型與 smoke test 維持邊界。"),("如何管理狀態？","Router 管理當次流程，analysisPackage 是跨頁核心資料，草稿／歷史按 owner 持久化；Session 到期或 owner mismatch 時必須清除私有狀態。"),("怎麼防 XSS？","不信任使用者與上游字串；插入 HTML 前 escape，URL 經 scheme allowlist，圖片來源使用專用驗證，外部連結加 noopener／noreferrer，並以 smoke test 防回歸。")]),
    ("認證與資安",[("Cookie 為什麼 HttpOnly？","降低 XSS 直接竊取 Session 的風險；前端只用 credentials=include，無法讀值。Secure 限 HTTPS，SameSite=Lax 降低跨站攜帶。"),("已有 SameSite 為何還要 CSRF？","SameSite 不是所有情境的完整保證；double-submit token 讓狀態改變請求同時具 Cookie 與 header 證明，並可明確拒絕。"),("owner pinning 是什麼？","伺服器以 Session actor 為準，前端提供的 expected actor 只用來偵測多帳號分頁或舊快取，不能拿來指定資料擁有者。")]),
    ("臉部分析與模型",[("BASIC 和 PRO 差異？","BASIC 以正面單張影像完成主要臉型與五官分類；PRO 要求多角度／側臉品質，提供進階特徵。兩者共用圖片安全、Job、結果契約，但品質門檻與模型不同。"),("CNN、DINOv2、規則怎麼選？","CNN 適合監督式 ROI 分類；DINOv2 使用預訓練視覺表徵與輕量 head，資料較少時可比較；幾何規則可解釋但對姿態與閾值敏感。以身分分組 CV、macro F1、各類 precision／recall、穩定度與部署成本選正式答案。"),("如何避免資料洩漏？","以人物 identity 分組切 train/val/test，不可讓同一人的不同照片跨集合；合併類別後重建切分與類別檔。")]),
    ("非同步與媒體",[("為什麼要 Job？","模型或渲染可能超過一般 HTTP 等待時間。建立 job 回 202，前端以受限輪詢查狀態，服務可控併發、重試、TTL 與清理。"),("如何保護臉照？","上傳先限大小、驗格式、安全解碼、移除 metadata 並重新編碼；GCS 私有、短 TTL、owner 授權，signed URL 短效，保存與刪除都需明確。"),("部分成功怎麼辦？","例如第三方渲染成功但 GCS／Firestore 失敗，要記錄可補償狀態並回收孤兒媒體；刪除需冪等，可重試且逐項回報。")]),
    ("部署與維運",[("如何部署？","服務以 Docker 建置不可變映像，Cloud Run 各自 revision；Gateway 與 Render 有專用 Dockerfile，前端用 Firebase Hosting。發布先測試、staging、少量流量、監控，再擴大。"),("如何回復？","Cloud Run 切回上一個已驗證 revision；模型必須連同 classes.json、mapping 與程式整包回復；Firebase Hosting 回上一 release。"),("看哪些指標？","可用率、p95/p99、冷啟動、上游錯誤、job 排隊／執行時間、跨 owner 拒絕、模型低信心／分布漂移、GCS／Firestore／Replicate 成本。")]),]
    for ch,qa in chapters:
        d.add_heading(ch,1)
        for q,a in qa:
            d.add_heading("Q："+q,2); add(d,"A："+a)
            add(d,"延伸追問：請準備一個實際 bug／權衡／測試證據，說明你如何發現、定位、修正與避免再發。","List Bullet")
    d.add_page_break(); d.add_heading("系統圖集",1)
    for p in figs:
        d.add_heading(p.stem.replace("_"," "),2); d.add_picture(str(p),width=Inches(8.3)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; d.add_page_break()
    d.add_heading("答辯準備清單",1)
    for x in ("能在 60 秒、3 分鐘、10 分鐘三種長度介紹系統。","每個服務說得出輸入、輸出、身分、錯誤、依賴、部署與監控。","能解釋一個模型選擇、一個資安設計、一個分散式失敗與一個跨組協調案例。","Demo 前準備正常流程與依賴故障時的備援說法。","所有數字只引用實測紀錄；沒有證據時明確說是設計目標，不虛構。"):
        add(d,"□ "+x)
    path=OUT/"12_專題面試技術答辯大全_含六大系統圖.docx"; d.save(path); return path

def diagram_book(figs):
    d=Document(); config(d,"系統分析與設計圖集"); title(d,"系統分析與設計圖集","系統架構圖、使用者流程圖、API 串接圖、使用案例圖、甘特圖與活動圖")
    notes=["呈現瀏覽器、Hosting、Gateway、AI 服務、外部資料服務與雲端儲存的責任和信任邊界。","呈現會員從登入到分析、推薦、渲染、比較與保存的完整旅程，以及品質／渲染失敗分支。","呈現 Session、CSRF、上游 API key、owner 授權與錯誤正規化在跨服務呼叫中的位置。","呈現訪客、會員、管理員與外部服務對主要系統能力的互動。","以 14 週專題節奏呈現需求、前端、模型、服務、整合、部署與面試文件工作；可按實際校程調整。","以活動與決策節點呈現臉部分析、品質判斷、建議、渲染及例外流程。"]
    for p,n in zip(figs,notes):
        d.add_heading(p.stem.replace("_"," "),1); add(d,n); d.add_picture(str(p),width=Inches(8.3)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; d.add_page_break()
    path=OUT/"13_系統分析與設計_六大圖集.docx"; d.save(path); return path

def md_sources():
    files=[]
    for base,label in [(ROOT,"後端／AI／整合專案"),(FRONT,"Web 前端專案")]:
        for p in base.glob("*.md"):
            if p.name.lower() not in {"claude.md"}:
                files.append((label,p))
        docs=base/"docs"
        if docs.exists():
            for p in docs.rglob("*.md"):
                if "DecorateMeAI_全端系統技術文件書" not in str(p): files.append((label,p))
    return sorted(files,key=lambda x:(x[0],x[1].name))

def add_markdown(d, text):
    in_code=False; code=[]
    for raw in text.splitlines():
        line=raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.25)
                r=p.add_run("\n".join(code)); r.font.name="Consolas"; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string("404040")
                code=[]; in_code=False
            else: in_code=True
            continue
        if in_code:
            code.append(line); continue
        if not line.strip():
            continue
        m=re.match(r"^(#{1,6})\s+(.*)",line)
        if m:
            level=min(len(m.group(1))+1,3)
            d.add_heading(m.group(2).strip(),level); continue
        if re.match(r"^\s*[-*+]\s+",line):
            add(d,re.sub(r"^\s*[-*+]\s+","",line),"List Bullet"); continue
        if re.match(r"^\s*\d+[.)]\s+",line):
            add(d,line,"List Number"); continue
        if line.lstrip().startswith(">"):
            p=add(d,line.lstrip()[1:].strip()); p.paragraph_format.left_indent=Inches(.25); continue
        if line.strip().startswith("|"):
            p=add(d,line); p.paragraph_format.left_indent=Inches(.2); continue
        # Strip the most common markdown emphasis while retaining code/backticks as visible text.
        clean=re.sub(r"\*\*(.*?)\*\*",r"\1",line)
        add(d,clean)

def md_compendium():
    d=Document(); config(d,"既有 Markdown 技術知識完整彙編"); title(d,"既有 Markdown 技術知識完整彙編","兩個專案的接口、模型、事故、資安、部署、交接與歷史決策原始資料庫")
    sources=md_sources()
    add(d,f"本冊彙整 {len(sources)} 份既有 Markdown。內容保留原始技術敘述，目的是提供可搜尋的面試證據與決策脈絡；若原文件描述的是歷史狀態，仍應以目前程式碼和最新分冊為準。")
    d.add_heading("來源索引",1)
    for label,p in sources: add(d,f"[{label}] {p.name}","List Bullet")
    for i,(label,p) in enumerate(sources,1):
        d.add_page_break(); d.add_heading(f"來源 {i}｜{p.name}",1)
        add(d,f"來源專案：{label}｜原始路徑：{p}")
        try: text=p.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError: text=p.read_text(encoding="cp950",errors="replace")
        add_markdown(d,text)
    path=OUT/"14_既有MD技術知識_完整可搜尋彙編.docx"; d.save(path); return path

def main():
    import sys
    mode = sys.argv[1] if len(sys.argv) > 1 else "all"
    if mode == "figures":
        print("\n".join(str(p) for p in all_figures())); return
    figs=sorted(FIG.glob("*.png")) if mode == "docs" else all_figures()
    if len(figs) != 6:
        raise RuntimeError("請先使用支援 matplotlib 的環境執行 figures 模式")
    docs=[api_book(),interview_book(figs),diagram_book(figs),md_compendium()]
    manifest=OUT/"文件清單.txt"
    current=manifest.read_text(encoding="utf-8") if manifest.exists() else ""
    manifest.write_text(current+"\n\n面試深度補充：\n"+"\n".join(p.name for p in docs)+"\n\n獨立圖檔：\n"+"\n".join("系統圖集_PNG/"+p.name for p in figs),encoding="utf-8")
    print("\n".join(str(x) for x in docs+figs))

if __name__=="__main__": main()
