from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MAIN_SOURCE = ROOT / "專題系統文件書_模型歷程詳版.docx"
MAIN_OUTPUT = ROOT / "專題系統文件書_範本格式_完整詳版.docx"


def set_font(run, name="標楷體", size=12, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill="D9EAD3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.first_child_found_in("w:tcMar")
    if node is None:
        node = OxmlElement("w:tcMar")
        tc_pr.append(node)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        child = node.find(qn(f"w:{name}"))
        if child is None:
            child = OxmlElement(f"w:{name}")
            node.append(child)
        child.set(qn("w:w"), str(value))
        child.set(qn("w:type"), "dxa")


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def write_cell(cell, text, bold=False, size=10):
    p = clear_cell(cell)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, size=9.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, text in enumerate(headers):
        write_cell(table.rows[0].cells[i], text, bold=True, size=size)
        set_cell_shading(table.rows[0].cells[i], "C6E0B4")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            write_cell(cells[i], text, size=size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return table


def add_body(doc, text, *, bold=False, size=12, align=None, indent=True, space_after=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, size=18 if level == 1 else 15 if level == 2 else 12.5, bold=True)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_font(run, size=11.5)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=10)


def add_toc_field(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（開啟文件後更新欄位即可顯示頁碼）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    set_font(run, size=12)
    return p


def enable_field_update(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.3)
    sec.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "標楷體"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for name, size in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 12.5)):
        style = doc.styles[name]
        style.font.name = "標楷體"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")
        style.font.size = Pt(size)
        style.font.bold = True

    for section in doc.sections:
        footer = section.footer
        if not footer.paragraphs:
            fp = footer.add_paragraph()
        else:
            fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.text = ""
        add_page_field(fp)


def classify_existing_paragraphs(doc):
    top_re = re.compile(r"^(摘要|附錄[一二三四五六七八九十]|[壹貳參肆伍陸柒捌玖拾]+、)")
    mid_re = re.compile(r"^(?:[一二三四五六七八九十]+、|[一二三四五六七八九十]+、\s)")
    sub_re = re.compile(r"^(?:\d+\.\d+\s|\d+\.\s)")
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if top_re.match(text):
            p.style = doc.styles["Heading 1"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                set_font(run, size=18, bold=True)
        elif sub_re.match(text):
            p.style = doc.styles["Heading 3"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                set_font(run, size=12.5, bold=True)
        elif mid_re.match(text):
            p.style = doc.styles["Heading 2"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                set_font(run, size=15, bold=True)
        else:
            p.style = doc.styles["Normal"]
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = Cm(0.74)
            for run in p.runs:
                set_font(run, size=12)


def insert_before(doc, target, elements):
    for element in elements:
        # 呼叫端保存的是 CT_P／CT_Tbl，本身就是可插入的 XML 元件。
        target._p.addprevious(element)


def is_conclusion_paragraph(paragraph):
    return re.match(r"^(?:伍|陸)、\s*結論及未來發展", paragraph.text.strip()) is not None


def move_existing_appendix_before_conclusion(doc):
    """修正前一版產物把附錄放在結論後面的順序問題。"""
    appendix = next((p for p in doc.paragraphs if p.text.strip().startswith("附錄一、臉部分析模型完整開發歷程")), None)
    target = next((p for p in doc.paragraphs if is_conclusion_paragraph(p)), None)
    if appendix is None or target is None:
        return
    body = doc._element.body
    if body.index(appendix._p) <= body.index(target._p):
        return
    start = appendix._p
    previous = start.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        if any(node.tag == qn("w:br") for node in previous.iter()):
            start = previous
    elements = []
    current = start
    while current is not None and current.tag != qn("w:sectPr"):
        nxt = current.getnext()
        elements.append(current)
        current = nxt
    for element in elements:
        target._p.addprevious(element)


def add_cover_and_toc(doc, title, subtitle):
    first = doc.paragraphs[0]
    elements = []
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(55)
    cover.paragraph_format.space_after = Pt(16)
    r = cover.add_run("專題系統文件書")
    set_font(r, size=24, bold=True)
    elements.append(cover._p)

    for text, size, gap in [
        (title, 22, 35),
        (subtitle, 15, 42),
        ("組別：＿＿＿＿＿＿＿＿＿＿＿＿", 13, 10),
        ("指導老師：＿＿＿＿＿＿＿＿老師", 13, 10),
        ("成員：＿＿＿＿＿＿＿＿＿＿＿＿", 13, 10),
        ("學年度：＿＿＿＿學年度", 13, 48),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(gap)
        r = p.add_run(text)
        set_font(r, size=size)
        elements.append(p._p)
    add_cover_note = doc.add_paragraph()
    add_cover_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_cover_note.paragraph_format.space_after = Pt(0)
    r = add_cover_note.add_run("（組別、老師與成員請依實際資料填寫）")
    set_font(r, size=10, color=(100, 100, 100))
    elements.append(add_cover_note._p)

    page = doc.add_paragraph()
    page.add_run().add_break(WD_BREAK.PAGE)
    elements.append(page._p)

    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_after = Pt(16)
    r = toc_title.add_run("目錄")
    set_font(r, size=18, bold=True)
    elements.append(toc_title._p)
    toc = add_toc_field(doc)
    elements.append(toc._p)
    page2 = doc.add_paragraph()
    page2.add_run().add_break(WD_BREAK.PAGE)
    elements.append(page2._p)
    insert_before(doc, first, elements)


def add_gap_appendix(doc):
    # 來源文件實際使用「伍、結論及未來發展」，但既有章節規劃曾寫成「陸」；
    # 兩種編號都視為結論章節，確保附錄一定放在結論之前。
    target = next((p for p in doc.paragraphs if is_conclusion_paragraph(p)), None)
    if target is None:
        return
    elements = []

    def h(text, level=1):
        p = doc.add_paragraph(style=f"Heading {level}")
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        set_font(r, size=18 if level == 1 else 15, bold=True)
        elements.append(p._p)

    def b(text):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0.74)
        r = p.add_run(text)
        set_font(r, size=12)
        elements.append(p._p)

    page = doc.add_paragraph()
    page.add_run().add_break(WD_BREAK.PAGE)
    elements.append(page._p)
    h("附錄二、目前我們端無法自行補齊的跨端資料", 1)
    b("這一份專題文件已經把我們自己能從程式、測試與地端模型整理出的內容補上；但有些資料只掌握在各端負責人手上，我不能自行推測。這些資料如果沒有補回來，文件就只能寫到流程和介面，不能負責任地宣稱完整驗收、正式契約或正式上線版本。以下先把缺口列清楚，詳細索取內容另外拆成四份文件交給各端。")
    table = add_table(
        doc,
        ["端別", "目前缺少、我們無法自行確認的內容", "補回來後要放在哪裡"],
        [
            ["Ollama 文字建議端", "正式可驗證的 HTTP 200 request／response、實際模型 tag／量化版、正式 schemaVersion、API key header 名稱與 Cloudflare／Gateway 轉送方式、timeout／錯誤碼、LLaVA 是否真的啟用。", "相關技術、三端契約、UAT、正式端到端驗收"],
            ["演算法端", "每個模型世代的 checkpoint／metrics／fold、資料與 class version、完整 ROI 規格、正式 final model hash、PRO 側面鼻型指標、回饋增量訓練前後比較。", "模型歷程、實驗結果、模型 manifest、訓練工作留"],
            ["爬蟲端", "實際爬取來源、robots／使用規範、欄位 mapping、selector／structured data、去重鍵、圖片授權與失敗重試、最近一次 run 統計、暫存到正式資料的審核規則。", "商品資料來源、爬蟲技術、UAT、商品品質驗收"],
            ["資料庫端", "Firestore collection schema、會員／商品 API 版本與 URL、欄位型別、ID 關聯、索引與分頁、IAM／secret 位置、備份保留、錯誤契約、管理員權限與正式 UAT 證據。", "系統架構、資料庫設計、資安、後台與端到端驗收"],
        ],
        widths=[3.2, 11.2, 6.0],
        size=8.7,
    )
    elements.append(table._tbl)
    h("補件原則", 2)
    for text in [
        "請提供版本、日期與環境，不要只回覆「已完成」；沒有版本就無法和本文件的模型、前端與 Gateway 對齊。",
        "請提供去除 API key、token、cookie、會員 email、真實照片後的範例；密鑰只需告知變數名稱與安全保存位置。",
        "若某一項尚未完成，請直接標示 BLOCKED、NOT RUN 或 NOT DONE，並寫出阻塞原因，不要用假資料補成 PASS。",
        "所有回覆最好帶一個可以核對的證據，例如 commit、Cloud Run revision、runId、測試檔名、截圖編號或 sanitized JSON。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        set_font(r, size=11.5)
        elements.append(p._p)
    insert_before(doc, target, elements)


def build_main():
    doc = Document(MAIN_SOURCE)
    configure_styles(doc)
    classify_existing_paragraphs(doc)
    move_existing_appendix_before_conclusion(doc)
    add_gap_appendix(doc)
    add_cover_and_toc(doc, "Decorate Me AI 個人化彩妝分析與推薦系統", "模型歷程、UAT 與跨端資料補件版")
    enable_field_update(doc)
    doc.save(MAIN_OUTPUT)
    return MAIN_OUTPUT


def build_request_doc(filename, recipient, title, summary, gap_intro, rows, evidence, response_items, handoff):
    path = ROOT / filename
    doc = Document()
    configure_styles(doc)
    cover_elements = []
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("專題系統文件書")
    set_font(r, size=24, bold=True)
    cover_elements.append(p._p)
    for text, size, gap in [
        (title, 20, 28),
        (f"資料提供對象：{recipient}", 14, 18),
        ("系統：Decorate Me AI 個人化彩妝分析與推薦系統", 12.5, 40),
        ("文件用途：跨端資料補件與系統文件彙整", 12.5, 15),
        ("日期：＿＿＿＿年＿＿月＿＿日", 12.5, 52),
    ]:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(gap)
        rr = q.add_run(text)
        set_font(rr, size=size)
        cover_elements.append(q._p)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = note.add_run("請以實際版本與可驗證證據回覆；密鑰、token、cookie 與真實個資不必放進文件。")
    set_font(rr, size=10, color=(100, 100, 100))
    cover_elements.append(note._p)
    pg = doc.add_paragraph()
    pg.add_run().add_break(WD_BREAK.PAGE)
    cover_elements.append(pg._p)
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = toc_title.add_run("目錄")
    set_font(rr, size=18, bold=True)
    cover_elements.append(toc_title._p)
    toc = add_toc_field(doc)
    cover_elements.append(toc._p)
    pg2 = doc.add_paragraph()
    pg2.add_run().add_break(WD_BREAK.PAGE)
    cover_elements.append(pg2._p)

    add_heading(doc, "摘要", 1)
    add_body(doc, summary)
    add_heading(doc, "壹、這份文件為什麼要請你們補", 1)
    add_body(doc, gap_intro)
    add_heading(doc, "貳、目前我們缺少的資料與請提供內容", 1)
    add_body(doc, "請依照下表逐項回覆。若某項目前沒有，請直接寫「尚未完成」以及原因；我們會把它放到專題文件的限制或待完成項目，不會自行猜測補上。")
    table = add_table(doc, ["資料項目", "為什麼需要", "請提供的內容", "建議格式"], rows, widths=[3.5, 5.3, 8.0, 3.2], size=8.4)
    add_heading(doc, "參、請附上的驗收證據", 1)
    for item in evidence:
        add_bullet(doc, item)
    add_heading(doc, "肆、請用這個格式回覆", 1)
    add_body(doc, "請不要只回覆「已完成」。我們需要知道資料是哪一版、在哪個環境、誰確認，以及還有沒有阻塞。可以直接複製以下格式回覆：", indent=False)
    add_table(
        doc,
        ["欄位", "回覆內容"],
        [["負責端／負責人", ""], ["資料版本／日期", ""], ["目前狀態", "PASS／BLOCKED／NOT RUN／NOT DONE"], ["實際內容或連結", ""], ["驗收證據", "commit／runId／revision／測試檔／sanitized JSON／截圖"], ["尚未完成原因", ""], ["需要我們配合的事項", ""], ["可放入文件的說法", "請用一至三句話說明完成狀態與限制"],
        ],
        widths=[5.0, 15.0],
        size=9.2,
    )
    add_heading(doc, "伍、資料交回後會怎麼使用", 1)
    add_body(doc, handoff)
    add_heading(doc, "陸、注意事項", 1)
    for item in [
        "不要把 API key、token、cookie、服務帳號私鑰、會員 email 或真實使用者照片放進回覆文件。",
        "範例 JSON 請去識別化，但欄位名稱、型別、錯誤碼和狀態要保留，否則無法對照前端與 Gateway。",
        "如果正式環境和本機環境不同，請分開寫清楚，不能用本機 PASS 代表 Cloud Run 或正式站 PASS。",
    ]:
        add_bullet(doc, item)

    first = doc.paragraphs[0]
    insert_before(doc, first, cover_elements)
    enable_field_update(doc)
    doc.save(path)
    return path


def build_all():
    outputs = [build_main()]

    outputs.append(build_request_doc(
        "待補資料請求_Ollama_文字建議端.docx",
        "Ollama／Gemma 3 文字建議端",
        "Ollama 文字建議端資料補件與契約驗收單",
        "目前文件可以說明 Ollama 在系統中的角色，但我們還缺少一份由文字建議端確認的正式 request／response、模型版本、認證轉送方式與實際 HTTP 200 證據。沒有這些資料，我們不能把前端的 overall／parts、schemaVersion 或正式端到端結果寫成已完成。",
        "我們可以從本機程式看到部分 endpoint 與錯誤處理，但不知道你們目前真正部署的 model tag、量化版本、正式回應欄位、API key header 以及 Cloudflare Access／Gateway 的實際串法。請協助提供下列資料，讓前端、Gateway、文件和 UAT 依同一份契約對齊。",
        [
            ["正式模型身分", "文件要說明到底使用哪個模型，之後才能追蹤品質變化。", "Gemma 3 的完整 model tag、版本日期、量化方式、context length、是否有 system prompt 版本；若有 LLaVA，請說明是否正式啟用。", "文字＋版本截圖／設定摘要"],
            ["正式服務位置", "本機 8010 wrapper、組員電腦服務、Cloud Run 或外部服務不是同一件事。", "正式 base URL、環境名稱、Cloud Run service／revision；URL 可遮網域，但要保留路由形狀。", "sanitized 設定表"],
            ["認證與轉送", "我們目前只能知道需要 key，不能自行猜 header 或由哪一層代送。", "API key 的環境變數名稱、header 名稱、Cloudflare Access 是否需要 service token、由 Gateway 還是上游驗證；不要提供密鑰值。", "欄位表，不含秘密"],
            ["成功 Request", "前端目前曾送 faceAnalysis、analysisPackage、style、language、userNote、model；要確認正式契約。", "一份去識別化且可重現的 HTTP request JSON／curl，標明必要與選填欄位。", "sanitized JSON／curl"],
            ["成功 Response", "目前沒有取得帶正確授權的 HTTP 200，因此不能宣稱 overall／parts 已完成。", "完整 HTTP 200 response：status、provider、suggestion、renderPromptEn、analysisPackageId、schemaVersion、modelVersion、錯誤時的欄位規則。", "sanitized JSON"],
            ["失敗契約", "上游不可用、授權失敗、格式錯誤的下一步不同，前端需要分辨。", "HTTP status、error code、message、retryable、是否可重試、是否允許 fallback；請列出 timeout 與 rate limit。", "錯誤碼表"],
            ["實際驗收證據", "只有 health 200 不代表建議成功，也不代表 renderPrompt 可以交給渲染端。", "一次帶正確授權的成功 200、同一 analysisPackageId 貫穿、前端顯示結果的截圖或測試 log；敏感值遮蔽。", "log／截圖／測試編號"],
        ],
        [
            "一份帶正確授權的正式 200 response，且 response 內的 analysisPackageId 與 request 相同。",
            "Gemma 3／LLaVA 的實際啟用狀態與版本；若尚未決定，請標示 BLOCKED。",
            "Cloudflare／Gateway 轉送一條成功、一條未授權、一條上游不可用的錯誤案例。",
            "確認前端是否可以移除 splitOllamaTwoPartSuggestion、固定 style.advice 與英文 prompt 展示；若不能，請說阻塞原因。",
        ],
        [
            "這次資料會填入系統文件的 Ollama 技術、三端契約、UAT 與目前限制章節。",
            "前端會依正式 response schema 再決定是否移除假建議與長文切割，不會在未驗證前自行改成已完成。",
            "Gateway 會依 modelVersion、schemaVersion、analysisPackageId 留下可追蹤欄位，供渲染與回饋追查。",
        ],
    ))

    outputs.append(build_request_doc(
        "待補資料請求_演算法_臉部分析模型端.docx",
        "演算法／臉部分析模型端",
        "臉部分析模型歷程、指標與部署成果補件單",
        "目前我們已經把地端能找到的模型世代、資料清理、identity split、五折比較、ConvNeXt 選型、DINOv2 實驗與 training_runs 整理進文件，但仍缺少由演算法端確認的正式版本資料。尤其是 final model、PRO 側面鼻型、每 fold 指標和回饋增量訓練前後比較，不能由我們只看一個 ONNX 檔自行推斷。",
        "我們需要把「我測過哪些模型」和「現在服務真的載入哪個模型」分開寫清楚。請提供每個任務的資料版本、類別版本、訓練設定、評估結果和部署 hash；若某些歷史紀錄已遺失，也請直接註明，不要用目前檔案猜回過去的實驗。",
        [
            ["任務與類別版本", "眉型曾有三類／四類、眼型曾不同合併、鼻型有 BASIC／PRO，混用會讓分數失真。", "face／brow／eye／nose／lip／pro_nose_side 的 class JSON、classVersion、各類樣本數、合併紀錄。", "JSON＋表格"],
            ["資料與身份切分", "random split 可能有身份洩漏，舊 identity map 也曾失效。", "datasetVersion、去重結果、identity map 產生方式、五折索引、holdout 檔名與每 fold 樣本數。", "manifest／摘要"],
            ["訓練設定", "沒有完整參數就無法重現，也不能公平比較模型。", "commit、seed、輸入尺寸、ROI specs、normalization、augmentation、optimizer、batch、epoch、learning rate、sampler、label smoothing。", "run summary"],
            ["模型世代與指標", "文件要能說明為什麼從規則、整臉 CNN、ROI CNN、DINOv2 走到 ConvNeXt。", "每個世代的模型檔／runId、每 fold Macro Accuracy、Balanced Accuracy、Macro-F1、confusion matrix、平均與標準差。", "metrics／圖表"],
            ["現行 final model", "training 完成不等於 runtime 已換模。", "正式模型檔名、SHA-256、manifest version、GCS prefix、Cloud Run revision、promotion ledger、rollback 對應版本。", "manifest＋hash"],
            ["PRO 側面鼻型", "目前有類別檔與模型資產，但缺完整正式評估指標。", "PRO 的資料量、每類數量、切分、指標、已部署 revision、是否和 BASIC feedback 分開。", "完整評估表"],
            ["回饋增量訓練", "使用者修正、管理員審核、worker 匯入和新模型指標要能串起來。", "一個去識別化 jobId→feedbackId→runId 範例、before／after metrics、採用樣本數、promotion 結果。", "流程證據"],
        ],
        [
            "至少一組五個 BASIC 部位同一 protocol 的完整表格，以及每類 confusion matrix。",
            "現行 manifest 指定的模型檔 hash 與 Cloud Run revision 對照。",
            "一個成功 run、一個失敗 run 與一個回滾／prepare-only run 的摘要。",
            "如果某個模型只存在本機或只是實驗／shadow，請明確標示，避免文件寫成正式線上模型。",
        ],
        [
            "這次資料會直接放進模型歷程章節、模型比較表、地端成果位置、限制與未來工作。",
            "版本資料會補到 manifest／promotion 的說明，讓 UAT 能驗證實際載入的模型不是只看本地檔案。",
            "若指標缺失，我們會在文件標示 BLOCKED／待補，不會用 acceptance rate 或其他任務數字替代。",
        ],
    ))

    outputs.append(build_request_doc(
        "待補資料請求_爬蟲_商品資料端.docx",
        "爬蟲／商品資料端",
        "商品爬蟲、資料清理與上架流程補件單",
        "目前文件只能確認系統使用 Requests、BeautifulSoup 或結構化資料取得商品，並且先進暫存區再由管理員審核；但商品來源、實際欄位 mapping、去重規則、圖片授權、爬蟲排程與最近一次結果不是我們端可以自行補齊的資料。這些缺口會直接影響商品推薦與後台 UAT。",
        "請提供爬蟲真正使用的來源與資料契約，而不是只提供一張商品頁截圖。系統要能說清楚商品從哪裡來、抓到什麼、怎麼判斷重複、什麼情況會被退回，以及目前資料是否可以被推薦服務使用。",
        [
            ["資料來源與使用規範", "沒有來源、robots／terms 和授權紀錄，就不能對外宣稱爬蟲合法且可長期執行。", "來源網域／頁面類型、robots.txt 與使用限制、抓取頻率、圖片與商品資訊使用規則。", "來源清單"],
            ["欄位 mapping", "爬到的名稱、品牌、價格、色號、圖片與來源 URL 必須和商品資料庫欄位對得上。", "原始欄位→標準欄位 mapping、必填欄位、資料型別、缺值處理、色號／LAB 欄位。", "欄位表"],
            ["selector 與 fallback", "網站 HTML 變動時，爬蟲可能成功回 200 但抓到空值或錯欄位。", "structured data 優先順序、CSS／XPath selector、備援 selector、欄位驗證與失敗條件。", "設定摘要"],
            ["去重與更新", "同一商品重抓不能產生多筆，價格／庫存更新也不能覆蓋錯商品。", "唯一鍵、canonical URL、SKU／色號去重規則、upsert 行為、下架／缺貨處理與時間欄位。", "規則＋範例"],
            ["圖片與色號", "商品推薦和 CIEDE2000 需要可信的圖片、色號與 LAB；外部圖片還有權限與 SSRF 風險。", "圖片來源、下載／代理方式、授權、尺寸與格式、LAB 如何取得、失敗時是否拒絕上架。", "欄位＋流程"],
            ["暫存與人工審核", "爬蟲資料不能直接混入正式推薦，否則錯商品會出現在使用者端。", "staging collection／狀態、管理員採用／退回欄位、審核者、審核時間、退回原因、正式寫入路徑。", "schema＋截圖"],
            ["最近一次執行結果", "需要知道爬蟲現在是真的能跑，還是只有程式存在。", "最近 run 日期、來源數、成功／失敗／跳過數、常見錯誤、最後寫入筆數與可重跑方式。", "run summary"],
        ],
        [
            "一筆去識別化的原始商品資料、標準化資料、審核採用後資料三階段對照。",
            "一個重複商品、缺圖片、缺價格、網站欄位改變時的失敗案例。",
            "最近一次爬蟲 run log 或 summary，標示實際環境與資料庫版本。",
            "管理員在後台預覽、採用、退回、重新抓取的 UAT 操作與結果。",
        ],
        [
            "這些內容會補進相關技術、商品資料庫、爬蟲流程、後台操作與推薦限制章節。",
            "商品資料端的欄位 mapping 會和推薦端、CIEDE2000 與前端顯示欄位互相核對。",
            "沒有通過人工審核的爬蟲資料會在文件中標示為暫存，不會寫成正式商品已上架。",
        ],
    ))

    outputs.append(build_request_doc(
        "待補資料請求_資料庫_會員商品資料端.docx",
        "資料庫／會員與商品服務端",
        "會員、商品、收藏與後台資料庫契約補件單",
        "目前文件可以描述 Firestore、GCS、會員、商品、收藏、分析紀錄與管理員後台的大致關係，也已經記錄部分 Gateway 錯誤處理；但我們缺少資料庫端確認的正式 schema、API 版本、權限、索引、備份與可驗收的端到端證據。這些內容不能由前端畫面反推。",
        "請提供不含個資與秘密的資料契約，讓我們確認前端欄位、Gateway 轉送、後台顯示、模型回饋與商品資料是不是同一套定義。若會員服務、商品服務、Firestore 與 GCS 分屬不同負責人，請在回覆中標明邊界。",
        [
            ["資料來源與版本", "目前環境可能有本機、Cloud Run、外部會員／商品 API，URL 指錯會被誤判成登入失效。", "服務名稱、環境、base URL／route 形狀、API／schema version、Cloud Run revision 或 commit。", "環境對照表"],
            ["Firestore schema", "後台要知道每個 collection 的欄位型別、狀態和關聯，不能靠畫面猜。", "collections、document ID、欄位型別、必填／選填、createdAt／updatedAt、jobId／feedbackId／runId 關係。", "schema 表／JSON"],
            ["會員與 owner", "圖片、渲染工作與回饋必須歸屬正確 owner，且不能把 email 洩漏給前端。", "ownerId 產生方式、會員查詢、訪客識別、本人／管理員權限、跨會員讀取限制。", "流程＋欄位"],
            ["商品、收藏與歷史", "前端 404／405、重複收藏、刪除或版本衝突都需要穩定契約。", "CRUD endpoints、HTTP method、request／response、分頁／排序、idempotency、409／422／404 行為。", "OpenAPI／sanitized JSON"],
            ["GCS 與 IAM", "模型回饋影像、妝前／妝後圖的保存與讀取不能讓瀏覽器直接拿金鑰。", "bucket／prefix、物件命名、metadata、保存期限、服務帳號角色、private／proxy 規則；不要提供私鑰。", "政策摘要"],
            ["索引、分頁與限制", "資料一多，後台列表與樣本掃描可能變慢或拿錯資料。", "Firestore indexes、limit 上限、cursor／descending 規則、GCS list 上限、單次下載大小。", "設定表"],
            ["備份與錯誤契約", "名冊讀不到、資料庫掛掉、刪除失敗時不能把空資料當成正常結果。", "backup／restore、retention、錯誤 code、retryable、transaction／對帳方式、刪除失敗保留規則。", "規格＋測試"],
            ["完整 UAT 證據", "健康檢查 200 不等於註冊、登入、收藏、商品管理與後台流程真的完成。", "測試帳號類型（不要給密碼）、成功／失敗案例、前後端版本、截圖／log、尚未執行的項目。", "UAT 表／log"],
        ],
        [
            "一份 sanitized schema／OpenAPI 或欄位表，至少涵蓋會員、商品、收藏、歷史、feedback、training runs。",
            "一條會員登入→分析紀錄→收藏／歷史，以及一條管理員商品審核→正式資料的完整 UAT 證據。",
            "會員名冊不可達、商品上游 401／403、資料不存在 404、版本衝突 409 的實際錯誤契約。",
            "GCS 私有權限、服務端 proxy 與刪除／保留政策的版本化摘要。",
        ],
        [
            "這些資料會補進系統架構、資料庫設計、後台、資安、UAT 和目前限制章節。",
            "Gateway 會依你們確認的 endpoint／method／error code 更新轉送與前端顯示對照。",
            "完整 E2E 若仍缺測試 session 或下游設定，文件會標示 NOT RUN，不會拿 health check 代替。",
        ],
    ))

    print("OUTPUTS")
    for path in outputs:
        check = Document(path)
        print(f"{path} | paragraphs={len(check.paragraphs)} tables={len(check.tables)}")


if __name__ == "__main__":
    build_all()
